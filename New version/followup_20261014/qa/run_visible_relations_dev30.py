"""Frozen legacy controls plus observable-domain hypotheses; score after locking."""
from pathlib import Path
from collections import Counter, defaultdict
import argparse
import hashlib
import os
import platform
import resource
import sys
import time

import openpyxl

ROOT = Path(__file__).resolve().parents[1]; REPO = ROOT.parents[1]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'qa'),
               str(ROOT / 'baseline/source/HybridGuard_V5_handoff_20260904/src')]
from run_selector_dev5 import read, save, sha
from build_candidates_dev16 import upstream
from formulaguard.workbook import WorkbookModel
from visible_relations_dev30 import predict as infer, VERSION
from run_string_execution_dev28 import CASES

RUN = ROOT / 'development/visible_relations_dev30_v1'
OLD = ROOT / 'development/candidate_expansion_dev18_data_v1'
CONTROL_TYPES = ['clean', 'operator_error', 'reference_error', 'coefficient_error',
                 'marked_adjustment', 'unmarked_adjustment', 'conflicting_instruction', 'duplicate_role']


def prepare():
    assert not RUN.exists()
    RUN.mkdir()
    sources = [Path(__file__), ROOT / 'scripts/visible_relations_dev30.py',
               ROOT / 'scripts/relation_algebra_dev30.py', ROOT / 'tests/test_visible_relations_dev30.py',
               ROOT / 'VISIBLE_RELATIONS_DEV30_PREREGISTRATION_20261018.md',
               ROOT / 'scripts/build_candidates_dev16.py', ROOT / 'scripts/selector_dev11.py']
    sources += sorted((ROOT / 'baseline/source').rglob('*.py'))
    save(RUN / 'source_lock.json', dict(version=VERSION, files={str(p.relative_to(ROOT)): sha(p) for p in sources}))
    index = [{**item, 'stratum': 'old213'} for item in read(OLD / 'public_index.json')]
    public = ROOT / 'development/public_candidate_pipeline_dev10'
    for item in read(public / 'public_index.json'):
        report = public / 'predictions' / (item['id'] + '.json.gz')
        status = read(report)['status']
        index.append(dict(case='public/' + item['id'], id=item['id'], stratum='public_original',
                          input=os.path.relpath(REPO / item['relative_path'], ROOT),
                          input_sha256=item['sha256'], upstream_status=status,
                          origin_group=item['origin'], public_report=str(report.relative_to(ROOT))))
    for identity, path in CASES:
        folder = {'cv500': 'natural_cv500_dev27_v1', 'ldas': 'natural_ldas_dev27_v1',
                  'hobyah': 'natural_hobyah_dev26_v1'}[identity]
        index.append(dict(case='natural/' + identity, id=identity, stratum='natural_original',
                          input=str(path.relative_to(ROOT)), input_sha256=sha(path), origin_group=identity,
                          public_report='development/' + folder + '/prediction.json.gz'))
    for item in read(ROOT / 'development/record_candidates_dev29_v1/index.json')[:16]:
        index.append(dict(case='string_control/' + item['id'], id=item['id'], stratum='string_control',
                          input=item['path'], input_sha256=item['sha256'], origin_group='authored_string_control'))
    truths = []
    for kind in CONTROL_TYPES:
        book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Ledger'
        sheet['A2'] = 'Total Cash Inflows'; sheet['A3'] = 'Total Cash Outflows'; sheet['A4'] = 'Net Cash Flow'
        expected, observed, changes = {}, {}, []
        for col in range(2, 10):
            letter = openpyxl.utils.get_column_letter(col)
            sheet[letter + '2'] = 100 + col * 11; sheet[letter + '3'] = 20 + col; sheet[letter + '5'] = 4 + col
            correct = '=' + letter + '2-' + letter + '3'; actual = correct
            if kind in ['operator_error', 'conflicting_instruction', 'duplicate_role']:
                actual = '=' + letter + '2+' + letter + '3'
            if kind == 'reference_error':
                actual = '=' + letter + '2-' + letter + '5'
            if kind == 'coefficient_error':
                actual = '=(' + correct[1:] + ')*1.1'
            if kind in ['marked_adjustment', 'unmarked_adjustment']:
                correct = actual = '=(' + correct[1:] + ')+10'
            address = letter + '4'; label = 'Ledger!' + address
            sheet[address] = actual; expected[label] = correct; observed[label] = actual
            if actual != correct:
                changes.append(dict(cell=label, original_formula=correct, injected_formula=actual))
        if kind == 'marked_adjustment':
            sheet['A7'] = 'Exception: a one-off adjustment of 10 applies to each period.'
        if kind == 'conflicting_instruction':
            sheet['A7'] = 'Net Cash Flow equals Total Cash Inflows plus Total Cash Outflows.'
        if kind == 'duplicate_role':
            sheet['A8'] = 'Total Cash Inflows'
            for col in range(2, 10):
                sheet.cell(8, col, 200 + col)
        path = RUN / 'controls' / (kind + '.xlsx'); path.parent.mkdir(exist_ok=True); book.save(path)
        index.append(dict(case='new_control/' + kind, id=kind, stratum='new_control',
                          input=str(path.relative_to(ROOT)), input_sha256=sha(path), origin_group='authored_role_control'))
        truths.append(dict(case='new_control/' + kind, expected_formulas=expected, observed_formulas=observed,
                           errors=[r['cell'] for r in changes], injection_records=changes,
                           scope='development-authored intent fixed before predictions; not natural or unseen final data'))
    assert len(index) == 251
    save(RUN / 'public_index.json', index)
    save(RUN / 'control_truth.json', dict(cases=truths, final_validation_eligible=False))
    save(RUN / 'preparation_lock.json', dict(index_sha256=sha(RUN / 'public_index.json'),
         source_lock_sha256=sha(RUN / 'source_lock.json'), control_truth_sha256=sha(RUN / 'control_truth.json'),
         inputs={i['input']: i['input_sha256'] for i in index},
         old_index_sha256=sha(OLD / 'public_index.json')))
    print('DEV30 prepared 251 fixed whole-workbook cases and eight intent controls', flush=True)


def verify():
    lock = read(RUN / 'preparation_lock.json')
    assert sha(RUN / 'source_lock.json') == lock['source_lock_sha256']
    assert sha(RUN / 'public_index.json') == lock['index_sha256']
    for rel, digest in {**read(RUN / 'source_lock.json')['files'], **lock['inputs']}.items():
        assert sha(ROOT / rel) == digest, rel


def baseline(item):
    if item['stratum'] == 'old213':
        if item['group'] == 'public_relations':
            pack = read(ROOT / item['upstream']['public_pack'])
            result, observed = pack['dev11'], pack['observed_formulas']
        else:
            models, observed = upstream(item); result = models['dev11']
    elif item['stratum'] in ('public_original', 'natural_original'):
        pack = read(ROOT / item['public_report'])
        if item['stratum'] == 'natural_original':
            result, observed = pack['dev11'], None
        else:
            result, observed = pack['predictions']['dev6'], None
            # Public dev11 qualified the same dev6 proposals in its frozen run.
            path = ROOT / 'development/selector_dev11_complete_ast_v1/predictions/public' / (item['id'] + '.json.gz')
            result = read(path)
    else:
        return [], None
    rows = [dict(cell=r['cell'], status='propose', suggested_formula=r['suggested_formula'])
            for r in result['rows'] if r.get('status') == 'propose' and r.get('suggested_formula')
            and (observed is None or r['suggested_formula'] != observed[r['cell']])]
    return rows, observed


def prediction(item):
    if item['stratum'] == 'public_original' and item['upstream_status'] != 'inference_complete_unscored':
        return dict(case=item['case'], status='retained_original_rejection', upstream_status=item['upstream_status'],
                    formula_count=None, dev11=[], relation_only=[], combined=[], diagnostics=[], labels_read=[])
    model = WorkbookModel.from_xlsx(ROOT / item['input'])
    base, observed = baseline(item)
    if observed is not None:
        assert observed == {s + '!' + a: f for (s, a), f in model.formulas.items()}
    result = infer(model)
    old = {r['cell'] for r in base}
    combined = base + [r for r in result['suggestions'] if r['cell'] not in old]
    assert len({r['cell'] for r in combined}) == len(combined)
    return dict(case=item['case'], status='complete', formula_count=len(model.formulas),
                dev11=base, relation_only=result['suggestions'], combined=combined,
                diagnostics=result['rows'], labels_read=[], original_workbook_modified=False)


def predict():
    verify(); assert not (RUN / 'prediction_lock.json').exists()
    active, accesses = [True], set()
    def guard(event, args):
        if not active[0] or event != 'open' or not args or not isinstance(args[0], (str, bytes, os.PathLike)):
            return
        name = os.fsdecode(args[0])
        if any(x in name.lower() for x in ['truth', 'training_labels', 'case_labels', 'author_diff',
                                          'product_sum_oracle', 'prepare_multisource', 'prepare_candidate',
                                          'certification', 'after.xlsx', 'natural_score', 'candidate_analysis']):
            raise PermissionError('DEV30 prediction cannot read scoring evidence: ' + name)
        if name.endswith('.xlsx') or name.endswith('.xlsm'):
            accesses.add(name)
    sys.addaudithook(guard)
    wall, cpu = time.perf_counter(), time.process_time(); files = {}; summaries = []
    for item in read(RUN / 'public_index.json'):
        result = prediction(item)
        path = RUN / 'predictions' / (item['case'].replace('/', '__') + '.json.gz'); save(path, result)
        files[str(path.relative_to(RUN))] = sha(path)
        summaries.append(dict(case=item['case'], status=result['status'], formulas=result['formula_count'],
                              suggestions={k: len(result[k]) for k in ['dev11', 'relation_only', 'combined']}))
        print(item['case'], result['status'], summaries[-1]['suggestions'], flush=True)
    save(RUN / 'runtime.json', dict(cpu_seconds=time.process_time() - cpu, wall_seconds=time.perf_counter() - wall,
         peak_rss_kib=resource.getrusage(resource.RUSAGE_SELF).ru_maxrss, python=platform.python_version(), gpu_used=False,
         scope='whole-file role extraction and domain hypothesis inference; immutable paper/dev11 predictions reused',
         cases=summaries))
    files['runtime.json'] = sha(RUN / 'runtime.json')
    save(RUN / 'prediction_lock.json', dict(files=files, labels_read=[], input_accesses=sorted(accesses),
         source_lock_sha256=sha(RUN / 'source_lock.json'), preparation_lock_sha256=sha(RUN / 'preparation_lock.json'),
         main_candidate_adopted=False, final_validation_eligible=False))
    active[0] = False; verify()
    print('DEV30 predictions locked; labels never read during inference', flush=True)


def totals(rows):
    keys = ['suggestions', 'correct', 'wrong', 'unknown', 'errors', 'correct_errors', 'unnecessary']
    result = {k: sum(r[k] for r in rows) for k in keys}
    result['precision_confirmed_lower_bound'] = result['correct'] / result['suggestions'] if result['suggestions'] else None
    result['coverage'] = result['correct_errors'] / result['errors'] if result['errors'] else None
    return result


def score():
    verify(); target = RUN / 'score.json'; assert not target.exists()
    lock = read(RUN / 'prediction_lock.json')
    for rel, digest in lock['files'].items():
        assert sha(RUN / rel) == digest
    from product_sum_oracle_dev17 import equivalent as public_eq
    from prepare_annotation_dev14 import equivalent as annotation_eq
    old = read(OLD / 'heldout_truth.json.gz')['cases']
    controls = {r['case']: r for r in read(RUN / 'control_truth.json')['cases']}
    natural = read(ROOT / 'development/natural_cv500_dev27_v1/natural_truth.json.gz')
    string_truth = {r['id']: r for r in read(ROOT / 'development/record_candidates_dev29_v1/controls_truth.json.gz')['cases']}
    cases = []
    for item in read(RUN / 'public_index.json'):
        pred = read(RUN / 'predictions' / (item['case'].replace('/', '__') + '.json.gz'))
        if item['stratum'] == 'old213':
            truth = old[item['case']]; expected = truth['expected_formulas']; errors = set(truth['errors'])
            workbook = openpyxl.load_workbook(ROOT / item['input']) if item['group'] == 'public_relations' else None
            if item['group'] == 'public_relations':
                eq = lambda a, b, c: public_eq(a, b, c.rsplit('!', 1)[0], workbook)
            elif item['group'] == 'lcoe':
                eq = lambda a, b, c: annotation_eq(a, b)
            else:
                eq = lambda a, b, c: a == b
        elif item['stratum'] == 'new_control':
            truth = controls[item['case']]; expected = truth['expected_formulas']; errors = set(truth['errors'])
            workbook = openpyxl.load_workbook(ROOT / item['input'])
            eq = lambda a, b, c: public_eq(a, b, c.rsplit('!', 1)[0], workbook)
        elif item['stratum'] == 'natural_original' and item['id'] == 'cv500':
            expected = {r['cell']: r['accepted_formula'] for r in natural['cases']}; errors = set(expected)
            eq = lambda a, b, c: a == b
        elif item['stratum'] == 'string_control':
            truth = string_truth[item['id']]; expected = truth['all_expected']; errors = set(truth['errors'])
            eq = lambda a, b, c: a == b
        else:
            expected, errors = {}, set(); eq = lambda a, b, c: False
        record = dict(case=item['case'], stratum=item['stratum'], origin=item['origin_group'], status=pred['status'],
                      family=truth['family'] if item['stratum'] == 'old213' else item['stratum'],
                      scenario=str(truth.get('scenario')) if item['stratum'] == 'old213' else item['id'])
        for name in ['dev11', 'relation_only', 'combined']:
            correct, wrong, unknown, unnecessary = [], [], [], []
            for row in pred[name]:
                cell = row['cell']
                if cell not in expected:
                    unknown.append(cell)
                elif cell not in errors:
                    wrong.append(cell); unnecessary.append(cell)
                elif eq(row['suggested_formula'], expected[cell], cell):
                    correct.append(cell)
                else:
                    wrong.append(cell)
            metric = dict(suggestions=len(pred[name]), correct=len(correct), wrong=len(wrong), unknown=len(unknown),
                          errors=len(errors), correct_errors=len(set(correct)), unnecessary=len(unnecessary),
                          correct_cells=correct, wrong_cells=wrong, unknown_cells=unknown)
            metric.update(precision_confirmed_lower_bound=len(correct) / len(pred[name]) if pred[name] else None,
                          coverage=len(set(correct)) / len(errors) if errors else None)
            record[name] = metric
        cases.append(record)
    summary = {}
    for stratum in sorted({r['stratum'] for r in cases}):
        summary[stratum] = {name: totals([r[name] for r in cases if r['stratum'] == stratum])
                            for name in ['dev11', 'relation_only', 'combined']}
    origins = {origin: {name: totals([r[name] for r in cases if r['stratum'] == 'old213' and r['origin'] == origin])
                        for name in ['dev11', 'relation_only', 'combined']}
               for origin in sorted({r['origin'] for r in cases if r['stratum'] == 'old213'})}
    breakdowns = {field: {key: {name: totals([r[name] for r in cases if r['stratum'] == 'old213' and r[field] == key])
                                   for name in ['dev11', 'relation_only', 'combined']}
                           for key in sorted({r[field] for r in cases if r['stratum'] == 'old213'})}
                  for field in ['family', 'scenario']}
    save(target, dict(cases=cases, by_stratum=summary, old213_by_origin=origins,
         old213_by_family=breakdowns['family'], old213_by_scenario=breakdowns['scenario'],
         prediction_lock_sha256=sha(RUN / 'prediction_lock.json'),
         old_truth_sha256=sha(OLD / 'heldout_truth.json.gz'), control_truth_sha256=sha(RUN / 'control_truth.json'),
         mode='revealed development hypothesis evaluation; no new-origin heldout claim',
         main_candidate_adopted=False, final_validation_complete=False, accepted_final_sources=0, accepted_final_families=0))
    for group, rows in summary.items():
        print('DEV30', group, rows, flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('phase', choices=['prepare', 'predict', 'score'])
    globals()[parser.parse_args().phase]()
