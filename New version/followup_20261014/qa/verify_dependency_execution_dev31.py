"""Independent inventory/graph/reference checks plus full deterministic replay."""
from collections import Counter
from pathlib import Path
import json
import math
import sys

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries, get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'qa'),
               str(ROOT / 'baseline/source/HybridGuard_V5_handoff_20260904/src')]
from run_selector_dev5 import read, save, sha
from run_dependency_execution_dev31 import RUN, verify_inputs
from dependency_execution_dev31 import Executor
from execution_reference import formula_declarations


def independently_expand(path, targets):
    book = openpyxl.load_workbook(path, data_only=False)
    sheets = {s.title.casefold(): s for s in book}
    formulas = {(s.title, c.coordinate): c.value for s in book for row in s for c in row if c.data_type == 'f'}
    cells, functions = set(), Counter()
    pending = list(targets)
    while pending:
        key = pending.pop()
        if key in cells: continue
        cells.add(key)
        if key not in formulas: continue
        for token in Tokenizer(formulas[key]).items:
            if token.type == 'FUNC' and token.subtype == 'OPEN':
                functions[token.value[:-1].upper()] += 1
            if token.type != 'OPERAND' or token.subtype != 'RANGE': continue
            value = token.value
            if '!' in value:
                sheet, value = value.rsplit('!', 1)
                if sheet.startswith("'"): sheet = sheet[1:-1].replace("''", "'")
            else: sheet = key[0]
            sheet = sheets[sheet.casefold()].title
            lo_col, lo_row, hi_col, hi_row = range_boundaries(value)
            assert all(v is not None for v in [lo_col, lo_row, hi_col, hi_row])
            pending.extend((sheet, get_column_letter(c) + str(r)) for r in range(lo_row, hi_row + 1)
                           for c in range(lo_col, hi_col + 1))
    book.close()
    return cells, cells & formulas.keys(), dict(functions)


def compare(actual, reference):
    if actual['kind'] != reference['kind']:
        return 'type_mismatch'
    if actual['kind'] == 'error':
        return 'same_error' if actual['error'] == reference['error'] else 'different_error'
    if actual['kind'] == 'number':
        return 'agreement' if math.isclose(actual['value'], reference['value'], rel_tol=1e-8, abs_tol=1e-8) else 'disagreement'
    if actual['kind'] in ['boolean', 'text']:
        return 'agreement' if actual['value'] == reference['value'] else 'disagreement'
    return 'not_comparable'


def main():
    verify_inputs(); execution = read(RUN / 'execution_lock.json'); ref_lock = read(RUN / 'reference_lock.json')
    assert sha(RUN / 'execution_lock.json') == ref_lock['execution_lock_sha256']
    for lock in [execution, ref_lock]:
        for rel, digest in lock['files'].items(): assert sha(RUN / rel) == digest, rel
    assert execution['labels_read'] == [] and execution['reference_outputs_read'] == []
    summary, replays, inventory_checks, all_unsupported = [], 0, 0, 0
    index = read(RUN / 'index.json'); graph = None
    for item in index:
        result = read(RUN / 'executions' / (item['id'] + '.json.gz'))
        if item['retained_rejection']:
            assert result['status'] == 'retained_original_rejection' and result['cells'] == []
            summary.append(dict(id=item['id'], retained_rejection=True, formulas=None)); continue
        path = ROOT / item['input']; raw = formula_declarations(path)
        engine = Executor.from_xlsx(path)
        assert set(raw) == {s + '!' + a for s, a in engine.model.formulas}, item['id']
        assert len(result['cells']) == len(raw); inventory_checks += len(raw)
        for row in result['cells']:
            s, a = row['cell'].rsplit('!', 1)
            replay = json.loads(json.dumps(engine.execute((s, a))))
            assert row == replay, (item['id'], row['cell'])
            replays += 1
        counts = Counter('executed_' + r['value']['kind'] if r['status'] == 'executed' else r['reason']
                         for r in result['cells'])
        assert dict(counts) == result['counts']
        unsupported = sum(r['status'] == 'unsupported' for r in result['cells']); all_unsupported += unsupported
        summary.append(dict(id=item['id'], formulas=len(raw), counts=dict(counts), unsupported=unsupported))
        if item['id'] == 'public10':
            targets = [('Debt', get_column_letter(c) + '10') for c in range(6, 45)]
            cells, formulas, functions = independently_expand(path, targets)
            observed, observed_formulas, hidden = set(), set(), set()
            for key in targets:
                audit = engine.audit(key); observed.update(audit['cells'])
                observed_formulas.update(audit['formula_cells']); hidden.update(audit['hidden_cells'])
            assert cells == observed and formulas == observed_formulas
            assert len(cells) == 317 and len(formulas) == 166
            actual_functions = Counter()
            for key in sorted(formulas): actual_functions.update(engine.parsed(key)[3])
            assert dict(actual_functions) == functions
            graph = dict(cells=len(cells), formulas=len(formulas), functions=functions,
                         targets=len(targets), hidden_cells=len(hidden), independent_graph_exact_agreement=True,
                         input_cells=len(cells - formulas), blank_missing_cells=[list(k) for k in sorted(cells - formulas - engine.inputs.keys())])
        print('DEV31 replay', item['id'], len(raw), 'formulas; unsupported', unsupported, flush=True)
    comparisons = []
    for identity in ['controls', 'public10']:
        result = read(RUN / 'executions' / (identity + '.json.gz'))
        reference = read(RUN / 'references' / (identity + '_uno.json.gz'))
        assert reference['status'] == 'complete' and reference['forced_calculate_all'] and reference['read_only']
        predictions = {r['cell']: r for r in result['cells']}
        for cell, raw in reference['cells'].items():
            row = predictions[cell]
            if row['status'] == 'unsupported':
                comparisons.append(dict(id=identity, cell=cell, comparison='retained_unsupported', reason=row['reason']))
                continue
            status = compare(row['value'], raw)
            comparisons.append(dict(id=identity, cell=cell, comparison=status, actual=row['value'],
                                    reference={k: raw[k] for k in ['kind', 'value', 'error'] if k in raw}))
    disagreements = [r for r in comparisons if r['comparison'] not in ['agreement', 'same_error', 'retained_unsupported']]
    save(RUN / 'independent_comparison.json', dict(rows=comparisons, disagreements=disagreements,
         counts=dict(Counter(r['comparison'] for r in comparisons)), graph=graph,
         full_replay_formulas=replays, independent_inventory_checks=inventory_checks,
         all_unsupported_retained=all_unsupported, workbook_saved=False, business_truth_claimed=False))
    verify_inputs()
    final = dict(passed=not disagreements, files=summary, graph=graph, full_replay_formulas=replays,
         independent_formula_inventory_checks=inventory_checks, comparisons=len(comparisons),
         independent_agreements=sum(r['comparison']=='agreement' for r in comparisons),
         same_errors=sum(r['comparison']=='same_error' for r in comparisons),
         reference_unsupported=sum(r['comparison']=='retained_unsupported' for r in comparisons),
         disagreements=disagreements, main_candidate_changed=False, repair_improvement_established=False,
         final_validation_eligible=False, accepted_final_sources=0, accepted_final_families=0,
         execution_lock_sha256=sha(RUN / 'execution_lock.json'), reference_lock_sha256=sha(RUN / 'reference_lock.json'),
         comparison_sha256=sha(RUN / 'independent_comparison.json'))
    save(ROOT / 'qa/DEPENDENCY_EXECUTION_DEV31_VERIFICATION_20261018.json', final)
    print(json.dumps({k: v for k, v in final.items() if k != 'files'}, ensure_ascii=False, indent=2), flush=True)
    assert not disagreements, 'Preserve failed v1 evidence; any change requires new registered version.'


if __name__ == '__main__':
    main()
