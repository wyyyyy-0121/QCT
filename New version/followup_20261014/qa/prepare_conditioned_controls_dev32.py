"""Author-only fixed controls and intent labels; never import during inference."""
from collections import deque
import openpyxl

PROFIT = ['clean','operator_error','reference_error','coefficient_error','marked_share','unmarked_share',
          'marked_adjustment','unmarked_adjustment','conflicting_instruction','duplicate_role',
          'hidden_input','unsupported_input','rate_outside','unknown_output']
DEBT = ['clean','operator_error','reference_error','coefficient_error','marked_adjustment','unmarked_adjustment',
        'positive_principal','duplicate_role','hidden_upstream','conflicting_instruction','non_ppmt','unsupported_upstream']


def create(family, kind, ordinal):
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Records' if family=='profit' else 'Schedule'
    expected, observed, changes = {}, {}, []
    if family=='profit':
        names = deque(['Units','Unit Price','Unit Cost','Shipping','Discount Rate','Tax Rate','Net Profit'])
        names.rotate(ordinal % len(names))
        bindings = {name:openpyxl.utils.get_column_letter(col) for col,name in enumerate(names,3)}
        for name,col in bindings.items():sheet[col+'4'] = name
        for row in range(7,13):
            values={'Units':row+4,'Unit Price':40+row,'Unit Cost':8,'Shipping':3,'Discount Rate':.1,'Tax Rate':.03}
            for name,value in values.items():sheet[bindings[name]+str(row)] = value
            q,p,c,s,d,t=(bindings[n]+str(row) for n in ['Units','Unit Price','Unit Cost','Shipping','Discount Rate','Tax Rate'])
            correct=f'={q}*{p}*(1-{d})*(1-{t})-{q}*{c}-{s}'; actual=correct
            if kind in ['operator_error','conflicting_instruction','duplicate_role','hidden_input','unsupported_input','rate_outside','unknown_output']:
                actual=correct.replace('-'+q+'*'+c,'+'+q+'*'+c)
            if kind=='reference_error':actual=correct.replace('*'+p,'*'+c,1)
            if kind=='coefficient_error':actual='=('+correct[1:]+')*1.19'
            if kind in ['marked_share','unmarked_share']:
                correct=actual=f'={q}*{p}*(1-{d})*(1-{t})*0.73-{q}*{c}-{s}'
            if kind in ['marked_adjustment','unmarked_adjustment']:correct=actual='=('+correct[1:]+')+17'
            if kind=='unsupported_input':sheet[q]='=SQRT(121)'
            if kind=='rate_outside':sheet[d]=1.2
            address=bindings['Net Profit']+str(row);sheet[address]=actual
            label=sheet.title+'!'+address;expected[label]=correct
            if actual!=correct:changes.append(dict(cell=label,original_formula=correct,injected_formula=actual))
        if kind=='marked_share':sheet['P18']='Contract: revenue share is 0.73.'
        if kind=='marked_adjustment':sheet['P18']='Adjustment: add 17 to each record.'
        if kind=='conflicting_instruction':sheet['P18']='Net Profit equals Units times Unit Price plus Units times Unit Cost.'
        if kind=='duplicate_role':sheet['O4']='Units'
        if kind=='hidden_input':sheet.column_dimensions[bindings['Unit Cost']].hidden=True
        if kind=='unknown_output':sheet[bindings['Net Profit']+'4']='Custom Metric'
    else:
        rates=book.create_sheet('Parameters');rates['B2']=.04;rates['C2']=.05;rates['E2']=1
        rates['G2']='=SQRT(0.0016)' if kind=='unsupported_upstream' else '=INDEX(B2:C2,E2)'
        for row,title in [(9,'Opening Balance'),(12,'Drawdowns'),(15,'Principal Repayment'),(18,'Closing Balance')]:
            sheet.cell(row,2,title);sheet.cell(row,3,'USD')
        start=6+ordinal%3
        for per,col in enumerate(range(start,start+6),1):
            c=openpyxl.utils.get_column_letter(col)
            sheet[c+'9']=900;sheet[c+'12']=80;sheet[c+'16']=33
            payment=f'=PPMT(Parameters!G2,{per},6,900)'
            if kind=='positive_principal':payment='=-PPMT(Parameters!G2,'+str(per)+',6,900)'
            if kind=='non_ppmt':payment='=-160'
            sheet[c+'15']=payment
            correct=f'=SUM({c}9,{c}12,{c}15)';actual=correct
            if kind in ['operator_error','duplicate_role','hidden_upstream','conflicting_instruction','unsupported_upstream']:
                actual=f'={c}9+{c}12-{c}15'
            if kind=='reference_error':actual=f'={c}9+{c}12+{c}16'
            if kind=='coefficient_error':actual='=('+correct[1:]+')*1.19'
            if kind in ['marked_adjustment','unmarked_adjustment']:correct=actual='=('+correct[1:]+')+17'
            address=c+'18';sheet[address]=actual;label=sheet.title+'!'+address;expected[label]=correct
            if actual!=correct:changes.append(dict(cell=label,original_formula=correct,injected_formula=actual))
        if kind=='marked_adjustment':sheet['B22']='Adjustment: add 17 to each period.'
        if kind=='duplicate_role':sheet['B25']='Drawdowns';sheet['C25']='USD'
        if kind=='hidden_upstream':rates.column_dimensions['C'].hidden=True
        if kind=='conflicting_instruction':sheet['B22']='Closing Balance subtracts Principal Repayment.'
    observed={s.title+'!'+c.coordinate:c.value for s in book for row in s for c in row if c.data_type=='f'}
    return book,dict(expected_formulas=expected,observed_formulas=observed,
                     errors=[r['cell'] for r in changes],injection_records=changes,
                     family=family,scenario=kind,scope='authored development local intent; upstream formulas not business-certified')
