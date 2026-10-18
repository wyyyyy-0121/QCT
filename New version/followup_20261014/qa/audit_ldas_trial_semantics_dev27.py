"""Independent constrained arithmetic execution and complete 4-box mask audit."""
from pathlib import Path
import ast
import itertools
import re
import sys
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'qa')]
from run_selector_dev5 import read,save,sha
import run_ldas_natural_dev27 as runner
RUN=runner.RUN


class Evaluator:
    def __init__(self,path,overrides):
        self.book=openpyxl.load_workbook(path,data_only=False);self.sheet=self.book.active
        self.names={name:next(value.destinations)[1].replace('$','') for name,value in self.book.defined_names.items()}
        self.overrides=overrides;self.cache={};self.active=set()

    def cell(self,name):
        key=self.names.get(name,name)
        if key in self.cache:return self.cache[key]
        assert key not in self.active,'circular formula'
        self.active.add(key);v=self.overrides.get(key,self.sheet[key].value)
        if isinstance(v,str) and v.startswith('='):
            text=re.sub(r'(?<![<>=!])=(?!=)','==',v[1:].replace('$',''))
            v=self.node(ast.parse(text,mode='eval').body)
        self.active.remove(key);self.cache[key]=v;return v

    def node(self,n):
        if isinstance(n,ast.Constant):return n.value
        if isinstance(n,ast.Name):return self.cell(n.id)
        if isinstance(n,ast.BinOp):
            a,b=self.node(n.left),self.node(n.right);a=0 if a is None else a;b=0 if b is None else b
            if isinstance(n.op,ast.Add):return a+b
            if isinstance(n.op,ast.Sub):return a-b
            if isinstance(n.op,ast.Mult):return a*b
        if isinstance(n,ast.Compare) and len(n.ops)==1 and isinstance(n.ops[0],ast.Eq):
            a,b=self.node(n.left),self.node(n.comparators[0]);return ('' if a is None else a)==('' if b is None else b)
        if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and not n.keywords:
            if n.func.id=='IF' and len(n.args)==3:return self.node(n.args[1] if self.node(n.args[0]) else n.args[2])
            if n.func.id=='SUM':return sum(0 if (v:=self.node(a)) is None else v for a in n.args)
        raise ValueError('unsupported independent execution node '+ast.dump(n))


def trial_rows(path):
    sheet=openpyxl.load_workbook(path,data_only=False).active
    rows=[]
    for row in range(1,sheet.max_row+1):
        forms=[sheet[f'{c}{row}'].value for c in ['B','D','F','H']]
        if all(isinstance(f,str) and f.startswith('=') for f in forms):rows.append(row)
    assert len(rows)==72
    return rows


def main():
    runner.engine.verify();diff=read(RUN/'author_diff.json.gz');results=[]
    before=RUN/'raw/before.xlsx';after=RUN/'raw/after.xlsx'
    rows_before=trial_rows(before);rows_after=trial_rows(after)
    subjects=['C11','E11','G11','I11'];cols=['B','D','F','H']
    # Cross-check the independent interpreter against stored caches on every
    # original numbered trial. Caches corroborate evaluation, not correctness.
    cache_checks=0
    for path,rows in [(before,rows_before),(after,rows_after)]:
        model=Evaluator(path,{})
        cache=openpyxl.load_workbook(path,data_only=True).active
        for row in rows:
            for col in cols:
                cell=f'{col}{row}';assert model.cell(cell)==cache[cell].value,(path.name,cell);cache_checks+=1
    for mask in itertools.product([0,1],repeat=4):
        if not any(mask):continue
        overrides={cell:(100+n if active else None) for n,(cell,active) in enumerate(zip(subjects,mask),1)}
        # Separate the actual missing input from formula logic: this restored
        # input is a counterfactual and never modifies the natural raw file.
        overrides['C10']=1
        old=Evaluator(before,overrides);new=Evaluator(after,overrides)
        rows=[]
        for version,model,indices in [('before',old,rows_before),('after',new,rows_after)]:
            count=0;bad=[];expected=0
            for r in indices:
                for c,active in zip(cols,mask):
                    if not active:continue
                    expected+=1;value=model.cell(f'{c}{r}');count+=1
                    if value!=expected:bad.append(dict(cell=f'{c}{r}',actual=value,expected=expected))
            rows.append(dict(version=version,active_trials=count,mismatches=len(bad),first_examples=bad[:4]))
        prefix=list(mask)==[1]*sum(mask)+[0]*(4-sum(mask))
        results.append(dict(mask=list(mask),within_documented_rightmost_empty_domain=prefix,versions=rows))
    in_domain=[r for r in results if r['within_documented_rightmost_empty_domain']]
    assert len(in_domain)==4 and all(v['mismatches']==0 for r in in_domain for v in r['versions'])
    old_raw=Evaluator(before,{});new_raw=Evaluator(after,{})
    assert old_raw.cell('C10') is None and new_raw.cell('C10')==1
    assert old_raw.cell('B15')==0 and new_raw.cell('B15')==1
    actual_different=[]
    for r in rows_before:
        # LBOX1=missing makes each first-box counter one smaller; this is
        # an input omission, not proof that its referencing formula is wrong.
        value=old_raw.cell(f'B{r}')
        if value!=(old_raw.sheet[f'A{r}'].value-1)*4+1:actual_different.append(f'B{r}')
    assert len(actual_different)==72
    target=RUN/'trial_semantics_audit.json';assert not target.exists()
    save(target,dict(independent_cached_trial_values_checked=cache_checks,subject_mask_cases=results,
        documented_valid_mask_count=4,all_documented_masks_both_versions_correct_after_input_completion=True,
        missing_original_input='ExperimentalWorksheet_v2_sample!C10',corrected_input_value=1,
        actual_first_box_counters_affected=actual_different,natural_formula_error_labels_established=0,
        corrected_version_not_assumed_oracle=True,all_subjects_in_template_are_non_numeric_placeholders=True,
        counterfactuals_not_natural_error_samples=True,precision=None,coverage=None,
        author_diff_sha256=sha(RUN/'author_diff.json.gz'),prediction_lock_sha256=sha(RUN/'prediction_lock.json'),
        auditor_sha256=sha(Path(__file__)),final_validation_eligible=False))
    print('LDAS independently checked',cache_checks,'cached trial values; 4 documented masks correct after completing missing input; no formula labels',flush=True)
    print('All masks',[(r['mask'],[v['mismatches'] for v in r['versions']]) for r in results],flush=True)


if __name__=='__main__':main()
