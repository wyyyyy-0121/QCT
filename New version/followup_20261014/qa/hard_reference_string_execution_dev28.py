"""Force a read-only calculation; preserve conversion-only evidence separately.

Independent execution compatibility and tokenizer audit, never repair scoring.
Full public-workbook text remains in ignored local storage.
"""
from pathlib import Path
import hashlib
import re
import subprocess
import sys

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'qa'),
               str(ROOT / 'baseline/source/HybridGuard_V5_handoff_20260904/src')]
from run_selector_dev5 import read, save, sha
from run_string_execution_dev28 import RUN, CASES, value_record
from run_public_execution_audit import run_worker
from string_execution_dev28 import parse, execute, provenance, Unsupported
from typed_execution_adapter import load_typed_inputs, ExcelError


def verify_execution():
    source = read(RUN / 'source_lock.json')
    for rel, digest in source['files'].items():
        assert sha(ROOT / rel) == digest, rel
    for name, path in CASES:
        assert sha(path) == source['inputs'][name], name
    lock = read(RUN / 'execution_lock.json')
    assert sha(RUN / 'source_lock.json') == lock['source_lock_sha256']
    for rel, digest in lock['files'].items():
        assert sha(RUN / rel) == digest, rel
    assert sha(RUN / 'local/executed_text.json.gz') == lock['local_text_sha256']


def reference_value(raw):
    if raw['kind'] == 'error':
        return ExcelError(raw['error'])
    return raw.get('value')


def token_references(formula, current_sheet):
    """Use the independent openpyxl lexer, preserving duplicate range tokens."""
    refs = []
    for token in Tokenizer(formula).items:
        if (token.type, token.subtype) != ('OPERAND', 'RANGE'):
            continue
        target = token.value
        if '!' in target:
            sheet, target = target.rsplit('!', 1)
            if sheet.startswith("'"):
                sheet = sheet[1:-1].replace("''", "'")
        else:
            sheet = current_sheet
        match = re.fullmatch(r'(\$?)([A-Za-z]+)(\$?)([0-9]+)', target)
        assert match, (formula, token.value)
        refs.append(dict(sheet=sheet, address=match[2].upper() + match[4],
                         column_absolute=bool(match[1]), row_absolute=bool(match[3])))
    return refs


def check_provenance(formula, sheet):
    actual = provenance(parse(formula), sheet)
    selected = [{k: r[k] for k in ['sheet', 'address', 'column_absolute', 'row_absolute']}
                for r in actual]
    assert selected == token_references(formula, sheet), formula
    assert [r['formula_offset'] for r in actual] == sorted(r['formula_offset'] for r in actual)
    for r in actual:
        # Offset must point at the reference itself (including a sheet qualifier).
        suffix = formula[r['formula_offset']:]
        assert suffix.startswith(("'", '$')) or suffix[0].isalpha(), (formula, r)
    return len(actual)


def smoke():
    path = RUN / 'local/hard_reference_smoke.xlsx'
    assert not path.exists()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Main'
    other = book.create_sheet("O'Brien")
    inputs = {'A1': 'x"好🙂', 'A2': 12, 'A3': True, 'A4': False,
              'A6': '#VALUE!', 'A7': 'line1\nline2', 'A8': '#VALUE!', 'A9': 1.25}
    for cell, value in inputs.items():
        sheet[cell] = value
    sheet['A6'].data_type = 's'
    sheet['A8'].data_type = 'e'
    other['A1'] = 'across'
    other['B2'] = 23
    formulas = {
        'B1': '=CONCATENATE("say """,A1,"""!")',
        'B2': '=("a"&CONCATENATE("b",(A1&"d")))&"e"',
        'B3': '=A2&A3&A4&-12', 'B4': '=A5&"blank"', 'B5': '=A5',
        'B6': '=A6&" text"', 'B7': '=A8&" error"', 'B8': '=A7&"🙂"',
        'B9': "='O''Brien'!$A1&'O''Brien'!B$2", 'B10': '=A2',
        'B11': '=""', 'B12': '="fake A100 and Sheet1!B9"&A1&A1',
        'B13': '=""&999999999999999', 'B14': '=""&-999999999999999',
        'B15': '=0', 'B16': '=CONCATENATE("outer",CONCATENATE("inner",A1))',
        'B17': '="_x000D_"&A7', 'B18': '="_x005F_x000D_"',
        # Explicit unsupported controls must remain in the audit.
        'B19': '=A9&" decimal"', 'B20': '=B1&" dependency"',
        'B21': '=SUM(A2,1)', 'B22': '=A2+1', 'B23': '=A1:A2',
        'B24': '=CONCAT(A1,A2)',
    }
    for cell, formula in formulas.items():
        sheet[cell] = formula
    book.save(path)
    before = sha(path)
    typed = load_typed_inputs(path)
    inventory = {('Main', cell): formula for cell, formula in formulas.items()}
    # Parse actual serialized formula declarations, not just the creation dict.
    loaded = openpyxl.load_workbook(path, data_only=False)
    assert {c.coordinate: c.value for row in loaded['Main'] for c in row if c.data_type == 'f'} == formulas
    ref = run_worker(path, ['Main!' + c for c in formulas], timeout=60)
    assert ref['status'] == 'complete', ref
    save(RUN / 'local/hard_reference_smoke_uno.json.gz', ref)
    rows = []
    for cell, formula in formulas.items():
        raw = ref['cells']['Main!' + cell]
        row = dict(cell=cell, formula=formula, reference=value_record(reference_value(raw)))
        try:
            node = parse(formula)
            row['references_checked'] = check_provenance(formula, 'Main')
            value = execute(node, 'Main', typed, formula_cells=inventory, sheet_names=book.sheetnames)
            # Compare typed results: Python bool must not match numeric 1/0.
            expected = reference_value(raw)
            row.update(status='executed', actual=value_record(value),
                       exact_agreement=(type(value) is type(expected) and value == expected)
                       or (type(value) in (int, float) and type(expected) in (int, float) and value == expected))
        except Unsupported as error:
            row.update(status='unsupported', reason=error.code)
        rows.append(row)
    assert sha(path) == before
    return dict(rows=rows, source_unchanged=True, source_sha256=before,
                forced_calculate_all=ref['forced_calculate_all'], read_only=ref['read_only'],
                workbook_saved=ref['workbook_saved'],
                executed=sum(r['status'] == 'executed' for r in rows),
                agreements=sum(r.get('exact_agreement', False) for r in rows),
                unsupported=sum(r['status'] == 'unsupported' for r in rows),
                reference_wall_seconds=ref['wall_seconds'])


def main():
    target = RUN / 'hard_recalculation_reference.json'
    assert not target.exists()
    verify_execution()
    auditor_lock = RUN / 'hard_reference_source_lock.json'
    assert not auditor_lock.exists()
    sources = [Path(__file__), ROOT / 'scripts/libreoffice_reference_worker.py',
               ROOT / 'scripts/run_public_execution_audit.py']
    save(auditor_lock, dict(files={str(p.relative_to(ROOT)): sha(p) for p in sources},
         execution_lock_sha256=sha(RUN / 'execution_lock.json'),
         protocol='read-only isolated UNO import, enableAutomaticCalculation, calculateAll, never save',
         conversion_diagnostic_sha256=sha(RUN / 'independent_reference.json'),
         comparison='exact text or typed scalar; no newline/OOXML escape normalization'))
    original = CASES[0][1]
    before = sha(original)
    local = [r for r in read(RUN / 'local/executed_text.json.gz')['rows'] if r['case'] == 'cv500']
    ref = run_worker(original, [r['cell'] for r in local], timeout=60)
    assert ref['status'] == 'complete', ref
    assert ref['forced_calculate_all'] and ref['read_only'] and not ref['workbook_saved']
    save(RUN / 'local/hard_reference_cv500_uno.json.gz', ref)
    book = openpyxl.load_workbook(original, data_only=False)
    rows = []
    for r in local:
        raw = ref['cells'][r['cell']]
        sheet, cell = r['cell'].rsplit('!', 1)
        rows.append(dict(cell=r['cell'], exact_agreement=raw['kind'] == 'text' and raw['value'] == r['value'],
                         reference=value_record(reference_value(raw)),
                         references_checked=check_provenance(book[sheet][cell].value, sheet)))
    assert len(rows) == 505
    assert sha(original) == before
    small = smoke()
    verify_execution()
    for rel, digest in read(auditor_lock)['files'].items():
        assert sha(ROOT / rel) == digest
    result = dict(rows=rows, compared=len(rows), exact_text_agreements=sum(r['exact_agreement'] for r in rows),
                  reference_wall_seconds=ref['wall_seconds'], smoke=small,
                  input_sha256=before, original_input_unchanged=True,
                  read_only=True, forced_calculate_all=True, workbook_saved=False,
                  source_lock_sha256=sha(auditor_lock),
                  preserved_conversion_only_diagnostic=dict(compared=505, agreements=500,
                       accepted_as_recalculation_evidence=False,
                       reason='format conversion did not explicitly request full recalculation; five newline mismatches'),
                  libreoffice_version=subprocess.check_output(['/usr/bin/libreoffice', '--version'], text=True).strip(),
                  business_truth_claimed=False, repair_suggestions_generated=False, final_validation_eligible=False)
    save(target, result)
    print('DEV28 hard calculateAll exact agreement', result['exact_text_agreements'], '/', len(rows), flush=True)
    print('DEV28 independent smoke', small['agreements'], '/', small['executed'],
          'executed;', small['unsupported'], 'retained unsupported', flush=True)
    assert result['exact_text_agreements'] == 505
    assert small['agreements'] == small['executed'] == 18 and small['unsupported'] == 6


if __name__ == '__main__':
    main()
