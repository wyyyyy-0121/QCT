"""Independent Python-AST interpreter for auditing dev23 witnesses/identities.

No production polynomial parser, arithmetic, cache, or normalization is used.
"""
import ast
from fractions import Fraction
import re
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries, get_column_letter
import sympy


class IndependentNumeric:
    def __init__(self, workbook, expected):
        self.workbook = workbook; self.expected = expected
        self.sheets = {name.casefold(): name for name in workbook.sheetnames}
        self.parsed = {}

    def formula(self, text, sheet, inputs=None, symbolic=False, stack=()):
        key = (text, sheet)
        if key not in self.parsed:
            pieces = []; operands = {}
            for token in Tokenizer(text).items:
                if token.type == 'WHITE-SPACE': continue
                if token.type == 'OPERAND':
                    assert token.subtype in ['NUMBER', 'RANGE']
                    name = 'operand_'+str(len(operands)); operands[name] = (token.subtype, token.value); pieces.append(name)
                elif token.type == 'FUNC' and token.subtype == 'OPEN':
                    assert token.value.casefold() == 'sum('; pieces.append('SUM(')
                elif token.type == 'FUNC' and token.subtype == 'CLOSE': pieces.append(')')
                elif token.type == 'PAREN': pieces.append('(' if token.subtype == 'OPEN' else ')')
                elif token.type == 'SEP' and token.subtype == 'ARG': pieces.append(',')
                elif token.type == 'OPERATOR-PREFIX' and token.value in ['+', '-']: pieces.append(token.value)
                elif token.type == 'OPERATOR-INFIX' and token.value in ['+', '-', '*', '/']: pieces.append(token.value)
                else: raise ValueError('unsupported independent token')
            self.parsed[key] = ast.parse(''.join(pieces), mode='eval').body, operands
        tree, operands = self.parsed[key]

        def scalar(value):
            if value is None: return sympy.Integer(0) if symbolic else Fraction(0)
            if isinstance(value, (str, list, bool)): raise ValueError('nonnumeric scalar')
            return value

        def reference(text):
            target_sheet = sheet
            if '!' in text:
                target_sheet, text = text.rsplit('!', 1)
                if target_sheet.startswith("'"): target_sheet = target_sheet[1:-1].replace("''", "'")
            assert '[' not in target_sheet and ':' not in target_sheet
            target_sheet = self.sheets[target_sheet.casefold()]
            text = text.replace('$', '').upper()
            assert re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*(?::[A-Z]{1,3}[1-9][0-9]*)?', text)
            left, top, right, bottom = range_boundaries(text)
            assert (right-left+1)*(bottom-top+1) <= 256
            values = []
            for row in range(top, bottom+1):
                for col in range(left, right+1):
                    address = f'{get_column_letter(col)}{row}'; name = target_sheet+'!'+address
                    assert name not in stack and len(stack) < 64
                    cell = self.workbook[target_sheet][address]
                    value = self.expected.get(name, cell.value)
                    if name in self.expected or isinstance(value, str) and value.startswith('='):
                        value = self.formula(value, target_sheet, inputs, symbolic, stack+(name,))
                    elif cell.data_type == 'e' or isinstance(value, bool): raise ValueError('error or boolean')
                    elif type(value) in (int, float): value = sympy.Symbol(name) if symbolic else inputs[name]
                    elif value is not None and not isinstance(value, str): raise ValueError('non numeric leaf')
                    values.append(value)
            return values if ':' in text else values[0]

        def visit(node):
            if isinstance(node, ast.Name):
                kind, text = operands[node.id]
                if kind == 'RANGE': return reference(text)
                return sympy.Rational(text) if symbolic else Fraction(text)
            if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
                value = scalar(visit(node.operand)); return value if isinstance(node.op, ast.UAdd) else -value
            if isinstance(node, ast.BinOp):
                a, b = scalar(visit(node.left)), scalar(visit(node.right))
                if isinstance(node.op, ast.Add): return a+b
                if isinstance(node.op, ast.Sub): return a-b
                if isinstance(node.op, ast.Mult): return a*b
                if isinstance(node.op, ast.Div): return a/b
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == 'SUM' and not node.keywords:
                total = sympy.Integer(0) if symbolic else Fraction(0)
                for argument in node.args:
                    value = visit(argument)
                    for member in value if isinstance(value, list) else [value]:
                        if member is None or isinstance(member, str): continue
                        total += scalar(member)
                return total
            raise ValueError('unsupported independent AST')
        return scalar(visit(tree))
