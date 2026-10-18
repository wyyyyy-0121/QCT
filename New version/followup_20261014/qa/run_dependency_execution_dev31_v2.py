"""Fixed all-original execution audit; lock before independent recalculation."""
from collections import Counter
from pathlib import Path
import argparse
import os
import platform
import resource
import sys
import time

import openpyxl

ROOT = Path(__file__).resolve().parents[1]; REPO = ROOT.parents[1]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'qa'),
               str(ROOT / 'baseline/source/HybridGuard_V5_handoff_20260904/src')]
from run_selector_dev5 import read, save, sha
from dependency_execution_dev31_v2 import Executor, VERSION
from run_public_execution_audit import run_worker

RUN = ROOT / 'development/dependency_execution_dev31_v2'
PUBLIC = ROOT / 'development/public_candidate_pipeline_dev10'


def controls():
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Cases'
    source = book.create_sheet('Input data')
    for address, value in {'A1': 10, 'B1': 20, 'A2': 30, 'B2': 40, 'A3': '123',
                           'B3': True, 'A5': '=1/0', 'B5': '=12', 'C1': 2}.items():
        source[address] = value
    source.row_dimensions[2].hidden = True
    formulas = []
    for rate in ['0', '0.05', '-0.02', '0.000000000001']:
        for when in [0, 1]:
            for fv in [0, 200]:
                key = f'rate={rate};type={when};fv={fv}'
                formulas.append((key + ';pmt', f'=PMT({rate},12,1000,{fv},{when})'))
                for per in [1, 6, 12]:
                    formulas.append((key + f';ppmt={per}', f'=PPMT({rate},{per},12,1000,{fv},{when})'))
    # Independently check aggregate amortization rather than only isolated values.
    for when in [0, 1]:
        terms = ','.join(f'PPMT(0.05,{per},12,1000,0,{when})' for per in range(1, 13))
        formulas.append((f'principal_sum;type={when}', '=SUM(' + terms + ')'))
    formulas += [
        ('index_horizontal', "=INDEX('Input data'!A1:B1,2)"),
        ('index_vertical', "=INDEX('Input data'!A1:A2,2)"),
        ('index_2d', "=INDEX('Input data'!A1:B2,2,2)"),
        ('index_computed_position', "=INDEX('Input data'!A1:B1,'Input data'!C1)"),
        ('index_lazy_error', "=INDEX('Input data'!A5:B5,2)"),
        ('index_blank', "=INDEX('Input data'!A4:B4,1)"),
        ('index_out_of_range', "=INDEX('Input data'!A1:B1,3)"),
        ('index_zero_unsupported', "=INDEX('Input data'!A1:B1,0)"),
        ('index_fraction_unsupported', "=INDEX('Input data'!A1:B1,1.5)"),
        ('index_array_unsupported', "=INDEX('Input data'!A1:B2,1)"),
        ('lazy_if', '=IF(1,12,1/0)'), ('lazy_false_if', '=IF(0,1/0,12)'),
        ('if_error', "=IF('Input data'!A5,12,13)"),
        ('if_omitted_false', '=IF(0,12)'),
        ('range_sum_ignore_types', "=SUM('Input data'!A1:B4)"),
        ('range_max_ignore_types', "=MAX('Input data'!A1:B4)"),
        ('range_min_ignore_types', "=MIN('Input data'!A1:B4)"),
        ('reference_sum_ignore_types', "=SUM('Input data'!A3,'Input data'!B3)"),
        ('range_error', "=SUM('Input data'!A1:A5)"),
        ('tiny_denominator', '=1/0.00000000000001'), ('zero_denominator', '=1/0'),
        ('negative_fractional_power', '=(-1)^0.5'), ('numeric_comparison', '=10>20'),
        ('text_numeric_unsupported', "='Input data'!A3+1"),
        ('unknown_function_inactive', '=IF(1,12,SQRT(9))'),
        ('chained_power_unsupported', '=2^3^2'),
        ('financial_invalid_period', '=PPMT(0.05,13,12,1000)'),
        ('financial_fractional_period', '=PPMT(0.05,1.5,12,1000)'),
        ('financial_rate_minus_one', '=PMT(-1,12,1000)'),
    ]
    for row, (name, formula) in enumerate(formulas, 1):
        sheet.cell(row, 1, name); sheet.cell(row, 2, formula)
    row = len(formulas) + 1
    sheet.cell(row, 1, 'inactive_circular'); sheet.cell(row, 2, f'=IF(1,12,B{row})')
    return book


def prepare():
    assert not RUN.exists(), 'Preserve evidence; register a new run for changes.'
    RUN.mkdir()
    path = RUN / 'controls.xlsx'; controls().save(path)
    index = []
    for item in read(PUBLIC / 'public_index.json'):
        pack = read(PUBLIC / 'predictions' / (item['id'] + '.json.gz'))
        index.append(dict(id=item['id'], input=os.path.relpath(REPO / item['relative_path'], ROOT),
                          sha256=item['sha256'], origin=item['origin'],
                          retained_rejection=pack['status'] != 'inference_complete_unscored',
                          old_status=pack['status']))
    index.append(dict(id='controls', origin='authored_execution_controls', input=str(path.relative_to(ROOT)),
                      sha256=sha(path), retained_rejection=False))
    sources = [Path(__file__), ROOT / 'scripts/dependency_execution_dev31.py',
               ROOT / 'tests/test_dependency_execution_dev31.py',
               ROOT / 'qa/verify_dependency_execution_dev31.py',
               ROOT / 'DEPENDENCY_EXECUTION_DEV31_PREREGISTRATION_20261018.md',
               ROOT / 'DEPENDENCY_EXECUTION_DEV31_COUNT_ADDENDUM_20261018.md',
               ROOT / 'scripts/typed_execution_adapter.py', ROOT / 'scripts/run_selector_dev5.py',
               ROOT / 'scripts/run_public_execution_audit.py', ROOT / 'scripts/libreoffice_reference_worker.py']
    sources += [ROOT / 'scripts/dependency_execution_dev31_v2.py',
                ROOT / 'qa/run_dependency_execution_dev31.py',
                ROOT / 'qa/verify_dependency_execution_dev31_v2.py',
                ROOT / 'tests/test_dependency_execution_dev31_v2.py',
                ROOT / 'DEPENDENCY_EXECUTION_DEV31_V2_REGISTRATION_20261018.md']
    sources += sorted((ROOT / 'baseline/source').rglob('*.py'))
    old = [ROOT / 'development' / folder / file for folder, file in [
        ('dependency_execution_dev31_v1', 'execution_lock.json'),
        ('dependency_execution_dev31_v1', 'reference_lock.json'),
        ('visible_relations_dev30_v1', 'prediction_lock.json'),
        ('string_execution_dev28_v2', 'execution_lock.json'),
        ('selector_dev11_complete_ast_v1', 'prediction_lock.json'),
        ('nested_calibration_dev25_v1/outer', 'prediction_lock.json')]]
    save(RUN / 'index.json', index)
    save(RUN / 'source_lock.json', dict(version=VERSION, files={str(p.relative_to(ROOT)): sha(p) for p in sources},
                                      old_locks={str(p.relative_to(ROOT)): sha(p) for p in old}))
    save(RUN / 'preparation_lock.json', dict(source_lock_sha256=sha(RUN / 'source_lock.json'),
         index_sha256=sha(RUN / 'index.json'), inputs={i['input']: i['sha256'] for i in index}))
    print('DEV31 prepared', len(index), 'files including 3 retained rejections', flush=True)


def verify_inputs():
    lock = read(RUN / 'preparation_lock.json'); source = read(RUN / 'source_lock.json')
    assert sha(RUN / 'source_lock.json') == lock['source_lock_sha256']
    assert sha(RUN / 'index.json') == lock['index_sha256']
    for rel, digest in {**source['files'], **source['old_locks'], **lock['inputs']}.items():
        assert sha(ROOT / rel) == digest, rel


def execute():
    verify_inputs(); assert not (RUN / 'execution_lock.json').exists()
    active = [True]; accesses = set()
    def guard(event, args):
        if active[0] and event == 'open' and args and isinstance(args[0], (str, bytes, os.PathLike)):
            name = os.fsdecode(args[0])
            if any(s in name.lower() for s in ['truth', 'after.xlsx', 'natural_score', 'author_diff',
                                              'natural_pair_docs', '_uno.json', 'control_labels']):
                raise PermissionError('DEV31 cannot read truth/reference: ' + name)
            if name.endswith(('.xlsx', '.xlsm')): accesses.add(name)
    sys.addaudithook(guard)
    wall, cpu = time.perf_counter(), time.process_time(); files, summaries = {}, []
    for item in read(RUN / 'index.json'):
        if item['retained_rejection']:
            result = dict(id=item['id'], status='retained_original_rejection', old_status=item['old_status'], cells=[])
        else:
            engine = Executor.from_xlsx(ROOT / item['input'])
            rows = [engine.execute(k) for k in sorted(engine.model.formulas)]
            result = dict(id=item['id'], status='complete', cells=rows,
                          counts=dict(Counter('executed_' + r['value']['kind'] if r['status'] == 'executed'
                                              else r['reason'] for r in rows)))
        target = RUN / 'executions' / (item['id'] + '.json.gz'); save(target, result)
        files[str(target.relative_to(RUN))] = sha(target)
        summaries.append(dict(id=item['id'], formulas=len(result['cells']), status=result['status'],
                              counts=result.get('counts', {})))
        print('DEV31', summaries[-1], flush=True)
    save(RUN / 'runtime.json', dict(wall_seconds=time.perf_counter()-wall, cpu_seconds=time.process_time()-cpu,
         peak_rss_kib=resource.getrusage(resource.RUSAGE_SELF).ru_maxrss, gpu_used=False,
         python=platform.python_version(), files=summaries))
    files['runtime.json'] = sha(RUN / 'runtime.json')
    save(RUN / 'execution_lock.json', dict(files=files, input_accesses=sorted(accesses), labels_read=[],
         reference_outputs_read=[], repair_suggestions_generated=False, final_validation_eligible=False,
         source_lock_sha256=sha(RUN / 'source_lock.json'), preparation_lock_sha256=sha(RUN / 'preparation_lock.json')))
    active[0] = False; verify_inputs()
    print('DEV31 executions locked before independent reference', flush=True)


def reference():
    verify_inputs(); assert not (RUN / 'reference_lock.json').exists()
    lock = read(RUN / 'execution_lock.json')
    for rel, digest in lock['files'].items(): assert sha(RUN / rel) == digest
    index = {i['id']: i for i in read(RUN / 'index.json')}; files = {}
    for identity in ['controls', 'public10']:
        output = read(RUN / 'executions' / (identity + '.json.gz'))
        if identity == 'controls':
            cells = [r['cell'] for r in output['cells']]
        else:
            target = [('Debt', openpyxl.utils.get_column_letter(c) + '10') for c in range(6, 45)]
            engine = Executor.from_xlsx(ROOT / index[identity]['input']); dependency = set()
            for key in target: dependency.update(engine.audit(key)['formula_cells'])
            cells = [s + '!' + a for s, a in sorted(dependency)]
            save(RUN / 'packt_target_graph.json', dict(targets=target, formula_cells=sorted(dependency)))
            files['packt_target_graph.json'] = sha(RUN / 'packt_target_graph.json')
        result = run_worker((ROOT / index[identity]['input']).resolve(), cells, timeout=90)
        path = RUN / 'references' / (identity + '_uno.json.gz'); save(path, result)
        files[str(path.relative_to(RUN))] = sha(path)
        print('DEV31 reference', identity, result['status'], 'requested', len(cells), flush=True)
        assert result['status'] == 'complete', result
        assert result['read_only'] and result['forced_calculate_all'] and not result['workbook_saved']
    verify_inputs()
    save(RUN / 'reference_lock.json', dict(files=files, execution_lock_sha256=sha(RUN / 'execution_lock.json')))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('phase', choices=['prepare', 'execute', 'reference'])
    globals()[parser.parse_args().phase]()
