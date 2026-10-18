"""Independent input/label audit, full replay, and calculateAll impact witness."""
from pathlib import Path
from collections import Counter
import hashlib
import json
import re
import sys

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'qa'),
               str(ROOT / 'baseline/source/HybridGuard_V5_handoff_20260904/src')]
from run_selector_dev5 import read, save, sha
from run_record_candidates_dev29 import RUN, verify
from record_candidates_dev29 import generate
from typed_execution_adapter import load_typed_inputs
from run_public_execution_audit import run_worker
from run_string_execution_dev28 import value_record
from hard_reference_string_execution_dev28 import reference_value, verify_execution


def token_changes(observed, candidate):
    before, after = Tokenizer(observed).items, Tokenizer(candidate).items
    assert len(before) == len(after)
    changed = []
    for a, b in zip(before, after):
        assert (a.type, a.subtype) == (b.type, b.subtype)
        if a.value == b.value:
            continue
        assert (a.type, a.subtype) == ('OPERAND', 'RANGE')
        aa = re.fullmatch(r"(.*?)(\$?[A-Za-z]+)(\$?)([0-9]+)", a.value)
        bb = re.fullmatch(r"(.*?)(\$?[A-Za-z]+)(\$?)([0-9]+)", b.value)
        assert aa and bb and aa.groups()[:3] == bb.groups()[:3] and not aa[3]
        assert abs(int(aa[4]) - int(bb[4])) == 1
        changed.append((a.value, b.value))
    assert len(changed) == 1


def main():
    target = ROOT / 'qa/RECORD_CANDIDATES_DEV29_VERIFICATION_20261018.json'
    assert not target.exists()
    verify(); verify_execution()
    lock = read(RUN / 'prediction_lock.json')
    assert lock['predictions_sha256'] == sha(RUN / 'predictions.json.gz')
    assert lock['preparation_lock_sha256'] == sha(RUN / 'preparation_lock.json')
    prediction = read(RUN / 'predictions.json.gz')
    assert prediction['labels_read'] == prediction['correction_versions_read'] == []
    assert not prediction['repair_suggestions_generated'] and prediction['status'] == 'candidates_only'
    index = {r['id']: r for r in read(RUN / 'index.json')}
    assert set(prediction['input_accesses']) == {str(ROOT / r['path']) for r in index.values()}
    truth = {r['id']: r for r in read(RUN / 'controls_truth.json.gz')['cases']}
    assert sha(RUN / 'controls_truth.json.gz') == read(RUN / 'preparation_lock.json')['controls_truth_sha256']
    analysis = read(RUN / 'candidate_analysis.json')
    assert analysis['prediction_lock_sha256'] == sha(RUN / 'prediction_lock.json')
    natural_path = ROOT / 'development/natural_cv500_dev27_v1/natural_truth.json.gz'
    assert analysis['natural_truth_sha256'] == sha(natural_path)
    natural = {r['cell']: r['accepted_formula'] for r in read(natural_path)['cases']}
    analyses = {r['id']: r for r in analysis['cases']}
    replayed, candidates, refs_checked = 0, 0, 0
    semantic_inputs, rebuilt_counts = {}, {}
    reference_book = openpyxl.Workbook(); reference_book.remove(reference_book.active)
    queries, comparison_cells = [], []
    for case in prediction['cases']:
        identity = case['id']; path = ROOT / index[identity]['path']
        book = openpyxl.load_workbook(path, data_only=False); typed = load_typed_inputs(path)
        formulas = {(s.title, c.coordinate): c.value for s in book for row in s for c in row if c.data_type == 'f'}
        assert len(formulas) == case['formulas'] == len(case['cells'])
        observed = {s + '!' + a: f for (s, a), f in formulas.items()}
        if identity in truth:
            semantic_inputs[identity] = [(s.title, c.coordinate, c.data_type, c.value)
                                         for s in book for row in s for c in row if c.value is not None]
            row_truth = truth[identity]
            for cell, injected in row_truth['errors'].items():
                assert observed[cell] == injected['injected_formula']
                assert row_truth['all_expected'][cell] == injected['original_formula']
                token_changes(injected['injected_formula'], injected['original_formula'])
            if not row_truth['errors']:
                assert observed == row_truth['all_expected']
            sheet = reference_book.create_sheet(identity)
            for row in book['Records']:
                for cell in row:
                    if cell.value is not None:
                        sheet[cell.coordinate] = cell.value
            if 'Other data' not in reference_book.sheetnames:
                other = reference_book.create_sheet('Other data')
                for row in book['Other data']:
                    for cell in row:
                        if cell.value is not None:
                            other[cell.coordinate] = cell.value
        counts = Counter()
        correct = harmful = reached = proposals = 0
        expected = natural if identity == 'cv500' else truth.get(identity, {}).get('all_expected')
        errors = set(natural) if identity == 'cv500' else set(truth.get(identity, {}).get('errors', {}))
        for row in case['cells']:
            sheet_name, address = row['cell'].rsplit('!', 1)
            f = formulas[(sheet_name, address)]
            assert row['formula_sha256'] == hashlib.sha256(f.encode()).hexdigest()
            result = generate(f, sheet_name, address, typed, formulas, book.sheetnames)
            # JSON object keys stringify integer row-number keys. Replay must
            # compare the exact persisted representation, not Python key types.
            replayed_record = dict(cell=row['cell'], formula_sha256=row['formula_sha256'], **result)
            assert row == json.loads(json.dumps(replayed_record))
            replayed += 1; counts[row['status']] += 1
            if identity in truth:
                queries.append(identity + '!' + address)
            if row['status'] != 'candidate_for_review':
                continue
            candidates += 1; proposals += 1
            assert not row['business_authority'] and not row['automatic_repair_allowed']
            token_changes(f, row['candidate_formula']); refs_checked += 1
            if expected is not None:
                correct += row['candidate_formula'] == expected[row['cell']]
                harmful += row['cell'] not in errors and row['candidate_formula'] != expected[row['cell']]
                reached += row['cell'] in errors and row['candidate_formula'] == expected[row['cell']]
            if identity in truth:
                col = 'U' if address.startswith('T') else 'F'
                candidate_address = col + re.search(r'[0-9]+$', address)[0]
                reference_book[identity][candidate_address] = row['candidate_formula']
                queries.append(identity + '!' + candidate_address)
                comparison_cells.append(dict(id=identity, original=identity + '!' + address,
                                             candidate=identity + '!' + candidate_address,
                                             error=row['cell'] in errors))
        assert dict(counts) == case['counts']
        if expected is not None:
            recomputed = dict(candidates=proposals, correct_candidates=correct, errors=len(errors),
                              harmful_candidates_on_clean_cells=harmful, correct_error_cells_reachable=reached,
                              candidate_precision=correct / proposals if proposals else None,
                              candidate_error_reachability=reached / len(errors) if errors else None)
            for key, value in recomputed.items():
                assert analyses[identity][key] == value, (identity, key)
            rebuilt_counts[identity] = recomputed
    assert replayed == 1195 and candidates == refs_checked == 649
    pairs = [('forward', 'legal_lead'), ('backward', 'legal_lag'),
             ('nested', 'legal_nested'), ('ampersand', 'legal_ampersand')]
    for positive, clean in pairs:
        assert semantic_inputs[positive] == semantic_inputs[clean]
        assert len(truth[positive]['errors']) == 12 and not truth[clean]['errors']
        assert all(truth[positive]['all_expected'][cell] != formula
                   for cell, formula in truth[clean]['all_expected'].items())
    # A single isolated calculation corroborates changed outputs, separately
    # from the authored intent labels. The original 19 inputs remain untouched.
    reference_path = RUN / 'control_output_witness.xlsx'; assert not reference_path.exists()
    reference_book.save(reference_path); reference_sha = sha(reference_path)
    reference = run_worker(reference_path, queries, timeout=60)
    assert reference['status'] == 'complete', reference
    assert reference['forced_calculate_all'] and reference['read_only'] and not reference['workbook_saved']
    assert reference_sha == sha(reference_path)
    save(RUN / 'control_output_witness_uno.json.gz', reference)
    impacts = []
    for pair in comparison_cells:
        a, b = [reference_value(reference['cells'][pair[k]]) for k in ('original', 'candidate')]
        assert isinstance(a, str) and isinstance(b, str) and a != b
        impacts.append(dict(**pair, output_changed=True, original_value=value_record(a), candidate_value=value_record(b)))
    assert len(impacts) == 144 and sum(not r['error'] for r in impacts) == 48
    save(RUN / 'control_output_impact.json', dict(rows=impacts, candidate_output_changes=144,
         changes_on_clean_cells=48, changed_outputs_do_not_supply_business_intent=True,
         reference_sha256=sha(RUN / 'control_output_witness_uno.json.gz'),
         input_sha256=reference_sha, auditor_sha256=sha(Path(__file__))))
    verify(); verify_execution()
    from run_selector_dev11 import verify as verify11
    from run_nested_calibration_dev25 import verify as verify25
    verify11(); verify25()
    full = (ROOT / 'qa/RECORD_CANDIDATES_DEV29_FULL_TESTS_20261018.log').read_text()
    assert 'Ran 278 tests' in full and full.rstrip().endswith('OK')
    save(target, dict(all_checks_passed=True, all_formula_declarations_replayed=1195,
         candidate_reference_edits_independently_checked=649, full_tests_passed=278,
         natural_revealed_development_errors=505, natural_errors_with_correct_candidate=505,
         natural_model_suggestions=0, natural_model_precision=None, natural_model_coverage=0,
         authored_controls=16, authored_error_cells=96, authored_correct_candidates=96,
         authored_clean_cells=96, harmful_candidates_on_authored_clean_cells=48,
         authored_candidate_precision=96 / 144, authored_candidate_reachability=1,
         semantically_identical_input_pairs_with_conflicting_intents=4,
         identical_pair_conflicting_formula_targets=48, independent_calculate_all_cells=len(queries),
         independent_output_change_witnesses=144, independent_clean_output_changes=48,
         old_dev11_dev25_dev28_locks_unchanged=True, autonomous_rule_rejected=True,
         main_candidate_changed=False, performance_improvement_established=False,
         accepted_final_sources=0, accepted_final_families=0, final_validation_complete=False, goal_status='active',
         prediction_lock_sha256=sha(RUN / 'prediction_lock.json'),
         analysis_sha256=sha(RUN / 'candidate_analysis.json'), verifier_sha256=sha(Path(__file__))))
    print('DEV29 VERIFIED: 1195 full declaration replay, 505/505 natural candidate reachability;',
          '96/96 authored error candidates plus 48 harmful clean candidates;',
          '4 identical-input opposite-intent pairs; 278 tests; no model adoption', flush=True)


if __name__ == '__main__':
    main()
