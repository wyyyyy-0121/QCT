"""Prepare controls, freeze candidates without labels, score in a separate phase."""
from pathlib import Path
from collections import Counter
import argparse
import hashlib
import os
import sys
import time

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'qa'),
               str(ROOT / 'baseline/source/HybridGuard_V5_handoff_20260904/src')]
from run_selector_dev5 import read, save, sha
from run_string_execution_dev28 import CASES
from typed_execution_adapter import load_typed_inputs
from record_candidates_dev29 import generate, VERSION

RUN = ROOT / 'development/record_candidates_dev29_v1'
ERROR_TYPES = ['forward', 'backward', 'nested', 'ampersand', 'escaped', 'absolute_column', 'repeated', 'other_columns']
CLEAN_TYPES = ['legal_lead', 'legal_lag', 'legal_nested', 'legal_ampersand', 'absolute_row', 'cross_sheet', 'two_anchors', 'two_dissent']


def formula(kind, row, corrected=False):
    anchor = row
    suspect = row if corrected else row + (-1 if kind in ('backward', 'legal_lag') else 1)
    a, b, c, d = [col + str(anchor) for col in ('A', 'B', 'C', 'D')]
    c = 'C' + str(suspect)
    if kind == 'other_columns':
        a, b, c, d = 'H' + str(row), 'J' + str(row), 'L' + str(suspect), 'N' + str(row)
    if kind == 'absolute_column':
        c = '$' + c
    if kind == 'absolute_row':
        c = 'C$' + str(suspect)
    if kind == 'cross_sheet':
        c = "'Other data'!" + c
    if kind == 'two_anchors':
        b = a
    if kind == 'two_dissent':
        d = 'D' + str(row + 1)
    if kind in ('nested', 'legal_nested'):
        return f'=CONCATENATE({a},CONCATENATE({b},{c}),{d})'
    if kind in ('ampersand', 'legal_ampersand'):
        return '=' + '&'.join([a, b, c, d])
    if kind == 'escaped':
        return f'=CONCATENATE("C{row + 1} says ""hello""",{a},{b},{c},{d})'
    if kind == 'repeated':
        return f'=CONCATENATE({a},{a},{b},{c},{d})'
    return f'=CONCATENATE({a},{b},{c},{d})'


def prepare():
    assert not RUN.exists()
    RUN.mkdir()
    files = [Path(__file__), ROOT / 'scripts/record_candidates_dev29.py',
             ROOT / 'tests/test_record_candidates_dev29.py', ROOT / 'scripts/string_execution_dev28_v2.py',
             ROOT / 'scripts/string_execution_dev28.py', ROOT / 'scripts/typed_execution_adapter.py',
             ROOT / 'RECORD_CANDIDATES_DEV29_PREREGISTRATION_20261018.md']
    save(RUN / 'source_lock.json', dict(version=VERSION, files={str(p.relative_to(ROOT)): sha(p) for p in files}))
    index, truth = [], []
    for kind in ERROR_TYPES + CLEAN_TYPES:
        book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Records'
        other = book.create_sheet('Other data')
        for row in range(1, 19):
            for col in ['A', 'B', 'C', 'D', 'H', 'J', 'L', 'N']:
                sheet[f'{col}{row}'] = f'{col}-record-{row}'
                other[f'{col}{row}'] = f'external-{col}-{row}'
        injected = kind in ERROR_TYPES
        output = 'T' if kind == 'other_columns' else 'E'
        errors, observed, expected = {}, {}, {}
        for row in range(5, 17):
            address = output + str(row)
            sheet[address] = formula(kind, row)
            observed[address] = sheet[address].value
            expected[address] = formula(kind, row, corrected=True) if injected else sheet[address].value
            if injected:
                errors['Records!' + address] = dict(original_formula=expected[address], injected_formula=observed[address])
        path = RUN / 'controls' / (kind + '.xlsx'); path.parent.mkdir(exist_ok=True); book.save(path)
        index.append(dict(id=kind, role='self-authored development control; no final eligibility',
                          path=str(path.relative_to(ROOT)), sha256=sha(path)))
        truth.append(dict(id=kind, errors=errors, all_expected={'Records!' + k: v for k, v in expected.items()},
                          intent='All fields belong to current record' if injected else
                          'Legitimate cross-row/absolute/external reference or insufficient structural evidence; preserve original formula',
                          truth_source='development author defined before mutation/prediction; not natural error evidence'))
    save(RUN / 'controls_truth.json.gz', dict(cases=truth, final_validation_eligible=False))
    index += [dict(id=name, path=str(path.relative_to(ROOT)), sha256=sha(path), role='revealed public development original') for name, path in CASES]
    save(RUN / 'index.json', index)
    save(RUN / 'preparation_lock.json', dict(index_sha256=sha(RUN / 'index.json'),
         source_lock_sha256=sha(RUN / 'source_lock.json'), controls_truth_sha256=sha(RUN / 'controls_truth.json.gz'),
         inputs={r['path']: r['sha256'] for r in index}))
    print('DEV29 prepared 16 full controls, 192 formulas, plus 1003 public original formulas', flush=True)


def verify():
    lock = read(RUN / 'preparation_lock.json')
    assert sha(RUN / 'index.json') == lock['index_sha256']
    assert sha(RUN / 'source_lock.json') == lock['source_lock_sha256']
    for rel, digest in {**read(RUN / 'source_lock.json')['files'], **lock['inputs']}.items():
        assert sha(ROOT / rel) == digest, rel


def predict():
    verify()
    assert not (RUN / 'prediction_lock.json').exists()
    index = read(RUN / 'index.json')
    active, accessed = [True], set()
    def guard(event, args):
        if not active[0] or event != 'open' or not args or not isinstance(args[0], (str, bytes, os.PathLike)):
            return
        name = os.fsdecode(args[0])
        if any(s in name.lower() for s in ['truth', 'natural_score', 'author_diff', 'after.xlsx', 'certification']):
            raise PermissionError('Candidate inference cannot read scoring evidence: ' + name)
        if name.endswith('.xlsx'):
            accessed.add(name)
    sys.addaudithook(guard)
    rows = []; wall, cpu = time.perf_counter(), time.process_time()
    for item in index:
        path = ROOT / item['path']; book = openpyxl.load_workbook(path, data_only=False); typed = load_typed_inputs(path)
        formulas = {(s.title, c.coordinate): c.value for s in book for row in s for c in row if c.data_type == 'f'}
        cells = []
        for (sheet, address), f in sorted(formulas.items()):
            result = generate(f, sheet, address, typed, formulas, book.sheetnames)
            cells.append(dict(cell=sheet + '!' + address, formula_sha256=hashlib.sha256(f.encode()).hexdigest(), **result))
        counts = dict(Counter(r['status'] for r in cells))
        rows.append(dict(id=item['id'], formulas=len(formulas), counts=counts, cells=cells))
        print('DEV29 candidates', item['id'], len(formulas), counts, flush=True)
    save(RUN / 'predictions.json.gz', dict(cases=rows, labels_read=[], correction_versions_read=[],
         input_accesses=sorted(accessed), cpu_seconds=time.process_time() - cpu, wall_seconds=time.perf_counter() - wall,
         status='candidates_only', repair_suggestions_generated=False, main_candidate_changed=False, final_validation_eligible=False))
    save(RUN / 'prediction_lock.json', dict(predictions_sha256=sha(RUN / 'predictions.json.gz'),
         preparation_lock_sha256=sha(RUN / 'preparation_lock.json'), source_lock_sha256=sha(RUN / 'source_lock.json')))
    active[0] = False
    verify()


def score():
    verify()
    target = RUN / 'candidate_analysis.json'; assert not target.exists()
    lock = read(RUN / 'prediction_lock.json')
    assert sha(RUN / 'predictions.json.gz') == lock['predictions_sha256']
    assert sha(RUN / 'controls_truth.json.gz') == read(RUN / 'preparation_lock.json')['controls_truth_sha256']
    controls = {r['id']: r for r in read(RUN / 'controls_truth.json.gz')['cases']}
    natural_path = ROOT / 'development/natural_cv500_dev27_v1/natural_truth.json.gz'
    natural = {r['cell']: r['accepted_formula'] for r in read(natural_path)['cases']}
    rows = []
    for case in read(RUN / 'predictions.json.gz')['cases']:
        candidates = [r for r in case['cells'] if r['status'] == 'candidate_for_review']
        expected = natural if case['id'] == 'cv500' else controls.get(case['id'], {}).get('all_expected')
        errors = set(natural) if case['id'] == 'cv500' else set(controls.get(case['id'], {}).get('errors', {}))
        if expected is None:
            rows.append(dict(id=case['id'], candidates=len(candidates), reliable_labels=False,
                             candidate_precision=None, candidate_error_reachability=None))
            continue
        correct = sum(r['candidate_formula'] == expected[r['cell']] for r in candidates)
        harmful = sum(r['cell'] not in errors and r['candidate_formula'] != expected[r['cell']] for r in candidates)
        reached = sum(r['cell'] in errors and r['candidate_formula'] == expected[r['cell']] for r in candidates)
        rows.append(dict(id=case['id'], candidates=len(candidates), correct_candidates=correct, errors=len(errors),
                         harmful_candidates_on_clean_cells=harmful, correct_error_cells_reachable=reached,
                         candidate_precision=correct / len(candidates) if candidates else None,
                         candidate_error_reachability=reached / len(errors) if errors else None,
                         model_repair_suggestions=0, model_precision=None, model_coverage=0 if errors else None))
    save(target, dict(cases=rows, prediction_lock_sha256=sha(RUN / 'prediction_lock.json'),
         control_truth_sha256=sha(RUN / 'controls_truth.json.gz'), natural_truth_sha256=sha(natural_path),
         natural_revealed_development_only=True, all_candidate_results_exploratory=True,
         main_candidate_adopted=False, autonomous_repair_rule_rejected=True,
         reason='Legitimate cross-row formulas generate harmful candidates; row coherence alone is not business authority',
         accepted_final_sources=0, accepted_final_families=0, final_validation_complete=False))
    for row in rows:
        print('DEV29 analysis', row, flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('phase', choices=['prepare', 'predict', 'score'])
    globals()[parser.parse_args().phase]()
