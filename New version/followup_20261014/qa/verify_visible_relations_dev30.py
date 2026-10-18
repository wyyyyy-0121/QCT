"""Replay all cases, rebuild score counts, and obtain independent execution witnesses."""
from pathlib import Path
from collections import Counter
import os
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'qa'),
               str(ROOT / 'baseline/source/HybridGuard_V5_handoff_20260904/src')]
from run_selector_dev5 import read, save, sha
from run_visible_relations_dev30 import RUN, OLD, verify, prediction
from relation_algebra_dev30 import equivalent as fraction_equivalent
from run_public_execution_audit import run_worker
from run_string_execution_dev28 import value_record
from hard_reference_string_execution_dev28 import reference_value


def ratio(num, den):
    return num / den if den else None


def main():
    target = ROOT / 'qa/VISIBLE_RELATIONS_DEV30_VERIFICATION_20261018.json'
    assert not target.exists()
    verify()
    lock = read(RUN / 'prediction_lock.json')
    assert lock['source_lock_sha256'] == sha(RUN / 'source_lock.json')
    assert lock['preparation_lock_sha256'] == sha(RUN / 'preparation_lock.json')
    assert lock['labels_read'] == []
    for rel, digest in lock['files'].items():
        assert sha(RUN / rel) == digest
    assert sha(RUN / 'control_truth.json') == read(RUN / 'preparation_lock.json')['control_truth_sha256']
    score = read(RUN / 'score.json')
    assert score['prediction_lock_sha256'] == sha(RUN / 'prediction_lock.json')
    assert score['old_truth_sha256'] == sha(OLD / 'heldout_truth.json.gz')
    auditor = RUN / 'audit_source_lock.json'; assert not auditor.exists()
    save(auditor, dict(files={str(Path(__file__).relative_to(ROOT)): sha(Path(__file__))},
         score_sha256=sha(RUN / 'score.json'), prediction_lock_sha256=sha(RUN / 'prediction_lock.json')))
    index = read(RUN / 'public_index.json'); assert len(index) == 251
    active, accesses = [True], set()
    def guard(event, args):
        if not active[0] or event != 'open' or not args or not isinstance(args[0], (str, bytes, os.PathLike)):
            return
        path = os.fsdecode(args[0])
        if any(x in path.lower() for x in ['truth', 'training_labels', 'case_labels', 'author_diff',
                                          'product_sum_oracle', 'certification', 'after.xlsx']):
            raise PermissionError('Replay cannot read truth/oracle: ' + path)
        if path.endswith(('.xlsx', '.xlsm')):
            accesses.add(path)
    sys.addaudithook(guard)
    predictions = {}
    for item in index:
        path = RUN / 'predictions' / (item['case'].replace('/', '__') + '.json.gz')
        prior = read(path); current = prediction(item)
        assert prior == current, item['case']
        predictions[item['case']] = prior
        print('DEV30 replayed', item['case'], prior['status'], len(prior['relation_only']), flush=True)
    active[0] = False
    assert accesses == set(lock['input_accesses'])
    rejected = [i['case'] for i in index if predictions[i['case']]['status'] != 'complete']
    assert len(rejected) == 3 and set(rejected) == {'public/public05', 'public/public06', 'public/public07'}
    old = read(OLD / 'heldout_truth.json.gz')['cases']
    controls = {r['case']: r for r in read(RUN / 'control_truth.json')['cases']}
    string = {r['id']: r for r in read(ROOT / 'development/record_candidates_dev29_v1/controls_truth.json.gz')['cases']}
    natural = {r['cell']: r['accepted_formula'] for r in read(ROOT / 'development/natural_cv500_dev27_v1/natural_truth.json.gz')['cases']}
    stored_cases = {r['case']: r for r in score['cases']}
    checked_metrics = 0
    for item in index:
        case = item['case']; stratum = item['stratum']; pred = predictions[case]
        expected = {}; errors = set(); use_fraction = False
        if stratum == 'old213':
            expected = old[case]['expected_formulas']; errors = set(old[case]['errors'])
            use_fraction = item['group'] == 'public_relations'
        elif stratum == 'new_control':
            expected = controls[case]['expected_formulas']; errors = set(controls[case]['errors']); use_fraction = True
            workbook = openpyxl.load_workbook(ROOT / item['input'])
            actual = {s.title + '!' + c.coordinate: c.value for s in workbook for row in s for c in row if c.data_type == 'f'}
            assert actual == controls[case]['observed_formulas']
            assert errors == {k for k, v in expected.items() if actual[k] != v}
            assert len(controls[case]['injection_records']) == len(errors)
        elif stratum == 'string_control':
            expected = string[item['id']]['all_expected']; errors = set(string[item['id']]['errors'])
        elif stratum == 'natural_original' and item['id'] == 'cv500':
            expected = natural; errors = set(expected)
        model = None
        if use_fraction and any(pred[m] for m in ['dev11', 'relation_only', 'combined']):
            from formulaguard.workbook import WorkbookModel
            model = WorkbookModel.from_xlsx(ROOT / item['input'])
        for mode in ['dev11', 'relation_only', 'combined']:
            right, wrong, unknown, unnecessary = [], [], [], []
            for row in pred[mode]:
                cell = row['cell']
                if cell not in expected:
                    unknown.append(cell)
                elif cell not in errors:
                    wrong.append(cell); unnecessary.append(cell)
                else:
                    matches = (fraction_equivalent(row['suggested_formula'], expected[cell], cell.rsplit('!', 1)[0], model)
                               if use_fraction else row['suggested_formula'] == expected[cell])
                    # No proposals are produced in LCOE, so its distinct
                    # old algebraic-equivalence scoring need not be guessed.
                    assert not (stratum == 'old213' and item['group'] == 'lcoe'), 'Unexpected LCOE proposals require its oracle'
                    (right if matches else wrong).append(cell)
            found = dict(suggestions=len(pred[mode]), correct=len(right), wrong=len(wrong), unknown=len(unknown),
                         errors=len(errors), correct_errors=len(set(right)), unnecessary=len(unnecessary),
                         precision_confirmed_lower_bound=ratio(len(right), len(pred[mode])), coverage=ratio(len(set(right)), len(errors)))
            stored = stored_cases[case][mode]
            for field, value in found.items():
                assert stored[field] == value, (case, mode, field)
                checked_metrics += 1
            assert stored['correct_cells'] == right and stored['wrong_cells'] == wrong and stored['unknown_cells'] == unknown
        base = {r['cell']: r for r in pred['dev11']}
        combined = {r['cell']: r for r in pred['combined']}
        assert all(combined[k] == v for k, v in base.items())
    # Rebuild every aggregation from independently verified per-case integers.
    for key, rows in [('by_stratum', score['by_stratum']), ('old213_by_origin', score['old213_by_origin']),
                      ('old213_by_family', score['old213_by_family']), ('old213_by_scenario', score['old213_by_scenario'])]:
        field = {'by_stratum': 'stratum', 'old213_by_origin': 'origin', 'old213_by_family': 'family', 'old213_by_scenario': 'scenario'}[key]
        for group, modes in rows.items():
            members = [r for r in score['cases'] if r[field] == group and (key == 'by_stratum' or r['stratum'] == 'old213')]
            for mode, metric in modes.items():
                for name in ['suggestions', 'correct', 'wrong', 'unknown', 'errors', 'correct_errors', 'unnecessary']:
                    assert metric[name] == sum(m[mode][name] for m in members)
                    checked_metrics += 1
                assert metric['precision_confirmed_lower_bound'] == ratio(metric['correct'], metric['suggestions'])
                assert metric['coverage'] == ratio(metric['correct_errors'], metric['errors'])
    old_metrics = score['by_stratum']['old213']; assert old_metrics['dev11']['correct'] == 14184
    assert old_metrics['dev11']['errors'] == old_metrics['combined']['errors'] == 18073
    assert old_metrics['relation_only']['suggestions'] == old_metrics['relation_only']['correct'] == 28
    assert old_metrics['combined']['correct'] == 14212 and old_metrics['combined']['wrong'] == 29
    negative = score['by_stratum']['new_control']['relation_only']
    assert negative['correct'] == 24 and negative['wrong'] == 8 and negative['errors'] == 40
    assert score['by_stratum']['natural_original']['combined']['correct_errors'] == 0
    # Consolidated authored controls: compare original and copied candidate
    # outputs with a completely separate spreadsheet engine, preserving all
    # original input files. This validates execution, not business intent.
    book = openpyxl.Workbook(); book.remove(book.active)
    queries, pairs = [], []
    for item in [i for i in index if i['stratum'] == 'new_control']:
        identity = item['id']; source = openpyxl.load_workbook(ROOT / item['input'])['Ledger']
        sheet = book.create_sheet(identity)
        for row in source:
            for c in row:
                if c.value is not None:
                    sheet[c.coordinate] = c.value
        for n, r in enumerate(predictions[item['case']]['relation_only'], 10):
            address = r['cell'].split('!')[1]
            col = openpyxl.utils.cell.coordinate_from_string(address)[0]
            target_address = col + str(n)
            sheet[target_address] = r['suggested_formula']
            queries += [identity + '!' + address, identity + '!' + target_address]
            pairs.append(dict(case=item['case'], original=identity + '!' + address, candidate=identity + '!' + target_address,
                              error=r['cell'] in controls[item['case']]['errors']))
    reference_path = RUN / 'control_execution_witness.xlsx'; assert not reference_path.exists(); book.save(reference_path)
    before = sha(reference_path); ref = run_worker(reference_path, queries, timeout=60)
    assert ref['status'] == 'complete' and ref['forced_calculate_all'] and ref['read_only'] and not ref['workbook_saved']
    assert sha(reference_path) == before
    changes = []
    for pair in pairs:
        a, b = [reference_value(ref['cells'][pair[k]]) for k in ['original', 'candidate']]
        assert type(a) in (int, float) and type(b) in (int, float) and a != b
        changes.append(dict(**pair, original_value=value_record(a), candidate_value=value_record(b), output_changed=True))
    assert len(changes) == 32 and sum(not r['error'] for r in changes) == 8
    save(RUN / 'control_execution_reference.json.gz', ref)
    save(RUN / 'control_execution_impact.json', dict(rows=changes, source_unchanged=True, input_sha256=before,
         reference_sha256=sha(RUN / 'control_execution_reference.json.gz'), business_intent_not_supplied_by_reference=True))
    # Cross-check raw historical sources/predictions, not just current metrics.
    from run_selector_dev11 import verify as verify11
    from run_nested_calibration_dev25 import verify as verify25
    from hard_reference_string_execution_dev28 import verify_execution
    verify11(); verify25(); verify_execution(); verify()
    full = (ROOT / 'qa/VISIBLE_RELATIONS_DEV30_FULL_TESTS_20261018.log').read_text()
    assert 'Ran 286 tests' in full and full.rstrip().endswith('OK')
    save(target, dict(all_checks_passed=True, cases_replayed=251, retained_original_rejections=rejected,
         independent_metric_checks=checked_metrics, exact_fraction_score_crosschecks=True,
         full_tests_passed=286, old213_new_correct=28, old213_new_wrong=0,
         old213_combined_precision=old_metrics['combined']['precision_confirmed_lower_bound'],
         old213_combined_coverage=old_metrics['combined']['coverage'],
         new_control_correct=24, new_control_wrong=8, new_control_errors=40,
         new_control_precision=.75, new_control_coverage=.6,
         independent_control_output_changes=32, independent_clean_output_changes=8,
         natural_development_errors=505, natural_development_correct_errors=0,
         old_dev11_dev25_dev28_locks_unchanged=True, main_candidate_adopted=False,
         generalization_improvement_established=False, accepted_final_sources=0, accepted_final_families=0,
         final_validation_complete=False, goal_status='active', verifier_sha256=sha(Path(__file__)),
         prediction_lock_sha256=sha(RUN / 'prediction_lock.json'), score_sha256=sha(RUN / 'score.json')))
    print('DEV30 VERIFIED: 251-case replay, independent counts and 286 tests;',
          'old213 +28 correct/+0 wrong; new controls 24 correct/8 wrong; natural zero coverage; no adoption', flush=True)


if __name__ == '__main__':
    main()
