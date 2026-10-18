from fractions import Fraction
from pathlib import Path
import sys
import unittest
from collections import defaultdict
from openpyxl import Workbook
import sympy
ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT/'qa'), str(ROOT/'scripts')]
from independent_numeric_relations_dev23 import IndependentNumeric
from tristate_relations_dev23 import NumericRelations, evaluate


class IndependentNumericTests(unittest.TestCase):
    def test_two_parsers_agree_with_direct_fraction_arithmetic(self):
        w = Workbook(); s = w.active; s.title = 'Main'
        s['A1'] = 4; s['B1'] = 7; s['C1'] = 8; s['D1'] = '=A1-B1/2'
        s['E1'] = 'text'
        independent = IndependentNumeric(w, {})
        production = NumericRelations(w)
        for a, b, c in [(Fraction(3), Fraction(7), Fraction(11)),
                        (Fraction(-4, 3), Fraction(2, 7), Fraction(0)),
                        (Fraction(0), Fraction(-2), Fraction(5, 9))]:
            inputs = {'Main!A1': a, 'Main!B1': b, 'Main!C1': c}
            cases = [('=A1+B1*C1', a+b*c), ('=(A1+B1)*C1', (a+b)*c),
                     ('=D1', a-b/2), ('=SUM(A1:C1,D1)', 2*a+b/2+c),
                     ('=SUM(A1:F1)', 2*a+b/2+c), ('=-A1/-2+B1', a/2+b),
                     ('=A1-(B1-C1)', a-b+c), ('=0.01*A1', a/100), ('=SUM()', Fraction(0))]
            for formula, expected in cases:
                self.assertEqual(independent.formula(formula, 'Main', inputs=inputs), expected)
                self.assertEqual(evaluate(production.compile(formula, 'Main'), inputs), expected)
        w.close()

    def test_independent_symbolic_expansion_and_corrected_dependency(self):
        w = Workbook(); s = w.active; s.title = 'Main'
        s['A1'] = 4; s['B1'] = 7; s['D1'] = '=99*A1'
        interpreter = IndependentNumeric(w, {'Main!D1': '=A1*2'})
        a = interpreter.formula('=SUM(D1,B1)', 'Main', symbolic=True)
        b = interpreter.formula('=2*A1+B1', 'Main', symbolic=True)
        self.assertEqual(sympy.expand(a-b), 0)
        w.close()

    def test_cancelled_inputs_can_be_completed_without_changing_witness(self):
        w = Workbook(); s = w.active; s.title = 'Main'
        s['A1'] = 3; s['B1'] = 7
        production = NumericRelations(w); independent = IndependentNumeric(w, {})
        result = production.certify('=A1+B1-B1', '=2*A1', 'Main')
        self.assertEqual(result['label'], 0)
        self.assertNotIn('Main!B1', result['witness']['inputs'])
        inputs = defaultdict(lambda: Fraction(1), {k: Fraction(v) for k, v in result['witness']['inputs'].items()})
        self.assertEqual(independent.formula('=A1+B1-B1', 'Main', inputs=inputs), Fraction(result['witness']['candidate']))
        w.close()


if __name__ == '__main__': unittest.main()
