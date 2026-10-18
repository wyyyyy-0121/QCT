import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from static_ooxml_view_dev1 import create_view, CT, XLSM


class StaticViewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)

    def package(self, name='Data', macro=True):
        wb = Workbook(); ws = wb.active; ws.title = name
        ws['A1'] = 3; ws['A2'] = '=A1*2'; ws['B2'] = '=A1+1'
        ws.row_dimensions[2].hidden = True
        ws.column_dimensions['B'].hidden = True
        hidden = wb.create_sheet('Hidden'); hidden.sheet_state = 'hidden'
        hidden['A1'] = '=Data!A1+5'
        buf = io.BytesIO(); wb.save(buf); wb.close()
        with zipfile.ZipFile(buf) as z:
            payload = {n: z.read(n) for n in z.namelist()}
        if macro:
            types = ET.fromstring(payload['[Content_Types].xml'])
            for x in types:
                if x.get('PartName') == '/xl/workbook.xml':
                    x.set('ContentType', XLSM)
            ET.SubElement(types, '{'+CT+'}Override', PartName='/xl/vbaProject.bin',
                          ContentType='application/vnd.ms-office.vbaProject')
            payload['[Content_Types].xml'] = ET.tostring(types)
            rels = ET.fromstring(payload['xl/_rels/workbook.xml.rels'])
            ET.SubElement(rels, '{'+rels.tag.split('}')[0][1:]+'}Relationship',
                          Id='rIdVBA', Type='http://schemas.microsoft.com/office/2006/relationships/vbaProject',
                          Target='vbaProject.bin')
            payload['xl/_rels/workbook.xml.rels'] = ET.tostring(rels)
            payload['xl/vbaProject.bin'] = b'opaque inert macro fixture; not executable'
        return payload

    def write(self, payload, filename='input.xlsm'):
        path = self.base / filename
        with zipfile.ZipFile(path, 'w') as z:
            for n, data in payload.items(): z.writestr(n, data)
        return path

    def test_macro_removed_all_worksheet_bytes_and_visibility_preserved(self):
        payload = self.package(); src = self.write(payload); dest = self.base/'view.xlsx'
        before = src.read_bytes(); receipt = create_view(src, dest)
        self.assertEqual(src.read_bytes(), before)
        self.assertEqual(set(receipt['removed_parts']), {'xl/vbaProject.bin'})
        with zipfile.ZipFile(dest) as z:
            self.assertNotIn('xl/vbaProject.bin', z.namelist())
            for n, data in payload.items():
                if n not in ['[Content_Types].xml', 'xl/_rels/workbook.xml.rels', 'xl/vbaProject.bin']:
                    self.assertEqual(z.read(n), data)
        wb = load_workbook(dest, data_only=False)
        self.assertEqual(wb['Data']['A2'].value, '=A1*2')
        self.assertTrue(wb['Data'].row_dimensions[2].hidden)
        self.assertTrue(wb['Data'].column_dimensions['B'].hidden)
        self.assertEqual(wb['Hidden'].sheet_state, 'hidden'); wb.close()
        self.assertFalse(receipt['recalculation'])
        self.assertFalse(receipt['functional_excel_equivalence_proven'])

    def test_shared_array_formula_and_cache_xml_never_rewritten(self):
        payload = self.package()
        payload['xl/worksheets/sheet1.xml'] = b'''<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1"><v>3</v></c><c r="B1"><f t="shared" si="0" ref="B1:B2">A1*2</f><v>6</v></c><c r="C1"><f t="array" ref="C1:C2">A1:A2*2</f><v>6</v></c></row><row r="2"><c r="B2"><f t="shared" si="0"/><v>8</v></c></row></sheetData></worksheet>'''
        src = self.write(payload); dest = self.base/'view.xlsx'
        create_view(src, dest)
        with zipfile.ZipFile(dest) as z:
            self.assertEqual(z.read('xl/worksheets/sheet1.xml'), payload['xl/worksheets/sheet1.xml'])

    def test_xlsx_passive_parts_unchanged_and_deterministic(self):
        payload = self.package(macro=False); src = self.write(payload, 'input.xlsx')
        a, b = self.base/'a.xlsx', self.base/'b.xlsx'
        ra, rb = create_view(src,a), create_view(src,b)
        self.assertEqual(ra['view_sha256'],rb['view_sha256'])
        self.assertFalse(ra['modified_package_metadata'])
        with zipfile.ZipFile(a) as z:
            self.assertEqual({n:z.read(n) for n in z.namelist()},payload)

    def test_protected_manifest_rejected_before_shared_strings_read(self):
        payload = self.package(name='Answer keys')
        payload['xl/sharedStrings.xml'] = b'not XML; must never read these answers'
        src = self.write(payload); opened=[]; original=zipfile.ZipFile.read
        def traced(z, name, *args, **kwargs):
            opened.append(name); return original(z,name,*args,**kwargs)
        with patch.object(zipfile.ZipFile,'read',traced):
            with self.assertRaisesRegex(ValueError,'answer/label'):
                create_view(src,self.base/'view.xlsx')
        self.assertEqual(opened,['xl/workbook.xml'])
        self.assertFalse((self.base/'view.xlsx').exists())

    def test_refuses_overwrite(self):
        src=self.write(self.package()); dest=self.base/'view.xlsx';dest.write_bytes(b'keep')
        with self.assertRaises(ValueError):create_view(src,dest)
        self.assertEqual(dest.read_bytes(),b'keep')

    def test_xls_and_xlsb_rejected(self):
        for name in ['old.xls','packed.xlsb']:
            with self.assertRaisesRegex(ValueError,'Only static OOXML'):
                create_view(self.base/name,self.base/'view.xlsx')

    def test_unsafe_member_rejected(self):
        payload=self.package();payload['../outside.xml']=b'x'
        with self.assertRaisesRegex(ValueError,'Unsafe package'):
            create_view(self.write(payload),self.base/'view.xlsx')

    def test_unknown_binary_refused(self):
        payload=self.package();payload['xl/embeddings/object.bin']=b'unknown'
        with self.assertRaisesRegex(ValueError,'Unrecognised binary'):
            create_view(self.write(payload),self.base/'view.xlsx')

    def test_external_workbook_relationship_refused(self):
        payload=self.package(); r=ET.fromstring(payload['xl/_rels/workbook.xml.rels'])
        ET.SubElement(r,r[0].tag,Id='ext',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink',Target='https://example.org/data.xlsx',TargetMode='External')
        payload['xl/_rels/workbook.xml.rels']=ET.tostring(r)
        with self.assertRaisesRegex(ValueError,'External non-hyperlink'):
            create_view(self.write(payload),self.base/'view.xlsx')

    def test_dtd_manifest_refused(self):
        payload=self.package();payload['xl/workbook.xml']=b'<!DOCTYPE x [<!ENTITY e "x">]><x/>'
        with self.assertRaisesRegex(ValueError,'DTD/entity'):
            create_view(self.write(payload),self.base/'view.xlsx')

    def test_broken_relationship_path_refused(self):
        payload=self.package();r=ET.fromstring(payload['xl/_rels/workbook.xml.rels'])
        r[0].set('Target','../../../outside.xml')
        payload['xl/_rels/workbook.xml.rels']=ET.tostring(r)
        with self.assertRaisesRegex(ValueError,'escapes package'):
            create_view(self.write(payload),self.base/'view.xlsx')


if __name__ == '__main__': unittest.main()
