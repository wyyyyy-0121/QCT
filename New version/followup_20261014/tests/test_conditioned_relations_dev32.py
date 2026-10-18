"""Actual repair output, unsupported domains and retained intent ambiguity."""
from copy import deepcopy
from pathlib import Path
import sys
import tempfile
import unittest

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'baseline/source/HybridGuard_V5_handoff_20260904/src')]
from conditioned_relations_dev32 import predict
from dependency_execution_dev31 import load_typed_inputs
from formulaguard.workbook import WorkbookModel


class ConditionedRelationsTests(unittest.TestCase):
    def run_book(self, book):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'input.xlsx'; book.save(path)
            model = WorkbookModel.from_xlsx(path); before = deepcopy((model.cells, model.formulas))
            result = predict(model, load_typed_inputs(path))
            self.assertEqual((model.cells, model.formulas), before)
            return result

    def profit(self, formula='=C8*F8*(1-E8)+B8*C8-D8'):
        book = openpyxl.Workbook(); s = book.active; s.title = 'Orders'
        for a, v in {'B5':'Unit Cost','C5':'Units','D5':'Shipping','E5':'Discount Rate',
                     'F5':'Unit Price','H5':'Net Profit','B8':3,'C8':12,'D8':2,'E8':.1,'F8':20,'H8':formula}.items():
            s[a] = v
        return book

    def debt(self, formula='=F8+F9-F10'):
        book = openpyxl.Workbook(); s = book.active; s.title = 'Ledger'; inputs = book.create_sheet('Rates')
        inputs['A1'] = .04; inputs['B1'] = .05; inputs['D1'] = 1
        for a, v in {'B8':'Opening Balance','B9':'Drawdowns','B10':'Principal Repayment','B11':'Closing Balance',
                     'C8':'kGBP','C9':'kGBP','C10':'kGBP','C11':'kGBP','F8':1000,'F9':100,
                     'F10':'=PPMT(INDEX(Rates!A1:B1,Rates!D1),1,6,1000)','F11':formula}.items():s[a]=v
        return book

    def test_profit_dynamic_binding_and_actual_qualified_output(self):
        p = self.run_book(self.profit())
        self.assertEqual(p['suggestions'][0]['suggested_formula'], '=C8*F8*(1-E8)-C8*B8-D8')
        self.assertFalse(p['suggestions'][0]['evidence']['author_intent_proven'])

    def test_clean_profit_algebra_not_changed(self):
        self.assertEqual(self.run_book(self.profit('=C8*(F8*(1-E8)-B8)-D8'))['suggestions'], [])

    def test_optional_tax_and_omitted_discount(self):
        b=self.profit('=C8*F8-C8*B8-D8');s=b.active;s['E5']='Tax Rate'
        self.assertEqual(self.run_book(b)['suggestions'][0]['suggested_formula'], '=C8*F8*(1-E8)-C8*B8-D8')

    def test_visible_notes_conflicting_instruction_and_duplicate_role_block(self):
        for note in ['Contract: revenue share applies.', 'Net Profit equals Units times Unit Price.']:
            b=self.profit();b.active['K20']=note
            self.assertEqual(self.run_book(b)['suggestions'], [])
        b=self.profit();b.active['K5']='Units';b.active['K8']=12
        self.assertEqual(self.run_book(b)['suggestions'], [])

    def test_unmarked_profit_share_and_adjustment_remain_harmful_hypotheses(self):
        for formula in ['=C8*F8*(1-E8)*0.73-C8*B8-D8', '=C8*F8*(1-E8)-C8*B8-D8+17']:
            self.assertEqual(len(self.run_book(self.profit(formula))['suggestions']), 1)

    def test_profit_domains_and_hidden_input_block(self):
        for value in [-.01,1.01]:
            b=self.profit();b.active['E8']=value
            self.assertEqual(self.run_book(b)['suggestions'], [])
        b=self.profit();b.active.column_dimensions['E'].hidden=True
        self.assertEqual(self.run_book(b)['suggestions'], [])

    def test_profit_unsupported_and_hidden_inactive_dependency_block(self):
        b=self.profit();b.active['C8']='=SQRT(144)'
        self.assertEqual(self.run_book(b)['suggestions'], [])
        b=self.profit();b.active['C8']='=IF(1,12,K8)';b.active['K8']=12;b.active.column_dimensions['K'].hidden=True
        self.assertEqual(self.run_book(b)['suggestions'], [])

    def test_non_profit_caption_does_not_authorize_profit(self):
        b=self.profit();b.active['H5']='Custom Metric'
        self.assertEqual(self.run_book(b)['suggestions'], [])

    def test_profit_section_scope_ends_at_next_nonformula_label(self):
        b=self.profit();b.active['H7']='Memo'
        self.assertEqual(self.run_book(b)['suggestions'], [])

    def test_debt_complete_index_ppmt_chain_generates_signed_sum(self):
        p=self.run_book(self.debt());row=p['suggestions'][0]
        self.assertEqual(row['suggested_formula'], '=SUM(F8,F9,F10)')
        execution=row['evidence']['inputs'][2]['execution']
        self.assertEqual(execution['audit']['functions']['PPMT'],1)
        self.assertEqual(execution['audit']['functions']['INDEX'],1)
        self.assertLess(execution['value']['value'],0)

    def test_clean_debt_equivalence_and_money_units(self):
        self.assertEqual(self.run_book(self.debt('=SUM(F8:F10)'))['suggestions'], [])
        for value in ['USD','unknown']:
            b=self.debt();b.active['C10']=value
            self.assertEqual(self.run_book(b)['suggestions'], [])

    def test_positive_constant_or_hidden_ppmt_chain_abstains(self):
        for formula in ['=-PPMT(.04,1,6,1000)','=-200','=SQRT(100)']:
            b=self.debt();b.active['F10']=formula
            self.assertEqual(self.run_book(b)['suggestions'], [])
        b=self.debt();b['Rates'].column_dimensions['B'].hidden=True
        self.assertEqual(self.run_book(b)['suggestions'], [])

    def test_unmarked_debt_adjustment_remains_a_harmful_hypothesis(self):
        self.assertEqual(len(self.run_book(self.debt('=SUM(F8:F10)+17'))['suggestions']),1)


if __name__=='__main__': unittest.main()
