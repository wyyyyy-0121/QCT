"""Independent token-stack symbolic recount; original XML input and prediction provenance checks."""
from pathlib import Path
from collections import Counter
import re
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from run_annotation_dev14 import RUN,DATA,NAMES,verify
from run_selector_dev5 import read,save,sha
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
import sympy

def symbolic(formula):
    values=[]; ops=[]
    precedence={'+':1,'-':1,'*':2,'/':2,'u+':3,'u-':3}
    def reduce():
        op=ops.pop()
        if op in ('u+','u-'):
            values.append(values.pop()*(1 if op=='u+' else -1)); return
        b,a=values.pop(),values.pop()
        values.append({'+':lambda:a+b,'-':lambda:a-b,'*':lambda:a*b,'/':lambda:a/b}[op]())
    for t in Tokenizer(formula).items:
        if t.type=='WHITE-SPACE':continue
        if t.type=='OPERAND':
            if t.subtype=='RANGE' and re.fullmatch(r'\$?[A-Za-z]{1,3}\$?[1-9][0-9]*',t.value):
                values.append(sympy.Symbol(t.value.replace('$','').upper()))
            elif t.subtype=='NUMBER':values.append(sympy.Rational(t.value))
            else:raise ValueError('unsupported_audit_operand')
        elif t.type=='OPERATOR-PREFIX' and t.value in ('+','-'):ops.append('u'+t.value)
        elif t.type=='PAREN' and t.subtype=='OPEN':ops.append('(')
        elif t.type=='PAREN' and t.subtype=='CLOSE':
            while ops and ops[-1]!='(':reduce()
            if not ops:raise ValueError('audit_parentheses')
            ops.pop()
        elif t.type=='OPERATOR-INFIX' and t.value in precedence:
            while ops and ops[-1]!='(' and precedence[ops[-1]]>=precedence[t.value]:reduce()
            ops.append(t.value)
        else:raise ValueError('unsupported_audit_token')
    while ops:reduce()
    if len(values)!=1:raise ValueError('audit_stack')
    return values[0]

def equals(a,b):
    try:return sympy.expand(symbolic(a)-symbolic(b))==0
    except (ValueError,IndexError,KeyError):return False

def main():
    verify()
    score=read(RUN/'score.json'); truths=read(DATA/'truth.json')['cases']
    counts={n:Counter() for n in NAMES}; candidates=[]; relation_count=0
    for item,t,scored in zip(read(RUN/'public_index.json'),truths,score['cases']):
        assert item['id']==t['id']==scored['id']
        w=load_workbook(ROOT/item['input'])
        forms={ws.title+'!'+c.coordinate:c.value for ws in w for row in ws for c in row if c.data_type=='f'}
        pred=read(RUN/'predictions'/(item['id']+'.json.gz'))
        assert forms==pred['observed_formulas']
        expected=t['expected_formulas']
        errors={k for k,f in expected.items() if not equals(forms[k],f)}
        assert errors==set(t['errors'])
        for n in NAMES:
            count=Counter(known_errors=len(errors))
            for r in pred[n]['rows']:
                suggestion=r.get('suggested_formula'); label=r['cell']
                if not suggestion or suggestion==forms[label] or (n!='paper' and r['status']!='propose'):continue
                count['suggestions']+=1
                if label not in expected:count['unknown']+=1
                elif equals(suggestion,expected[label]):
                    count['confirmed_correct']+=1
                    count['correct_errors']+=label in errors
                else:
                    count['confirmed_wrong']+=1
                    candidates.append({'id':item['id'],'method':n,'cell':label,'observed':forms[label],
                                       'suggested':suggestion,'expected':expected[label]})
                if n=='note_only':
                    s,a=label.rsplit('!',1); ws=w[s]; e=r['annotation_evidence']
                    note=e['note_cell'].rsplit('!',1)[1]
                    assert ws[note].value==e['note_text']
                    assert not ws.column_dimensions[ws[note].column_letter].hidden
                    assert not ws.row_dimensions[ws[note].row].hidden
                    terms=re.split(r'\s+[x*]\s+|\s*×\s*',ws[note].value.strip())
                    assert len(terms)==len(e['bindings'])
                    for term,b in zip(terms,e['bindings']):
                        assert term.strip().casefold()==str(ws[b['caption_cell']].value).strip().casefold()
                    product='='+'*'.join(b['value_cell'] for b in e['bindings'])
                    assert equals(suggestion,product)
                    for source in r['candidate_sources']:
                        if source['kind']=='same_row_peer':
                            ps,pa=source['cell'].rsplit('!',1)
                            assert Translator(w[ps][pa].value,origin=pa).translate_formula(a)==suggestion
                    relation_count+=1
            for key in ['suggestions','confirmed_correct','confirmed_wrong','unknown','known_errors','correct_errors']:
                assert count[key]==scored[n][key],(item['id'],n,key)
                counts[n][key]+=count[key]
        w.close()
    for n in NAMES:
        for k,v in counts[n].items():assert score['all_scenarios'][n][k]==v
        c=counts[n]; total=c['suggestions']; errors=c['known_errors']
        assert score['all_scenarios'][n]['precision_confirmed_lower_bound']==(c['confirmed_correct']/total if total else None)
        assert score['all_scenarios'][n]['known_injected_error_coverage']==(c['correct_errors']/errors if errors else None)
    out=ROOT/'qa/ANNOTATION_DEV14_INDEPENDENT_AUDIT_V2_20261018.json'
    assert not out.exists()
    save(out,{'cases':len(truths),'engine':'independent openpyxl token stack + exact SymPy rational expressions',
              'counts':counts,'all_metric_counts_match':True,'all_formula_observations_match_openpyxl':True,
              'annotation_candidates_provenance_checked':relation_count,'confirmed_wrong_details':candidates,
              'prediction_lock_sha256':sha(RUN/'prediction_lock.json'),'audit_script_sha256':sha(Path(__file__)),
              'main_candidate_adopted':False,'final_validation_eligible':False})
    print('independent counts and provenance verified',len(truths),relation_count)
if __name__=='__main__':main()
