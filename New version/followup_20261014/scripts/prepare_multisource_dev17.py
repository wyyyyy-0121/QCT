"""Author-only reviewed relation scopes and complete workbook variants."""
from pathlib import Path
import shutil
import sys
import warnings
ROOT=Path(__file__).resolve().parents[1];REPO=ROOT.parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from run_selector_dev5 import read,save,sha
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as letter
from product_sum_oracle_dev17 import equivalent,witness
DATA=ROOT/'development/multisource_dev17_inputs_v1'
IDS=['public03','public09','public10']
SCENARIOS=['clean_original','clean_populated','isolated_operator','isolated_reference','coefficient',
           'common_region','multiple_regions','legitimate_exception','hidden_region','wrong_text_with_error']

def registration(book,w):
    expected={};records=[];primary=[];secondary=[]
    def add(ws,a,relation,kind,role_cells):
        assert ws[a].data_type=='f',(book,ws.title,a)
        assert equivalent(ws[a].value,relation,ws.title,w),(book,ws.title,a,ws[a].value,relation)
        label=ws.title+'!'+a;expected[label]=ws[a].value
        records.append({'cell':label,'business_relation':relation,'kind':kind,
                        'visible_role_cells':{c:ws[c].value for c in role_cells}})
    if book=='public03':
        ws=w['Cash Flow Budget']
        assert [ws[a].value.strip().casefold() for a in ['A9','A26','A29','A52','A56','A57','A58']]==[
            'cash inflows','total cash inflows','cash outflows','total cash outflows',
            'net cash flow','beginning cash balance','cumulative net cash flow']
        assert ws['N9'].value=='All Periods'
        for r in list(range(10,26))+list(range(30,52)):
            add(ws,'N'+str(r),'=SUM('+','.join(letter(c)+str(r) for c in range(2,14))+')',
                'sum_each_item_across_all_periods',['N9','N29'])
        for col in range(2,15):
            c=letter(col)
            for target,start,end in [(26,10,25),(52,30,51)]:
                terms=[f'{c}{r}' for r in range(start,end+1)] if col<14 else [f'{letter(k)}{target}' for k in range(2,14)]
                add(ws,c+str(target),'=SUM('+','.join(terms)+')','sum_inflow_or_outflow',['A'+str(target),'N9'])
            add(ws,c+'56',f'={c}26-{c}52','net_cash_equals_inflow_minus_outflow',['A26','A52','A56'])
            primary.append(ws.title+'!'+c+'56')
            relation=(f'={c}56+{c}57' if col==2 else f'={letter(col-1)}58+{c}56') if col<14 else '=N57+N56'
            add(ws,c+'58',relation,'cash_cumulative_rollforward',['A56','A57','A58'])
            secondary.append(ws.title+'!'+c+'58')
        add(ws,'N57','=B57','initial_cash_for_all_periods',['A57','N55'])
        assert len(expected)==91
    elif book=='public09':
        sheets=[('Enter Your Model Here',18,19,6,[11,13,15,17]),
                ('Optimistic Inucrring Technical ',12,13,5,[8,9,10,11]),
                ('Realistic Steady-State',12,13,5,[8,9,10,11])]
        for name,total,cost,base,percentages in sheets:
            ws=w[name]
            assert ws[f'D{total}'].value=='Yearly % Base:'
            assert ws[f'D{cost}'].value=='Yearly Cost'
            assert ws[f'B{base}'].value=='Base Annual Development Cost:'
            assert all('% of Base' in ws[f'B{r}'].value for r in percentages)
            add(ws,f'E{total}','='+'+'.join(f'E{r}' for r in percentages),'sum_four_annual_cost_percentages',
                [f'B{r}' for r in percentages]+[f'D{total}'])
            add(ws,f'E{cost}',f'=E{total}*E{base}','annual_percent_times_signed_base_cost',
                [f'D{total}',f'D{cost}',f'B{base}'])
            primary.append(name+f'!E{cost}');secondary.append(name+f'!E{total}')
        assert len(expected)==6
    else:
        ws=w['Debt']
        assert [ws[f'B{r}'].value for r in [8,9,10,11]]==[
            'Opening Balance','Drawdowns','Principal Repayment','Closing Balance']
        assert all(ws[f'C{r}'].value=='k£' for r in [8,9,10,11])
        for col in range(6,45):
            c=letter(col);prev=letter(col-1)
            assert 'PPMT(' in ws[c+'10'].value
            add(ws,c+'8',f'={prev}11','opening_equals_previous_period_closing',['B8','B11','C8'])
            add(ws,c+'11',f'={c}8+{c}9+{c}10','closing_equals_opening_drawdown_signed_principal',
                ['B8','B9','B10','B11','C11'])
            primary.append('Debt!'+c+'11');secondary.append('Debt!'+c+'8')
        assert len(expected)==78
    return expected,records,primary,secondary

def prepare():
    if DATA.exists():raise SystemExit('Preserve prepared development data')
    source_index={i['id']:i for i in read(ROOT/'development/public_candidate_pipeline_dev10/public_index.json') if i['id'] in IDS}
    index=[];truths=[];sources=[];scopes={}
    for bi,book in enumerate(IDS,1):
        source=source_index[book];p=REPO/source['relative_path'];assert sha(p)==source['sha256']
        w=load_workbook(p);expected,records,primary,secondary=registration(book,w)
        scopes[book]={'business_scope':records,'known_formula_count':len(expected),
                      'primary_cells':primary,'secondary_cells':secondary,
                      'scope_limit':'reviewed local relations only; upstream and other formulas not certified'}
        w.close();sources.append(source)
        for vi,scenario in enumerate(SCENARIOS,1):
            w=load_workbook(p);truth=dict(expected);changes=[];metadata=[]
            def change(s,a,v,kind):
                ws=w[s];old=ws[a].value;ws[a]=v
                changes.append({'sheet':s,'cell':a,'before':old,'after':v,'kind':kind})
            if scenario!='clean_original' and book=='public03':
                s='Cash Flow Budget'
                for r in list(range(10,26))+list(range(30,52)):
                    for c in range(2,14):
                        change(s,letter(c)+str(r),100+r*7+c*11,'numeric_input_scenario')
                change(s,'B57',321,'numeric_input_scenario')
            def mutate(label,kind):
                s,a=label.rsplit('!',1);f=w[s][a].value
                if kind=='coefficient':new=f+'*1.1'
                elif kind=='reference':
                    if book=='public03':new=f.replace('26','25',1)
                    elif book=='public09':new=f.replace('*E','*F',1)
                    else:new=f.replace('8:', '7:',1)
                elif book=='public03':new=f.replace('-','+',1)
                elif book=='public09':new=f.replace('*','+',1)
                else:
                    # Replace a SUM with a same-reference, syntactically valid
                    # sign error; do not change upstream PPMT semantics.
                    c=w[s][a].column_letter
                    new=f'={c}8+{c}9-{c}10'
                assert new!=f,(label,kind)
                change(s,a,new,'formula_error_injection')
            if scenario in ['isolated_operator','hidden_region','wrong_text_with_error']:
                mutate(primary[0],'operator')
            elif scenario=='isolated_reference':mutate(primary[0],'reference')
            elif scenario=='coefficient':mutate(primary[0],'coefficient')
            elif scenario=='common_region':
                for label in primary:mutate(label,'operator')
            elif scenario=='multiple_regions':
                for label in [primary[0],primary[-1]]:mutate(label,'coefficient')
                label=secondary[len(secondary)//2];s,a=label.rsplit('!',1)
                change(s,a,'=('+w[s][a].value[1:]+')*1.1','formula_error_injection')
            elif scenario=='legitimate_exception':
                label=primary[0];s,a=label.rsplit('!',1)
                new='=('+w[s][a].value[1:]+')+10'
                change(s,a,new,'legitimate_one_off_adjustment');truth[label]=new
                note=f'Exception: {a} includes a one-off adjustment of 10 in this scenario.'
                change(s,'P'+str(w[s][a].row) if book=='public03' else 'B200',note,'visible_exception_note')
            if scenario=='hidden_region':
                s,a=primary[0].rsplit('!',1);row=w[s][a].row
                w[s].row_dimensions[row].hidden=True
                metadata.append({'sheet':s,'row':row,'hidden':True})
            if scenario=='wrong_text_with_error':
                s,a=primary[0].rsplit('!',1)
                # Deliberately false, visible assertion, never oracle authority.
                text={'public03':'Net Cash Flow equals Total Cash Inflows plus Total Cash Outflows.',
                      'public09':'Yearly Cost equals Yearly % Base plus Base Annual Development Cost.',
                      'public10':'Closing Balance subtracts the signed negative Principal Repayment.'}[book]
                change(s,'P'+str(w[s][a].row) if book=='public03' else 'B200',text,'intentionally_wrong_business_text')
            case=f'ms{bi:02d}_{vi:02d}';out=DATA/'workbooks'/(case+'.xlsx')
            out.parent.mkdir(parents=True,exist_ok=True)
            if scenario=='clean_original':shutil.copyfile(p,out)
            else:w.save(out)
            observed={ws.title+'!'+c.coordinate:c.value for ws in w for row in ws for c in row if c.data_type=='f'}
            errors=[];witnesses={}
            for label,f in truth.items():
                s,a=label.rsplit('!',1)
                if not equivalent(observed[label],f,s,w):
                    errors.append(label);witnesses[label]=witness(observed[label],f,s,w)
            index.append({'id':case,'input':str(out.relative_to(ROOT)),'input_sha256':sha(out),
                          'parallel_sheets':[],'origin':source['origin'],'source_id':book,
                          'source_sha256':source['sha256'],'final_validation_eligible':False})
            truths.append({'id':case,'origin':source['origin'],'source_id':book,'scenario':scenario,
                           'expected_formulas':truth,'observed_formulas':observed,'errors':errors,
                           'rational_difference_witnesses':witnesses,'changes':changes,'metadata_changes':metadata,
                           'known_scope_count':len(truth),'unknown_formula_count':len(observed)-len(truth)})
            w.close();assert sha(p)==source['sha256']
            print(case,book,scenario,'known',len(truth),'errors',len(errors),flush=True)
    save(DATA/'public_index.json',index)
    save(DATA/'truth.json',{'cases':truths,'sources':sources,'scopes':scopes,
          'final_validation_eligible':False,'business_evidence':'visible original headings and locally reviewed accounting identities; not expert certification',
          'oracle':'independent openpyxl token stack and exact symbolic rational arithmetic with SUM blank/text semantics'})
    save(DATA/'preparation_lock.json',{'files':{str(p.relative_to(DATA)):sha(p) for p in sorted(DATA.rglob('*')) if p.is_file()},
         'author_script_sha256':sha(Path(__file__)),'oracle_sha256':sha(ROOT/'scripts/product_sum_oracle_dev17.py')})
if __name__=='__main__':prepare()
