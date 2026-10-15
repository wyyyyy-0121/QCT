"""Registered, revealed development cases outside the old profit template.

Four business mechanisms, not 32 independent families. prepare/predict/score are
separate processes. No input is a final holdout. An independent Fraction oracle
checks every expected formula, including roll-forwards and summary rows.
"""
from collections import Counter
from fractions import Fraction
from pathlib import Path
import argparse
import gzip
import hashlib
import importlib.util
import json
import math
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'baseline'
RUN = ROOT / 'development/crossdomain_dev1'
sys.path.insert(0, str(BASE / 'source/HybridGuard_V5_handoff_20260904/src'))
from formulaguard.workbook import WorkbookModel
from formulaguard.v5_structural_guard import v5_structural_guard_scores
from structuralguard.detector import analyze_workbook
from selector_dev2 import select as dev2_select
from selector_dev4 import review_guard

FAMILIES = ['inventory_rollforward', 'loan_amortization', 'budget_allocation', 'unit_conversion']
VARIANTS = ['clean', 'sparse_error', 'whole_column_error', 'unmarked_exception',
            'marked_exception', 'exception_and_error', 'copied_error_majority', 'multirow_regions']
HEADERS = {
    'inventory_rollforward': ['Period', 'Opening inventory [units]', 'Receipts [units]',
                             'Issues [units]', 'Write-offs [units]', 'Closing inventory [units]'],
    'loan_amortization': ['Period', 'Opening loan [USD]', 'Monthly rate [1/month]',
                         'Interest [USD]', 'Payment [USD]', 'Closing loan [USD]'],
    'budget_allocation': ['Department', 'Budget [USD]', 'Allocation weight',
                         'Total weight', 'Adjustment [USD]', 'Allocated budget [USD]'],
    'unit_conversion': ['Sample', 'Temperature [C]', 'Scale [F/C]', 'Offset [F]',
                        'Sensor correction [F]', 'Temperature [F]'],
}
POLICIES = {
    'inventory_rollforward': 'Row {row}: write-offs are recorded for disclosure only; exclude them from closing inventory.',
    'loan_amortization': 'Row {row}: payment is deferred; the quoted payment is not deducted from the closing loan.',
    'budget_allocation': 'Row {row}: this department receives half of its ordinary weighted allocation, plus the adjustment.',
    'unit_conversion': 'Row {row}: sensor correction is not applied to this calibration sample.',
}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(data, ensure_ascii=False, sort_keys=True, indent=2).encode()
    if path.suffix == '.gz':
        with gzip.GzipFile(filename=str(path), mode='wb', mtime=0) as handle:
            handle.write(payload)
    else:
        path.write_bytes(payload + b'\n')


def read(path):
    return json.loads(gzip.decompress(path.read_bytes()) if path.suffix == '.gz' else path.read_bytes())


def correct_formula(family, row, exception=False):
    r = row
    if family == 'inventory_rollforward':
        return f'=B{r}+C{r}-D{r}' + ('' if exception else f'-E{r}')
    if family == 'loan_amortization':
        return f'=B{r}+D{r}' + ('' if exception else f'-E{r}')
    if family == 'budget_allocation':
        return '=' + ('0.5*' if exception else '') + f'B{r}*C{r}/D{r}+E{r}'
    return f'=B{r}*C{r}+D{r}' + ('' if exception else f'+E{r}')


def corrupt_formula(family, row):
    formula = correct_formula(family, row)
    term = {'inventory_rollforward': f'-D{row}', 'loan_amortization': f'-E{row}',
            'budget_allocation': f'/D{row}', 'unit_conversion': f'+D{row}'}[family]
    replacement = {'inventory_rollforward': f'+D{row}', 'loan_amortization': f'+E{row}',
                   'budget_allocation': f'*D{row}', 'unit_conversion': f'-D{row}'}[family]
    return formula.replace(term, replacement)


def build_case(family, variant):
    from openpyxl import Workbook
    book = Workbook(); book.remove(book.active)
    expected = {}; numeric = {}; mutation_log = []
    first = 5 if variant == 'multirow_regions' else 2
    last = first + 23
    exception_row, error_row = first + 8, first + 14
    for sheet_index, sheet in enumerate(['A', 'B', 'C', 'D', 'E']):
        ws = book.create_sheet(sheet)
        if variant == 'multirow_regions':
            ws['A1'] = family.replace('_', ' ').title()
            ws['B2'] = 'Inputs'; ws['F2'] = 'Calculated output'
            ws['A3'] = 'Regional schedule, followed by a separate reconciliation block'
        for col, header in zip('ABCDEF', HEADERS[family]):
            ws[f'{col}{first-1}'] = header
        exception = sheet == 'A' and variant in ['unmarked_exception', 'marked_exception', 'exception_and_error']
        if exception and variant != 'unmarked_exception':
            ws['H1'] = 'Business policy'
            ws['H2'] = POLICIES[family].format(row=exception_row)
        previous = Fraction(5000 if family == 'loan_amortization' else 300)
        total_weight = sum(Fraction(1 + (r-first) % 5) for r in range(first, last+1))
        for row in range(first, last+1):
            i = row-first
            ws[f'A{row}'] = i+1
            is_exception = exception and row == exception_row
            b = previous if family in ['inventory_rollforward', 'loan_amortization'] else Fraction(1000 + sheet_index*70 + i*13)
            c = (Fraction(20+i % 4) if family == 'inventory_rollforward' else
                 Fraction(1, 100) if family == 'loan_amortization' else
                 Fraction(1+i % 5) if family == 'budget_allocation' else Fraction(9, 5))
            if family == 'unit_conversion':
                b = Fraction(-15+i*3+sheet_index)
            d = (Fraction(10+i % 3) if family == 'inventory_rollforward' else
                 b*c if family == 'loan_amortization' else
                 total_weight if family == 'budget_allocation' else Fraction(32))
            e = (Fraction(1+i % 2) if family == 'inventory_rollforward' else
                 Fraction(120) if family == 'loan_amortization' else
                 Fraction(i % 3-1) if family == 'budget_allocation' else Fraction(1, 2))
            for col, value in [('B', b), ('C', c), ('D', d), ('E', e)]:
                ws[f'{col}{row}'] = float(value)
                numeric[sheet, f'{col}{row}'] = value
            if family in ['inventory_rollforward', 'loan_amortization'] and row > first:
                ws[f'B{row}'] = f'=F{row-1}'
                expected[f'{sheet}!B{row}'] = ws[f'B{row}'].value
            if family == 'loan_amortization':
                ws[f'D{row}'] = f'=B{row}*C{row}'
                expected[f'{sheet}!D{row}'] = ws[f'D{row}'].value
            elif family == 'budget_allocation':
                ws[f'D{row}'] = f'=SUM($C${first}:$C${last})'
                expected[f'{sheet}!D{row}'] = ws[f'D{row}'].value
            # An independent arithmetic oracle, not generated by evaluating formula text.
            if family == 'inventory_rollforward':
                out = b + c - d - (0 if is_exception else e)
            elif family == 'loan_amortization':
                out = b * (1+c) - (0 if is_exception else e)
            elif family == 'budget_allocation':
                out = b * (Fraction(1,2) if is_exception else 1) * c / total_weight + e
            else:
                out = b * Fraction(9,5) + 32 + (0 if is_exception else e)
            numeric[sheet, f'F{row}'] = out
            standard = correct_formula(family, row, is_exception)
            observed = standard
            if ((variant == 'sparse_error' and sheet == 'A' and row == error_row) or
                (variant == 'whole_column_error' and sheet == 'A') or
                (variant == 'exception_and_error' and sheet == 'A' and row == error_row) or
                (variant == 'copied_error_majority' and sheet in ['A', 'B', 'C', 'D']) or
                (variant == 'multirow_regions' and sheet == 'A' and row == error_row)):
                observed = corrupt_formula(family, row)
            ws[f'F{row}'] = observed
            expected[f'{sheet}!F{row}'] = standard
            if observed != standard:
                mutation_log.append({'cell': f'{sheet}!F{row}', 'before': standard, 'after': observed})
            previous = out
        if variant == 'multirow_regions':
            summary = last+2
            ws[f'E{summary}'] = 'Output total'
            ws[f'F{summary}'] = f'=SUM(F{first}:F{last})'
            expected[f'{sheet}!F{summary}'] = ws[f'F{summary}'].value
            numeric[sheet, f'F{summary}'] = sum(numeric[sheet, f'F{r}'] for r in range(first,last+1))
            # A genuinely different region in the same output column must be preserved.
            region = summary+4
            ws[f'E{region-1}'] = 'USD amount'; ws[f'F{region-1}'] = 'USD cents'
            for r in range(region, region+4):
                ws[f'E{r}'] = r+3
                ws[f'F{r}'] = f'=E{r}*100'
                expected[f'{sheet}!F{r}'] = ws[f'F{r}'].value
                numeric[sheet, f'F{r}'] = Fraction((r+3)*100)
    return book, expected, numeric, mutation_log


def prepare():
    if RUN.exists():
        raise SystemExit('Preserve existing crossdomain development evidence')
    public = []; truths = []; receipts = []
    # Registration is written before model predictions, and is not a model input.
    save(RUN / 'registration.json', {'role': 'revealed authored development only',
         'families': FAMILIES, 'variants': VARIANTS, 'case_count': 32,
         'new_final_sources': 0, 'new_final_families': 0,
         'truth_validation': 'Fraction arithmetic independently checked against every expected formula',
         'copied_majority': 'Four sheets intentionally share one corrupted ancestor; not independent votes',
         'scoring': 'strict expected formula match; every output and intermediate formula retained'})
    for fi, family in enumerate(FAMILIES, 1):
        for vi, variant in enumerate(VARIANTS, 1):
            case = f'x{fi}{vi}'
            book, expected, numeric, mutations = build_case(family, variant)
            path = RUN / 'inputs' / f'{case}.xlsx'; path.parent.mkdir(parents=True,exist_ok=True)
            book.save(path)
            model = WorkbookModel.from_xlsx(path)
            observed = {s+'!'+a: v for (s,a),v in model.formulas.items()}
            assert set(observed) == set(expected)
            overrides = {tuple(k.rsplit('!',1)): v for k,v in expected.items()}
            values, errors = model.evaluate(overrides=overrides)
            assert not errors, (case,errors)
            for label in expected:
                key = tuple(label.rsplit('!',1))
                assert math.isclose(values[key],float(numeric[key]),rel_tol=1e-10,abs_tol=1e-8), (case,key)
            # The corrupted inputs must still evaluate normally: silent formula errors.
            _, observed_errors = model.evaluate()
            assert not observed_errors, (case,observed_errors)
            public.append({'id':case, 'input':str(path.relative_to(RUN)), 'sha256':sha(path),
                           'parallel_sheets':['A','B','C','D','E']})
            truths.append({'id':case,'family':family,'scenario':variant,'expected_formulas':expected,
                           'errors':[k for k in expected if expected[k]!=observed[k]],
                           'mutation_log':mutations})
            receipts.append({'id':case,'checked_formulas':len(expected),'silent_input':True,
                             'correct_oracle':'independent Fraction calculation','input_sha256':sha(path)})
    save(RUN / 'public_index.json', public)
    save(RUN / 'truth.json', {'role':'revealed development; not final validation','cases':truths})
    save(RUN / 'qa/input_oracle_check.json', receipts)
    print('Prepared',len(public),'cases in four development families; independent oracles passed')


def predict():
    if (RUN / 'predictions').exists():
        raise SystemExit('Preserve existing prediction evidence')
    for name in ['source_lock.json','prediction_lock.json']:
        for relative, expected in read(BASE / 'protocol' / name)['files'].items():
            assert sha(BASE / relative)==expected,relative
    prior = ROOT / 'development/selector_dev4'
    lock = read(prior / 'prediction_lock.json')
    assert sha(prior/'source_lock.json')==lock['source_lock_sha256']
    for relative,expected in read(prior/'source_lock.json')['files'].items():
        assert sha(ROOT/relative)==expected,relative
    sources = [Path(__file__), BASE/'scripts/predict.py', ROOT/'scripts/selector_dev2.py',
               ROOT/'scripts/selector_dev3.py',ROOT/'scripts/selector_dev4.py']
    sources += sorted((BASE/'source').rglob('*.py'))
    save(RUN/'source_lock.json', {'files':{str(p.relative_to(ROOT)):sha(p) for p in sources}})
    spec = importlib.util.spec_from_file_location('crossdomain_paper_adapter',BASE/'scripts/predict.py')
    adapter=importlib.util.module_from_spec(spec);spec.loader.exec_module(adapter)
    accesses=set()
    def guard(event,args):
        if event=='open' and isinstance(args[0],(str,bytes)):
            name=str(args[0])
            if any(x in name.lower() for x in ['truth.json','registration.json','input_oracle',
                                               'mutation_log','failure_audit','/score','/results/']):
                raise PermissionError('Prediction cannot read scoring evidence: '+name)
            if str(RUN) in name: accesses.add(name)
    sys.addaudithook(guard)
    files={}
    for entry in read(RUN/'public_index.json'):
        path=RUN/entry['input']; assert sha(path)==entry['sha256']
        t=time.perf_counter(); model=WorkbookModel.from_xlsx(path);before=dict(model.formulas)
        structural=analyze_workbook(path)['results']
        fg=[{'rank':i,'cell':v.cell_label,'score':v.score,'candidate_formula':v.candidate_formula,
             'evidence':v.evidence} for i,v in enumerate(v5_structural_guard_scores(model),1)]
        forms={s+'!'+a:v for (s,a),v in model.formulas.items()}
        paper=adapter.hybrid((model,structural,fg,forms,entry['parallel_sheets']),{})
        dev2=dev2_select(model,paper,fg,structural,entry['parallel_sheets'])
        dev4=review_guard(model,dev2)
        assert model.formulas==before
        target=RUN/'predictions'/(entry['id']+'.json.gz')
        save(target,{'paper':paper,'dev4':dev4,'observed_formulas':forms,
                     'runtime_seconds':time.perf_counter()-t,'automatic_edit_applied':False})
        files[str(target.relative_to(RUN))]=sha(target)
        print(entry['id'],'proposals',sum(r['status']=='propose' for r in dev4['rows']),flush=True)
    save(RUN/'prediction_lock.json',{'files':files,'source_lock_sha256':sha(RUN/'source_lock.json'),
         'public_index_sha256':sha(RUN/'public_index.json'),'labels_read':[],'accesses':sorted(accesses),
         'role':'revealed development only; does not satisfy final separation'})


def score():
    if (RUN/'score.json').exists(): raise SystemExit('Preserve existing score evidence')
    lock=read(RUN/'prediction_lock.json')
    assert sha(RUN/'source_lock.json')==lock['source_lock_sha256']
    assert sha(RUN/'public_index.json')==lock['public_index_sha256']
    for relative,expected in lock['files'].items(): assert sha(RUN/relative)==expected,relative
    for relative,expected in read(RUN/'source_lock.json')['files'].items(): assert sha(ROOT/relative)==expected,relative
    for entry in read(RUN/'public_index.json'): assert sha(RUN/entry['input'])==entry['sha256']
    cases=[]
    for truth in read(RUN/'truth.json')['cases']:
        pred=read(RUN/'predictions'/(truth['id']+'.json.gz'))
        result={k:truth[k] for k in ['id','family','scenario']};result['errors']=len(truth['errors'])
        for name in ['paper','dev4']:
            rows=pred[name]['rows']; assert len(rows)==len({r['cell'] for r in rows})==len(truth['expected_formulas'])
            assert {r['cell'] for r in rows}==set(truth['expected_formulas'])
            proposals=[r for r in rows if r.get('suggested_formula') and
                       r['suggested_formula']!=pred['observed_formulas'][r['cell']] and
                       (name=='paper' or r['status']=='propose')]
            correct=sum(r['suggested_formula']==truth['expected_formulas'][r['cell']] for r in proposals)
            result[name]={'proposals':len(proposals),'correct':correct,'wrong':len(proposals)-correct,
                          'harmful_correct_cells':sum(r['cell'] not in truth['errors'] for r in proposals),
                          'precision':correct/len(proposals) if proposals else None,
                          'coverage':correct/len(truth['errors']) if truth['errors'] else None,
                          'statuses':dict(Counter(r.get('status','unspecified') for r in rows)),
                          'reason_counts':dict(Counter(r.get('reason','unspecified') for r in rows)),
                          'examples':[{'cell':r['cell'],'observed':pred['observed_formulas'][r['cell']],
                                      'proposed':r['suggested_formula'],'expected':truth['expected_formulas'][r['cell']]}
                                      for r in proposals[:3]]}
        result['runtime_seconds']=pred['runtime_seconds'];cases.append(result)
        print(truth['id'],truth['family'],truth['scenario'],'errors',result['errors'],
              'paper',result['paper']['correct'],result['paper']['wrong'],
              'dev4',result['dev4']['correct'],result['dev4']['wrong'],flush=True)
    totals={}
    for family in FAMILIES+['all_development_families']:
        selected=[c for c in cases if c['family']==family or family=='all_development_families']
        totals[family]={'errors':sum(c['errors'] for c in selected)}
        for name in ['paper','dev4']:
            p=sum(c[name]['proposals'] for c in selected); k=sum(c[name]['correct'] for c in selected)
            e=totals[family]['errors']
            totals[family][name]={'proposals':p,'correct':k,'wrong':p-k,'precision':k/p if p else None,
                                 'coverage':k/e if e else None,
                                 'harmful_correct_cells':sum(c[name]['harmful_correct_cells'] for c in selected)}
    save(RUN/'score.json',{'role':'revealed cross-domain diagnostic only; not acceptance',
         'cases':cases,'totals':totals,'prediction_lock_sha256':sha(RUN/'prediction_lock.json')})
    print(json.dumps(totals,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('stage',choices=['prepare','predict','score'])
    args=parser.parse_args();{'prepare':prepare,'predict':predict,'score':score}[args.stage]()
