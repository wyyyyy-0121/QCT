"""Post-lock human-readable business-scope review; frozen score is never rewritten.

The new scope is explicitly selected after viewing model outputs, so it is
development evidence only and must not be used to claim independent validation.
"""
from pathlib import Path
import sys
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'baseline/source/HybridGuard_V5_handoff_20260904/src')]
from run_signed_learning_dev19 import verify,RUN,BASE,case_path
from run_selector_dev5 import read,save,sha
from product_sum_oracle_dev17 import equivalent,witness
from audit_multisource_dev17 import equals

def main():
    verify();score=read(RUN/'score.json');assert score['origin_group']['uk_government']['learned_only']['unknown']==80
    score_hash=sha(RUN/'score.json');rows=[];unique=set()
    for item in read(BASE/'public_index.json'):
        if item['origin_group']!='uk_government':continue
        case=item['case'];w=load_workbook(ROOT/item['input'])
        pred=read(RUN/'predictions/source__uk_government'/(case_path(case)+'.json.gz'))
        for r in pred['learned_only']:
            s,a=r['cell'].rsplit('!',1);ws=w[s];column=ws[a].column_letter
            assert ws[a].row==45
            assert ws['B42'].value=='Capital cost after accounting for cost adjustment'
            assert ws['B43'].value=='Infrastructure cost'
            assert ws['B45'].value=='Totl construction costs'
            assert ws['E45'].value=='Sum of capital costs + infrastructure costs'
            assert all(ws[f'C{n}'].value=='£m' for n in [42,44,45])
            expected='='+column+'42+'+column+'44'
            assert equivalent(ws[a].value,expected,s,w) and equals(ws[a].value,expected,s,w)
            assert not equivalent(r['suggested_formula'],expected,s,w)
            assert not equals(r['suggested_formula'],expected,s,w)
            proof=witness(r['suggested_formula'],expected,s,w)
            unique.add((s,a))
            rows.append({'case':case,'cell':r['cell'],'observed_formula':ws[a].value,
                         'expected_formula':expected,'suggested_formula':r['suggested_formula'],
                         'correct_before':True,'repair_correct':False,'rational_relation_witness':proof,
                         'evidence':'Visible construction-total description requires sum; both contributing costs and total are in GBP million. Product has incompatible squared currency units.',
                         'scope':'local additive cost relationship only; upstream costs and actual Excel numeric impact not certified'})
        w.close()
    assert len(rows)==80 and len(unique)==5 and sha(RUN/'score.json')==score_hash
    path=ROOT/'qa/SIGNED_LEARNING_DEV19_POSTHOC_BUSINESS_REVIEW_20261018.json';assert not path.exists()
    save(path,{'outputs_reviewed':80,'unique_original_cells':5,'confirmed_harmful_suggestions':80,
         'rows':rows,'prediction_lock_sha256':sha(RUN/'prediction_lock.json'),'frozen_score_sha256':score_hash,
         'frozen_score_unchanged':True,'scope_selected_after_output_inspection':True,'final_validation_eligible':False,
         'supplemental_posthoc_learner_counts':{'suggestions':87,'correct':7,'wrong':80,'unknown':0,'precision':7/87},
         'review_script_sha256':sha(Path(__file__))})
    print('Posthoc review: 80 harmful suggestions across 5 original cells; frozen score untouched',flush=True)

if __name__=='__main__':main()
