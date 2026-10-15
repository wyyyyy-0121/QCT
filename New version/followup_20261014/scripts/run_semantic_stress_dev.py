"""Revealed development challenge: headers do not establish business policy.

Stages preserve inputs, lock actual paper/dev2 predictions, then score separately.
This deliberately includes an observationally identical pair with different external
policies. Case results are diagnostic; pooling them is not an acceptance metric.
"""
from pathlib import Path
import argparse
import hashlib
import importlib.util
import json
import sys

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / 'development/semantic_stress_dev1'
BASE = ROOT / 'baseline'
sys.path.insert(0, str(BASE / 'source/HybridGuard_V5_handoff_20260904/src'))
from formulaguard.workbook import WorkbookModel
from formulaguard.v5_structural_guard import v5_structural_guard_scores
from structuralguard.detector import analyze_workbook
from selector_dev2 import select


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')


def read(path):
    return json.loads(path.read_text())


def prepare():
    from openpyxl import Workbook
    if RUN.exists():
        raise SystemExit('Preserve existing challenge evidence')
    scenarios = [
        ('c01', 'expense_sign_error', 'cost_sign', True, 'Net Profit', ''),
        ('c02', 'missing_shipping_error', 'omit_shipping_one', True, 'Net Profit', ''),
        ('c03', 'legitimate_shipping_waiver', 'omit_shipping_one', False, 'Net Profit',
         'Row 10: seller absorbs shipping; exclude Shipping from this profit calculation.'),
        ('c04', 'legitimate_discount_exemption', 'omit_discount', False, 'Net Profit',
         'Sheet A: contractual full-price sales; Discount Rate is quoted but not applied.'),
        ('c05', 'legitimate_revenue_factor', 'revenue_factor', False, 'Net Profit', ''),
        ('c06', 'erroneous_revenue_factor_same_observed_input', 'revenue_factor', True, 'Net Profit', ''),
        ('c07', 'unknown_output_role_sign_error', 'cost_sign', True, 'Custom Metric', ''),
        ('c08', 'legitimate_cost_reimbursement', 'cost_sign', False, 'Net Profit',
         'Sheet A: Unit Cost is fully reimbursed as an additional receipt in this contract.'),
    ]
    public = []; truths = []
    for case, scenario, mode, corrupted, output, policy_note in scenarios:
        book = Workbook(); book.remove(book.active)
        expected = {}; observed = {}
        for sheet in ['A', 'B', 'C', 'D', 'E']:
            ws = book.create_sheet(sheet)
            for address, value in {'A1': 'Units', 'B1': 'Unit Price', 'C1': 'Unit Cost',
                                   'D1': 'Shipping', 'E1': 'Discount Rate', 'F1': output}.items():
                ws[address] = value
            if sheet == 'A' and policy_note:
                ws['H1'] = 'Business policy'; ws['H2'] = policy_note
            for row in range(2, 26):
                for col, value in [('A', 4 + row % 3), ('B', 20), ('C', 5), ('D', 2), ('E', .1)]:
                    ws[f'{col}{row}'] = value
                standard = f'=A{row}*B{row}*(1-E{row})-A{row}*C{row}-D{row}'
                formula = standard
                if sheet == 'A':
                    if mode == 'cost_sign':
                        formula = formula.replace(f'-A{row}*C{row}', f'+A{row}*C{row}')
                    elif mode == 'omit_shipping_one' and row == 10:
                        formula = formula.removesuffix(f'-D{row}')
                    elif mode == 'omit_discount':
                        formula = formula.replace(f'*(1-E{row})', '')
                    elif mode == 'revenue_factor':
                        formula = formula.replace(f'A{row}*B{row}', f'0.85*A{row}*B{row}')
                ws[f'F{row}'] = formula
                label = f'{sheet}!F{row}'
                observed[label] = formula
                expected[label] = standard if corrupted else formula
        path = RUN / 'inputs' / f'{case}.xlsx'; path.parent.mkdir(parents=True, exist_ok=True)
        book.save(path)
        public.append({'id': case, 'input': str(path.relative_to(RUN)),
                       'sha256': sha(path), 'parallel_sheets': ['A', 'B', 'C', 'D', 'E']})
        truths.append({'id': case, 'scenario': scenario, 'expected_formulas': expected,
                       'errors': [k for k in expected if expected[k] != observed[k]]})
    # Make the ambiguity pair byte-identical, including OOXML creation metadata.
    (RUN / 'inputs/c06.xlsx').write_bytes((RUN / 'inputs/c05.xlsx').read_bytes())
    public[5]['sha256'] = sha(RUN / 'inputs/c06.xlsx')
    save(RUN / 'public_index.json', public)
    save(RUN / 'truth.json', {'cases': truths, 'role': 'revealed development only'})
    save(RUN / 'registration.json', {
        'role': 'development microfixtures; derived financial pattern; zero new final families',
        'hypotheses': ['A familiar header does not require every listed input to affect profit.',
                       'Visible policy notes must not be overridden by peer majority.',
                       'Identical observed inputs can correspond to different external intended policies.'],
        'case_count': len(public), 'pooled_acceptance_metric': False,
        'ambiguity_pair': ['c05', 'c06'],
        'warning': 'Fixtures are authored during development, not unseen final synthetic data.'})
    print('Prepared', len(public), 'development cases')


def predict():
    if (RUN / 'predictions').exists():
        raise SystemExit('Preserve locked challenge predictions')
    for name in ['source_lock.json', 'prediction_lock.json']:
        for relative, expected in read(BASE / 'protocol' / name)['files'].items():
            assert sha(BASE / relative) == expected, relative
    for relative, expected in read(ROOT / 'development/selector_dev2/source_lock.json')['files'].items():
        assert sha(ROOT / relative) == expected, relative
    sources = [Path(__file__), ROOT / 'scripts/selector_dev2.py', BASE / 'scripts/predict.py']
    sources += sorted((BASE / 'source').rglob('*.py'))
    save(RUN / 'source_lock.json', {'files': {str(p.relative_to(ROOT)): sha(p) for p in sources}})
    spec = importlib.util.spec_from_file_location('stress_paper_adapter', BASE / 'scripts/predict.py')
    adapter = importlib.util.module_from_spec(spec); spec.loader.exec_module(adapter)
    accesses = set()
    def guard(event, args):
        if event == 'open' and isinstance(args[0], (str, bytes)):
            name = str(args[0])
            if any(x in name.lower() for x in ['truth.json', 'registration.json', 'failure_audit', '/score', '/results/']):
                raise PermissionError('Prediction cannot read scoring or case policy: ' + name)
            if str(RUN) in name:
                accesses.add(name)
    sys.addaudithook(guard)
    files = {}
    for entry in read(RUN / 'public_index.json'):
        path = RUN / entry['input']; assert sha(path) == entry['sha256']
        model = WorkbookModel.from_xlsx(path); before = dict(model.formulas)
        structural = analyze_workbook(path)['results']
        fg = [{'rank': i, 'cell': v.cell_label, 'score': v.score,
               'candidate_formula': v.candidate_formula, 'evidence': v.evidence}
              for i, v in enumerate(v5_structural_guard_scores(model), 1)]
        forms = {s + '!' + a: value for (s, a), value in model.formulas.items()}
        paper = adapter.hybrid((model, structural, fg, forms, entry['parallel_sheets']), {})
        dev2 = select(model, paper, fg, structural, entry['parallel_sheets'])
        assert model.formulas == before
        target = RUN / 'predictions' / (entry['id'] + '.json')
        save(target, {'paper': paper, 'dev2': dev2, 'observed_formulas': forms})
        files[str(target.relative_to(RUN))] = sha(target)
        print(entry['id'], 'dev2 proposals', sum(r['status'] == 'propose' for r in dev2['rows']), flush=True)
    save(RUN / 'prediction_lock.json', {'files': files, 'source_lock_sha256': sha(RUN / 'source_lock.json'),
        'public_index_sha256': sha(RUN / 'public_index.json'), 'labels_read': [], 'accesses': sorted(accesses)})


def score():
    if (RUN / 'score.json').exists():
        raise SystemExit('Preserve existing scores')
    lock = read(RUN / 'prediction_lock.json')
    assert sha(RUN / 'source_lock.json') == lock['source_lock_sha256']
    assert sha(RUN / 'public_index.json') == lock['public_index_sha256']
    for relative, expected in lock['files'].items():
        assert sha(RUN / relative) == expected, relative
    for relative, expected in read(RUN / 'source_lock.json')['files'].items():
        assert sha(ROOT / relative) == expected, relative
    cases = []
    for truth in read(RUN / 'truth.json')['cases']:
        prediction = read(RUN / 'predictions' / (truth['id'] + '.json'))
        result = {'id': truth['id'], 'scenario': truth['scenario'], 'errors': len(truth['errors'])}
        for name in ['paper', 'dev2']:
            proposals = [r for r in prediction[name]['rows'] if r.get('suggested_formula')
                and r['suggested_formula'] != prediction['observed_formulas'][r['cell']]
                and (name == 'paper' or r['status'] == 'propose')]
            correct = sum(r['suggested_formula'] == truth['expected_formulas'][r['cell']] for r in proposals)
            result[name] = {'proposals': len(proposals), 'correct': correct,
                            'wrong': len(proposals) - correct,
                            'precision': correct / len(proposals) if proposals else None,
                            'coverage': correct / len(truth['errors']) if truth['errors'] else None,
                            'examples': [{'cell': r['cell'], 'proposal': r['suggested_formula'],
                               'expected': truth['expected_formulas'][r['cell']]} for r in proposals[:2]]}
        cases.append(result)
        print(result['scenario'], {n: {k: v for k, v in result[n].items() if k != 'examples'} for n in ['paper', 'dev2']})
    p5 = read(RUN / 'predictions/c05.json'); p6 = read(RUN / 'predictions/c06.json')
    assert p5 == p6, 'Identical observed inputs should produce identical predictions'
    save(RUN / 'score.json', {'role': 'development diagnostics only; no pooled acceptance score',
        'cases': cases, 'identical_input_pair_predictions_equal': True,
        'prediction_lock_sha256': sha(RUN / 'prediction_lock.json')})


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('stage', choices=['prepare', 'predict', 'score'])
    args = parser.parse_args()
    {'prepare': prepare, 'predict': predict, 'score': score}[args.stage]()
