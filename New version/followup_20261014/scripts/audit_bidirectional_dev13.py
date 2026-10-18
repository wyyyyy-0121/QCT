"""Check new candidate provenance against independently loaded public cells."""
from pathlib import Path
import json
import sys

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from run_bidirectional_dev13 import RUN, read, sha, verify, load_case, VARIANTS
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator


def main():
    verify()
    score = read(RUN/'score.json')
    unchanged = 0
    for group,report in score['groups'].items():
        for case in report['cases']:
            assert all(case[v] == case['dev11'] for v in VARIANTS)
            unchanged += 1
    cases, details = [], []
    for item in read(RUN/'public_index.json'):
        path = (ROOT if item['input_scope']=='root' else REPO)/item['input']
        assert sha(path) == item['input_sha256']
        pred = read(RUN/'predictions'/item['group']/(item['id']+'.json.gz'))
        if item['group'] != 'public':
            continue
        if pred['status'] == 'input_rejected':
            cases.append({'id':item['id'],'status':'input_rejected',
                          'formula_bodies_read_by_audit':False})
            continue
        observed,prior = load_case(item)
        prior_proposals = {r['cell'] for r in prior['rows'] if r.get('status')=='propose'}
        workbook = load_workbook(path,data_only=False,read_only=False)
        array_formula_count = 0
        def formula_text(value):
            nonlocal array_formula_count
            if isinstance(value,str):
                return value
            assert value.__class__.__name__ == 'ArrayFormula'
            array_formula_count += 1
            return value.text
        mapping = {ws.title+'!'+cell.coordinate:formula_text(cell.value)
                   for ws in workbook for row in ws for cell in row if cell.data_type=='f'}
        assert mapping == observed, item['id']
        count = 0
        for variant,rows in pred['variants'].items():
            for label,evidence in rows.items():
                sheet,address = label.rsplit('!',1)
                ws = workbook[sheet]
                candidate = evidence['formula']
                before = [(t.type,t.subtype,t.value) for t in Tokenizer(observed[label]).items if t.type!='WHITE-SPACE']
                after = [(t.type,t.subtype,t.value) for t in Tokenizer(candidate).items if t.type!='WHITE-SPACE']
                assert len(before) == len(after)
                differences = [(a,b) for a,b in zip(before,after) if a!=b]
                assert len(differences) == 1
                assert all(t[:2] == ('OPERAND','RANGE') for t in differences[0])
                axes = evidence.get('axes',{'single':evidence})
                for axis in axes.values():
                    for peer in axis['peers']:
                        ps,pa = peer['cell'].rsplit('!',1)
                        actual = Translator(workbook[ps][pa].value,origin=pa).translate_formula(address)
                        assert actual == peer['translated_formula']
                cell = ws[address]
                context = {ws.cell(r,c).coordinate:ws.cell(r,c).value
                           for r in range(max(1,cell.row-2),cell.row+3)
                           for c in range(max(1,cell.column-2),cell.column+3)
                           if ws.cell(r,c).value is not None}
                details.append({'id':item['id'],'variant':variant,'cell':label,
                                'original_formula':observed[label],'candidate_formula':candidate,
                                'additional_to_dev11':label not in prior_proposals,
                                'context_5x5':context,'business_truth_available':False})
                count += 1
        workbook.close()
        assert sha(path) == item['input_sha256']
        cases.append({'id':item['id'],'status':pred['status'],
                      'all_loaded_formulas_match_frozen_observation':True,
                      'candidate_records_independently_checked':count,
                      'array_formula_objects_compared_via_text':array_formula_count,
                      'original_input_unchanged':True})
    output = ROOT/'qa/BIDIRECTIONAL_DEV13_INDEPENDENT_AUDIT_20261018.json'
    assert not output.exists()
    result = {'prediction_lock_sha256':sha(RUN/'prediction_lock.json'),
              'audit_script_sha256':sha(Path(__file__)),
              'labelled_cases_without_metric_change':unchanged,
              'public_cases':cases,'public_candidate_details':details,
              'public_candidate_count':len(details),
              'scope':'openpyxl original-formula/provenance and token-difference check; no business labels',
              'main_candidate_adopted':False,'final_validation_eligible':False}
    output.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str)+'\n')
    print(json.dumps({'labelled_cases_without_metric_change':unchanged,
                      'public_candidates_checked':len(details)}))


if __name__ == '__main__':
    main()
