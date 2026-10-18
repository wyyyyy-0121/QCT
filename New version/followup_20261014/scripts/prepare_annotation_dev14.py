"""Author-only development inputs and independent AST oracle; never imported by prediction."""
from pathlib import Path
from fractions import Fraction
import ast
import re
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from run_selector_dev5 import read,save,sha
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import sympy

DATA = ROOT/'development/annotation_dev14_inputs_v1'
SOURCE = REPO/'data/external/hybridguard_candidate_audit_dev1/workbooks/uk_government/lcoe_calculator.xlsx'
SOURCE_SHA = 'b0f15382e8c9c9b74d944a0ab8301a3326fb3bfbf4a950896a66f29c52000357'
REGIONS = [('Offshore LCOE',27,15,44),('CCGT LCOE calc',27,13,37),('CCGT LCOE calc',81,13,37)]
SCENARIOS = ['clean_original','clean_input_scenario','isolated_operator','isolated_reference',
             'isolated_coefficient','common_operator_row','common_coefficient_row',
             'common_reference_row','multiple_regions','missing_note','hidden_note',
             'duplicate_caption','documented_exception','wrong_note_clean',
             'wrong_note_polluted_peers','wrong_note_with_error']
REF = re.compile(r'(?<![A-Za-z0-9_])\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)(?![A-Za-z0-9_])')

def expression(text, ref):
    """Independent Python AST, exact literals, no FormulaGuard parser/selector."""
    cleaned = REF.sub(lambda m:m[1].upper()+m[2],text.lstrip('='))
    tree = ast.parse(cleaned,mode='eval')
    def convert(n):
        if isinstance(n,ast.Name):return ref(n.id)
        if isinstance(n,ast.Constant) and type(n.value) in (int,float):
            literal=ast.get_source_segment(cleaned,n)
            f=Fraction(literal)
            return sympy.Rational(f.numerator,f.denominator)
        if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.UAdd,ast.USub)):
            return convert(n.operand)*(1 if isinstance(n.op,ast.UAdd) else -1)
        if isinstance(n,ast.BinOp):
            a,b=convert(n.left),convert(n.right)
            if isinstance(n.op,ast.Add):return a+b
            if isinstance(n.op,ast.Sub):return a-b
            if isinstance(n.op,ast.Mult):return a*b
            if isinstance(n.op,ast.Div):return a/b
        raise ValueError('unsupported_independent_oracle_grammar')
    return convert(tree.body)

def equivalent(a,b):
    try:return sympy.expand(expression(a,sympy.Symbol)-expression(b,sympy.Symbol))==0
    except (ValueError,SyntaxError,TypeError):return False

def evaluate(ws,text,stack=()):
    def ref(a):
        if a in stack:raise ValueError('oracle_cycle')
        value=ws[a].value
        if isinstance(value,str) and value.startswith('='):
            return evaluate(ws,value,stack+(a,))
        if type(value) not in (int,float):raise ValueError('oracle_nonnumeric:'+a)
        return sympy.Rational(str(value))
    return expression(text,ref)

def original_truth(w):
    expected={}; evidence=[]
    for sheet,row,start,end in REGIONS:
        ws=w[sheet]
        if row==27:
            assert ws['B27'].value=='Generation' and ws['C27'].value=='MWh'
            assert ws['E27'].value=='Net power output x Gross load factor x Availability x Hours per year'
            assert [(ws[f'B{r}'].value,ws[f'C{r}'].value) for r in [11,25,26,14]]==[
                ('Net power output','MW'),('Gross load factor','%'),('Availability','%'),('Hours per year','hours')]
        else:
            assert ws['B81'].value=='Carbon captured' and ws['C81'].value=='KgCO2'
            assert ws['E81'].value=='Carbon emissions x CO2 scrubbing'
            assert [(ws[f'B{r}'].value,ws[f'C{r}'].value) for r in [80,78]]==[
                ('Carbon emissions','KgCO2'),('CO2 scrubbing','%')]
        for col in range(start,end+1):
            c=get_column_letter(col); a=f'{c}{row}'
            product=f'=$D$11*{c}25*{c}26*$D$14' if row==27 else f'={c}78*{c}80'
            assert equivalent(ws[a].value,product)
            expected[sheet+'!'+a]=ws[a].value
        evidence.append({'sheet':sheet,'row':row,'columns':[start,end],
                         'note_cell':'E'+str(row),'note':ws['E'+str(row)].value,
                         'independent_relation':'power * load * availability * hours' if row==27 else 'emissions * scrubbing',
                         'scope':'local author-reviewed visible business relation, not external expert certification'})
    assert len(expected)==80
    return expected,evidence

def prepare():
    if DATA.exists():raise SystemExit('Preserve prepared inputs; use new version')
    assert sha(SOURCE)==SOURCE_SHA
    w=load_workbook(SOURCE); originals,evidence=original_truth(w); w.close()
    index=[]; truths=[]
    for number,scenario in enumerate(SCENARIOS,1):
        w=load_workbook(SOURCE); out=w['Offshore LCOE']; gas=w['CCGT LCOE calc']
        expected=dict(originals); changes=[]; metadata=[]
        def change(ws,a,value,kind):
            before=ws[a].value; ws[a]=value
            changes.append({'sheet':ws.title,'cell':a,'before':before,'after':value,'kind':kind})
        if scenario!='clean_original':
            for col in range(13,38):change(gas,f'{get_column_letter(col)}78',0.25,'numeric_scenario_input')
        def inject(ws,a,formula):change(ws,a,formula,'formula_error_injection')
        if scenario in ['isolated_operator','missing_note','hidden_note','duplicate_caption','wrong_note_with_error']:
            inject(out,'O27',out['O27'].value.replace('*','+',1))
        elif scenario=='isolated_reference':
            inject(out,'P27',out['P27'].value.replace('P25','P26'))
        elif scenario=='isolated_coefficient':
            inject(out,'Q27',out['Q27'].value+'*1.1')
        elif scenario=='common_operator_row':
            for col in range(15,45):
                a=f'{get_column_letter(col)}27'; inject(out,a,out[a].value.replace('*','+',1))
        elif scenario=='common_coefficient_row':
            for col in range(13,38):
                a=f'{get_column_letter(col)}27'; inject(gas,a,gas[a].value+'*1.1')
        elif scenario=='common_reference_row':
            for col in range(13,38):
                c=get_column_letter(col); inject(gas,f'{c}81',gas[f'{c}81'].value.replace(c+'78',c+'79'))
        elif scenario=='multiple_regions':
            inject(out,'O27',out['O27'].value.replace('*','+',1))
            inject(gas,'N27',gas['N27'].value+'*1.1')
            inject(gas,'O81',gas['O81'].value.replace('O78','O79'))
        if scenario=='missing_note':change(out,'E27',None,'note_removed')
        if scenario=='hidden_note':
            before=out.column_dimensions['E'].hidden; out.column_dimensions['E'].hidden=True
            metadata.append({'sheet':out.title,'column':'E','hidden_before':before,'hidden_after':True})
        if scenario=='duplicate_caption':
            for a,v in [('B200','Gross load factor'),('C200','%'),('D200',0.5)]:
                change(out,a,v,'duplicate_caption')
        if scenario=='documented_exception':
            change(out,'E27',out['E27'].value+' except O27 includes a 0.9 adjustment','legitimate_exception_note')
            new=out['O27'].value+'*0.9'
            change(out,'O27',new,'legitimate_formula_exception')
            expected[out.title+'!O27']=new
        if scenario.startswith('wrong_note'):
            change(out,'E27',out['E27'].value.replace('Gross load factor','Net efficiency (LHV)'),'intentionally_wrong_note')
        if scenario=='wrong_note_polluted_peers':
            for col in range(16,45):
                c=get_column_letter(col); a=c+'27'
                inject(out,a,out[a].value.replace(c+'25','$D$12'))
        case_id=f'annotation{number:02d}'
        path=DATA/'workbooks'/f'{case_id}.xlsx'; path.parent.mkdir(parents=True,exist_ok=True)
        if scenario=='clean_original':shutil.copyfile(SOURCE,path)
        else:w.save(path)
        observed={ws.title+'!'+cell.coordinate:cell.value for ws in w for row in ws for cell in row if cell.data_type=='f'}
        errors=[label for label,f in expected.items() if not equivalent(observed[label],f)]
        numeric={}
        for label in expected:
            s,a=label.rsplit('!',1); ws=w[s]
            good=evaluate(ws,expected[label],(a,)); actual=evaluate(ws,observed[label],(a,))
            assert label not in errors or good!=actual,(case_id,label)
            numeric[label]={'expected':str(good),'observed':str(actual)}
        truths.append({'id':case_id,'scenario':scenario,'expected_formulas':expected,'errors':errors,
                       'numeric_oracle':numeric,'observed_formulas':observed,
                       'changes':changes,'metadata_changes':metadata,'known_scope_formula_count':80,
                       'unknown_scope_formula_count':len(observed)-80})
        index.append({'id':case_id,'input':str(path.relative_to(ROOT)),'input_sha256':sha(path),
                      'group':'public_development_annotation','parallel_sheets':[],
                      'final_validation_eligible':False,'origin':'uk_government'})
        w.close()
        print(case_id,scenario,'errors',len(errors),flush=True)
    save(DATA/'public_index.json',index)
    save(DATA/'truth.json',{'source_sha256':SOURCE_SHA,'cases':truths,'scope_evidence':evidence,
                           'independent_oracle':'Python AST + exact SymPy rational arithmetic',
                           'final_validation_eligible':False,'original_source_unchanged':sha(SOURCE)==SOURCE_SHA})
    save(DATA/'preparation_lock.json',{'files':{str(p.relative_to(DATA)):sha(p) for p in sorted(DATA.rglob('*')) if p.is_file()},
                                     'author_script_sha256':sha(Path(__file__))})
if __name__=='__main__':prepare()
