"""Author-only exact numeric relation certificates; never a predictor oracle.

This is a bounded polynomial abstraction, not full Excel/business semantics.
References are expanded so a counterexample respects supported dependencies.
"""
from collections import Counter
from fractions import Fraction
import hashlib
import json
import re
from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.utils.cell import range_boundaries, get_column_letter

BUDGETS = dict(formula_chars=4096, tokens=512, range_cells=256, depth=32,
               terms=128, degree=8, literal_chars=64, visits=4096, rational_bits=4096)
SEEDS = (1, 3, 7, 13)
REF = re.compile(r'\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)(?::\$?([A-Za-z]{1,3})\$?([1-9][0-9]*))?\Z')


class Unsupported(ValueError):
    pass


def checked(poly):
    poly = {monomial: coefficient for monomial, coefficient in poly.items() if coefficient}
    if len(poly) > BUDGETS['terms']: raise Unsupported('term_budget')
    if any(len(m) > BUDGETS['degree'] for m in poly): raise Unsupported('degree_budget')
    if any(max(abs(c.numerator).bit_length(), c.denominator.bit_length()) > BUDGETS['rational_bits'] for c in poly.values()):
        raise Unsupported('rational_budget')
    return poly


def add(a, b, sign=1):
    result = dict(a)
    for key, value in b.items(): result[key] = result.get(key, Fraction(0))+sign*value
    return checked(result)


def multiply(a, b):
    result = {}
    for left, x in a.items():
        for right, y in b.items():
            key = tuple(sorted(left+right)); result[key] = result.get(key, Fraction(0))+x*y
    return checked(result)


def constant(number):
    return {} if number == 0 else {(): Fraction(number)}


def canonical(poly):
    return [[list(m), str(c)] for m, c in sorted(poly.items())]


def digest(poly):
    return hashlib.sha256(json.dumps(canonical(poly), separators=(',', ':')).encode()).hexdigest()


def evaluate(poly, assignments):
    result = Fraction(0)
    for monomial, coefficient in poly.items():
        value = coefficient
        for name in monomial: value *= assignments[name]
        result += value
    return result


class NumericRelations:
    def __init__(self, workbook, expected=None):
        self.workbook = workbook
        self.sheets = {s.casefold(): s for s in workbook.sheetnames}
        if len(self.sheets) != len(workbook.sheetnames): raise Unsupported('ambiguous_sheet')
        self.expected = expected or {}
        self.cache = {}
        self.compiled = {}
        self.counts = Counter()

    def location(self, sheet, address):
        if sheet.casefold() not in self.sheets: raise Unsupported('missing_sheet')
        sheet = self.sheets[sheet.casefold()]
        match = REF.fullmatch(address)
        if not match or match[3]: raise Unsupported('unsupported_ref')
        address = match[1].upper()+match[2]
        bounds = range_boundaries(address)
        if bounds[0] > 16384 or bounds[1] > 1048576: raise Unsupported('reference_bounds')
        return sheet, address

    def reference(self, value, sheet, stack, depth, visits):
        if '!' in value:
            prefix, value = value.rsplit('!', 1)
            if prefix.startswith("'") and prefix.endswith("'"): prefix = prefix[1:-1].replace("''", "'")
            sheet = prefix
        if '[' in value or '[' in sheet or ':' in sheet: raise Unsupported('external_or_3d_ref')
        match = REF.fullmatch(value)
        if not match: raise Unsupported('unsupported_ref')
        ref = value.replace('$', '')
        left, top, right, bottom = range_boundaries(ref)
        if left > right or top > bottom or right > 16384 or bottom > 1048576: raise Unsupported('reference_bounds')
        if (right-left+1)*(bottom-top+1) > BUDGETS['range_cells']: raise Unsupported('range_budget')
        values = [self.cell(sheet, f'{get_column_letter(col)}{row}', stack, depth+1, visits)
                  for row in range(top, bottom+1) for col in range(left, right+1)]
        return ('range', values) if ':' in value else values[0]

    def cell(self, sheet, address, stack, depth, visits):
        sheet, address = self.location(sheet, address); key = sheet+'!'+address
        visits[0] += 1
        if visits[0] > BUDGETS['visits']: raise Unsupported('visit_budget')
        if depth > BUDGETS['depth']: raise Unsupported('depth_budget')
        if key in stack: raise Unsupported('cycle')
        # Cache successes only when their depth/visits cannot bypass budgets.
        # Pure numeric leaves are constant-cost and safe to cache.
        if key in self.cache: return self.cache[key]
        cell = self.workbook[sheet][address]
        value = self.expected.get(key, cell.value)
        if key in self.expected or (isinstance(value, str) and value.startswith('=')):
            return self.formula(value, sheet, stack+(key,), depth+1, visits)
        if cell.data_type == 'e': raise Unsupported('error_input')
        if value is None: result = ('blank', None)
        elif isinstance(value, bool): raise Unsupported('boolean_input')
        elif type(value) in (int, float):
            # Every numeric leaf is a symbolic input, not a cached formula value.
            if str(value).lower() in ['nan', 'inf', '-inf']: raise Unsupported('nonfinite_input')
            result = ('poly', {(key,): Fraction(1)})
        elif isinstance(value, str): result = ('text', value)
        else: raise Unsupported('nonnumeric_input')
        self.cache[key] = result
        return result

    def formula(self, formula, sheet, stack=(), depth=0, visits=None):
        if visits is None: visits = [0]
        if depth > BUDGETS['depth']: raise Unsupported('depth_budget')
        if not isinstance(formula, str) or not formula.startswith('='): raise Unsupported('not_formula')
        if len(formula) > BUDGETS['formula_chars']: raise Unsupported('formula_budget')
        try: tokens = [t for t in Tokenizer(formula).items if t.type != 'WHITE-SPACE']
        except (TokenizerError, IndexError): raise Unsupported('tokenizer')
        if len(tokens) > BUDGETS['tokens']: raise Unsupported('token_budget')
        position = [0]

        def scalar(value):
            if value[0] == 'poly': return value[1]
            if value[0] == 'blank': return {}
            raise Unsupported('nonscalar_numeric_operand')

        def peek(): return tokens[position[0]] if position[0] < len(tokens) else None

        def take():
            token = peek()
            if token is None: raise Unsupported('incomplete_formula')
            position[0] += 1; return token

        def primary():
            token = take()
            if token.type == 'OPERATOR-PREFIX' and token.value in ['+', '-']:
                value = scalar(primary())
                return ('poly', value if token.value == '+' else {m: -c for m, c in value.items()})
            if token.type == 'OPERAND':
                if token.subtype == 'NUMBER':
                    if len(token.value) > BUDGETS['literal_chars']: raise Unsupported('literal_budget')
                    if 'e' in token.value.casefold():
                        exponent = token.value.casefold().split('e', 1)[1]
                        if len(exponent.lstrip('+-')) > 4 or abs(int(exponent)) > 1024: raise Unsupported('rational_budget')
                    try: return ('poly', checked(constant(Fraction(token.value))))
                    except (ValueError, OverflowError, ZeroDivisionError): raise Unsupported('numeric_literal')
                if token.subtype == 'RANGE': return self.reference(token.value, sheet, stack, depth, visits)
                raise Unsupported('unsupported_operand')
            if token.type == 'PAREN' and token.subtype == 'OPEN':
                value = expression(); closing = take()
                if closing.type != 'PAREN' or closing.subtype != 'CLOSE': raise Unsupported('parentheses')
                return ('poly', scalar(value))
            if token.type == 'FUNC' and token.subtype == 'OPEN' and token.value.casefold() == 'sum(':
                result = {}; first = peek()
                if first and first.type == 'FUNC' and first.subtype == 'CLOSE': take(); return ('poly', result)
                while True:
                    value = expression()
                    arguments = value[1] if value[0] == 'range' else [value]
                    for argument in arguments:
                        if argument[0] in ['text', 'blank']: continue
                        result = add(result, scalar(argument))
                    separator = take()
                    if separator.type == 'FUNC' and separator.subtype == 'CLOSE': break
                    if separator.type != 'SEP' or separator.subtype != 'ARG': raise Unsupported('sum_separator')
                return ('poly', result)
            raise Unsupported('unsupported_grammar')

        def product():
            left = primary()
            while peek() and peek().type == 'OPERATOR-INFIX' and peek().value in ['*', '/']:
                op = take().value; right = scalar(primary()); left_poly = scalar(left)
                if op == '*': left = ('poly', multiply(left_poly, right))
                else:
                    if set(right) != {()} or right[()] == 0: raise Unsupported('nonconstant_or_zero_divisor')
                    left = ('poly', checked({m: c/right[()] for m, c in left_poly.items()}))
            return left

        def expression():
            left = product()
            while peek() and peek().type == 'OPERATOR-INFIX' and peek().value in ['+', '-']:
                op = take().value; right = product()
                left = ('poly', add(scalar(left), scalar(right), 1 if op == '+' else -1))
            return left

        result = expression()
        if position[0] != len(tokens): raise Unsupported('unconsumed_tokens')
        return ('poly', checked(scalar(result)))

    def compile(self, formula, sheet, target=None):
        key = (formula, sheet, target)
        if key in self.compiled:
            result = self.compiled[key]
            if isinstance(result, str): raise Unsupported(result)
            return result
        stack = (target,) if target else ()
        try: result = self.formula(formula, sheet, stack)[1]
        except Unsupported as error:
            result = str(error)
        if len(self.compiled) >= 256: self.compiled.clear()
        self.compiled[key] = result
        if isinstance(result, str): raise Unsupported(result)
        return result

    def certify(self, candidate, expected, sheet, target=None, in_scope=True):
        if not in_scope: return dict(label=None, reason='outside_reviewed_scope')
        try:
            left = self.compile(candidate, sheet, target)
            right = self.compile(expected, sheet, target)
            left_sha, right_sha = digest(left), digest(right)
            if left == right:
                return dict(label=1, reason='identical_exact_polynomial', candidate_sha256=left_sha, expected_sha256=right_sha)
            names = sorted({name for monomial in list(left)+list(right) for name in monomial})
            for seed in SEEDS:
                assignment = {name: Fraction(31+seed*n+seed, 7+n) for n, name in enumerate(names)}
                a, b = evaluate(left, assignment), evaluate(right, assignment)
                if a != b:
                    return dict(label=0, reason='exact_numeric_relation_counterexample', candidate_sha256=left_sha,
                        expected_sha256=right_sha, witness=dict(seed=seed, inputs={name: str(value) for name, value in assignment.items()},
                        candidate=str(a), expected=str(b)), scope='bounded numeric dependency abstraction, not actual Excel values')
            return dict(label=None, reason='no_counterexample_in_registered_assignments')
        except (Unsupported, ValueError, KeyError, TypeError, IndexError, ZeroDivisionError, RecursionError) as error:
            return dict(label=None, reason=str(error) if isinstance(error, Unsupported) else 'unsupported_dependency')
