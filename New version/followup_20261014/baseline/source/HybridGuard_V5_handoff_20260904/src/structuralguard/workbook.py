"""Read-only workbook snapshot with explicit protected-sheet exclusion."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Sequence

from openpyxl import load_workbook

CellKey = tuple[str, str]
PROTECTED_SHEET_RE = re.compile(r"answer|label|ground.?truth|secret|post.?run", re.I)


@dataclass(frozen=True)
class WorkbookSnapshot:
    path: Path
    cells: dict[CellKey, object]
    formulas: dict[CellKey, str]
    visible_sheets: frozenset[str]
    sheet_order: tuple[str, ...]

    @classmethod
    def load(
        cls,
        path: str | Path,
        *,
        exclude_sheets: Sequence[str] = (),
        progress: Callable[[str, int, int], None] | None = None,
    ) -> "WorkbookSnapshot":
        source = Path(path).resolve()
        if source.suffix.lower() != ".xlsx":
            raise ValueError("StructuralGuard accepts .xlsx only; export macro workbooks safely first")
        if not source.is_file():
            raise FileNotFoundError(source)
        excluded = {name.casefold(): name for name in exclude_sheets}
        workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        available = {sheet.title.casefold(): sheet.title for sheet in workbook.worksheets}
        missing = [name for key, name in excluded.items() if key not in available]
        if missing:
            workbook.close()
            raise ValueError("excluded sheet(s) not found: " + ", ".join(sorted(missing)))
        protected = [
            sheet.title for sheet in workbook.worksheets
            if PROTECTED_SHEET_RE.search(sheet.title) and sheet.title.casefold() not in excluded
        ]
        if protected:
            workbook.close()
            raise ValueError(
                "potential answer/label sheet must be explicitly excluded: "
                + ", ".join(protected)
            )
        included = [sheet for sheet in workbook.worksheets if sheet.title.casefold() not in excluded]
        cells: dict[CellKey, object] = {}
        formulas: dict[CellKey, str] = {}
        visible: set[str] = set()
        order: list[str] = []
        for index, sheet in enumerate(included, 1):
            order.append(sheet.title)
            if sheet.sheet_state == "visible":
                visible.add(sheet.title)
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    key = (sheet.title, cell.coordinate.upper())
                    cells[key] = cell.value
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas[key] = cell.value
            if progress:
                progress("read_workbook", index, len(included))
        workbook.close()
        return cls(source, cells, formulas, frozenset(visible), tuple(order))

    @property
    def visible_formula_keys(self) -> tuple[CellKey, ...]:
        return tuple(sorted(key for key in self.formulas if key[0] in self.visible_sheets))
