#!/usr/bin/env python3
"""Bounded static-label mapping without target formula or numeric content."""
import argparse
import json
import os
from pathlib import Path
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

from spreadsheetbench2_scoring import sha256


def snapshot(path):
    book = load_workbook(path, read_only=False, data_only=False)
    result = {'sheets': {}}
    try:
        for sheet in book:
            cells = {}
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'f':
                        cells[cell.coordinate] = {'kind': 'formula'}
                    elif cell.value is not None:
                        if cell.data_type == 's' and isinstance(cell.value, str):
                            cells[cell.coordinate] = {'kind': 'text', 'text': cell.value}
                        elif cell.data_type == 'n':
                            cells[cell.coordinate] = {'kind': 'number'}
                        else:
                            cells[cell.coordinate] = {'kind': 'other'}
            result['sheets'][sheet.title] = cells
    finally:
        book.close()
    return result


def norm(text):
    return ' '.join(text.split()).rstrip('= ').strip()


def address_parts(address):
    column, row = coordinate_from_string(address)
    return row, column_index_from_string(column)


def label_match(label, term):
    if label.lstrip().startswith(('=', '≈', '*')) or len(label) > 160:
        return False
    clean, query = norm(label), norm(term)
    # A one-token scientific symbol has case-sensitive meaning.
    if re.fullmatch(r'\w+', query):
        return clean == query or query in [norm(part) for part in re.findall(r'\(([^()]+)\)', label)]
    return clean.casefold() == query.casefold()


def map_field(data, query):
    if set(query) - {'term', 'unit', 'sheet'} or not isinstance(query.get('term'), str) or not query['term'].strip():
        return dict(status='unresolved', reason='invalid_query', candidates=[])
    candidates = []
    for sheet, cells in data['sheets'].items():
        if query.get('sheet') and sheet != query['sheet']:
            continue
        headers = []
        for address, cell in cells.items():
            if cell['kind'] == 'text' and norm(cell['text']).casefold() == 'parameter':
                row, column = address_parts(address)
                values, units = [], []
                for other, content in cells.items():
                    r, c = address_parts(other)
                    if r == row and c > column and content['kind'] == 'text':
                        if norm(content['text']).casefold() == 'calculated value':
                            values.append(c)
                        elif norm(content['text']).casefold() == 'units':
                            units.append(c)
                if values and units:
                    headers.append((row, column, values, units))
        for label_address, content in cells.items():
            if content['kind'] != 'text' or not label_match(content['text'], query['term']):
                continue
            row, column = address_parts(label_address)
            value_columns = [column + 1]
            applicable = [header for header in headers if header[0] < row and header[1] == column]
            if applicable:
                value_columns = max(applicable, key=lambda header: header[0])[2]
            for value_column in value_columns:
                address = get_column_letter(value_column) + str(row)
                if cells.get(address, {}).get('kind') not in {'number', 'formula'}:
                    continue
                unit_address = get_column_letter(value_column + 1) + str(row)
                unit_cell = cells.get(unit_address, {})
                unit = norm(unit_cell['text']) if unit_cell.get('kind') == 'text' else None
                candidates.append(dict(sheet=sheet, cell=address, label_cell=label_address,
                                       label=content['text'], unit=unit, unit_cell=unit_address,
                                       unit_origin='static' if unit is not None else 'unknown',
                                       layout='header' if applicable else 'right_adjacent'))
    unique = {(row['sheet'], row['cell']): row for row in candidates}
    candidates = [unique[key] for key in sorted(unique)]
    if not candidates:
        return dict(status='unresolved', reason='no_structural_field_match', candidates=[])
    if query.get('unit') is not None:
        # Unknown-unit alternatives cannot be silently filtered out.
        if any(row['unit'] is None for row in candidates):
            return dict(status='unresolved', reason='unit_not_statically_verified', candidates=candidates)
        matching = [row for row in candidates if row['unit'] == norm(query['unit'])]
        if not matching:
            return dict(status='unresolved', reason='unit_mismatch', candidates=candidates)
        candidates = matching
    if len(candidates) != 1:
        return dict(status='unresolved', reason='ambiguous_fields', candidates=candidates)
    return dict(status='mapped', reason='unique_structural_match', candidates=candidates,
                reference=dict(sheet=candidates[0]['sheet'], cell=candidates[0]['cell']))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('request', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    request = json.loads(args.request.read_text())
    source = Path(request['workbook']).resolve()
    metadata_path = Path(request['metadata']).resolve()
    output = args.output.resolve()
    if output.exists() or output in (source, metadata_path, args.request.resolve()):
        raise SystemExit('Output must be new and distinct')
    if sha256(source) != request['workbook_sha256'] or sha256(metadata_path) != request['metadata_sha256']:
        raise SystemExit('Input changed')

    def guard(event, values):
        if event != 'open' or not values or not isinstance(values[0], (str, bytes, os.PathLike)):
            return
        path = Path(os.fsdecode(values[0])).resolve()
        if 'scoring' in path.parts or any(name in path.name.lower() for name in ('expected.json', 'golden', 'after.xlsx')):
            raise PermissionError('Scoring evidence denied')
        if path.suffix.lower() in {'.xlsx', '.xls', '.xlsm', '.ods'}:
            if request['mode'] == 'map' or path != source:
                raise PermissionError('Mapping only reads formula-free metadata')
            mode, flags = values[1:3]
            if (isinstance(mode, str) and any(c in mode for c in 'wa+')) or (
                    isinstance(flags, int) and flags & (os.O_WRONLY | os.O_RDWR | os.O_TRUNC | os.O_CREAT)):
                raise PermissionError('Original is read-only')

    sys.addaudithook(guard)
    data = json.loads(metadata_path.read_text())
    result = dict(status='abstain', suggested_formula=None, field_mapping={}, formula_content_used_for_mapping=False,
                  final_validation_eligible=False, automatic_field_query_generation=False, labels_read=[])
    if any(key in request for key in ('references', 'relations', 'constraints', 'candidates')):
        result.update(status='unsupported', reason='manual_address_or_relation_forbidden')
    elif request['mode'] == 'map':
        result.update(map_field(data, request['query']))
    else:
        result['field_mapping'] = {symbol: map_field(data, query) for symbol, query in request['fields'].items()}
        if any(row['status'] != 'mapped' for row in result['field_mapping'].values()):
            result['reason'] = 'field_mapping_incomplete'
        else:
            # Only after metadata mapping completes may paper inference read formulas.
            from source_relations_m7 import extract, load_source
            from relation_candidates_m6 import review
            from trace_paper_candidates import predict
            from e3_contract import Unsupported
            try:
                references = {}
                for symbol, row in result['field_mapping'].items():
                    reference = row['reference']
                    key = reference['cell'] if reference['sheet'] == request['sheet'] else (
                        "'" + reference['sheet'].replace("'", "''") + "'!" + reference['cell'])
                    if key in references:
                        raise Unsupported('different_fields_share_one_cell')
                    references[key] = symbol
                book = load_workbook(source, read_only=True, data_only=False)
                try:
                    relation = extract(load_source(request, book), request['aliases'], request['symbols'])
                finally:
                    book.close()
                paper = predict(source)
                target = request['sheet'] + '!' + request['cell']
                observed = paper['observed_formulas'][target]
                native = {row['formula'] for row in paper['candidate_trace']['fg'].get(target, [])}
                native.update(row['formula'] for row in paper['candidate_trace']['sg'].get(target, {}).get('candidates', []))
                paper_row = next(row for row in paper['paper']['rows'] if row['cell'] == target)
                if paper_row.get('suggested_formula'):
                    native.add(paper_row['suggested_formula'])
                compiled = dict(request, relations=[relation], references=references)
                result.update(review(compiled, observed, sorted(native)))
                result.update(automatically_mapped_references=references, observed_formula=observed)
            except (Unsupported, ValueError, KeyError, TypeError, SyntaxError) as error:
                result.update(status='unsupported', reason=str(error), suggested_formula=None)
    if sha256(metadata_path) != request['metadata_sha256']:
        raise SystemExit('Metadata changed')
    result.update(request_sha256=sha256(args.request), metadata_sha256=sha256(metadata_path),
                  workbook_sha256=request['workbook_sha256'], mapping_semantics='Structural candidate under declared layout; not business certification')
    with output.open('x') as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2, sort_keys=True)
        handle.write('\n')


if __name__ == '__main__':
    main()
