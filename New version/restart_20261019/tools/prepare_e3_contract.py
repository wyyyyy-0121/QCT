#!/usr/bin/env python3
"""Create and lock the preregistered E3 development inputs and score key."""
import hashlib
import json
from pathlib import Path
import platform
import sys

from openpyxl import Workbook, load_workbook
import openpyxl
import sympy

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / 'source_evidence'
EPA = EVIDENCE / 'raw/epa_pairs/xlsx_export/oxidizers_before.xlsx'
OUT = ROOT / 'e3_contract_v2'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def dump(path, value):
    Path(path).write_text(json.dumps(value, sort_keys=True, indent=2) + '\n')


def main():
    if OUT.exists():
        allowed_partial = {'manifest.jsonl', 'workbooks/counterfactual.xlsx',
                           'scoring/expected.json'}
        allowed_partial.update(f'requests/{case}.json' for case in (
            'epa_none', 'epa_ratio', 'counter_none', 'counter_units', 'counter_ratio',
            'counter_affine', 'counter_conflict', 'counter_missing'))
        existing = {path.relative_to(OUT).as_posix() for path in OUT.rglob('*') if path.is_file()}
        if existing - allowed_partial or (OUT / 'input_lock.json').exists() or (OUT / 'score.json').exists():
            raise SystemExit(f'Preserving existing E3 evidence: {OUT}')
    if not EPA.is_file():
        raise SystemExit(f'Missing registered EPA development input: {EPA}')
    OUT.mkdir(exist_ok=True)
    workbook_dir = OUT / 'workbooks'
    request_dir = OUT / 'requests'
    expected_dir = OUT / 'scoring'
    for directory in (workbook_dir, request_dir, expected_dir):
        directory.mkdir(exist_ok=True)

    micro = workbook_dir / 'counterfactual.xlsx'
    if not micro.exists():
        book = Workbook()
        sheet = book.active
        sheet.title = 'Model'
        sheet['B2'] = 10
        sheet['C2'] = 1
        sheet['D2'] = '=B2*(2-C2)'
        book.save(micro)
        book.close()
    else:
        book = load_workbook(micro, data_only=False, read_only=True)
        try:
            sheet = book['Model']
            if (sheet['B2'].value, sheet['C2'].value, sheet['D2'].value) != (10, 1, '=B2*(2-C2)'):
                raise SystemExit('Existing partial counterfactual does not match preregistration')
        finally:
            book.close()

    epa_sheet = 'Design Parameters'
    epa_cell = 'D36'
    book = load_workbook(EPA, data_only=False, read_only=True)
    try:
        observed = book[epa_sheet][epa_cell].value
    finally:
        book.close()
    if not isinstance(observed, str) or not observed.startswith('=IF('):
        raise SystemExit('EPA D36 no longer matches the preregistered IF case')
    guard = 'D35="Not applicable"'
    inactive = "'Data Inputs'!B23"
    corrected = f'=IF({guard},{inactive},({inactive}/D34))'
    original = observed
    if "'Data Inputs'!B23" not in observed or 'D34' not in observed:
        raise SystemExit('EPA D36 formula differs from registered symbols')

    symbols = {'q': {'units': {'flow': 1}}, 'd': {'units': {}}}
    cases = [
        ('epa_none', EPA, epa_sheet, epa_cell,
         {"'Data Inputs'!B23": 'q', 'D34': 'd'}, symbols, [], [original, corrected],
         {'guard': guard, 'inactive': inactive}, 'abstain', None, 'missing_business_relation'),
        ('epa_ratio', EPA, epa_sheet, epa_cell,
         {"'Data Inputs'!B23": 'q', 'D34': 'd'}, symbols, ['q/d'], [original, corrected],
         {'guard': guard, 'inactive': inactive}, 'propose', corrected,
         'candidate_satisfies_explicit_relation'),
        ('counter_none', micro, 'Model', 'D2', {'B2': 'q', 'C2': 'd'}, symbols, [],
         ['=B2/C2', '=B2*(2-C2)'], None, 'abstain', None, 'missing_business_relation'),
        ('counter_units', micro, 'Model', 'D2', {'B2': 'q', 'C2': 'd'}, symbols, [],
         ['=B2/C2', '=B2*(2-C2)'], None, 'abstain', None, 'missing_business_relation'),
        ('counter_ratio', micro, 'Model', 'D2', {'B2': 'q', 'C2': 'd'}, symbols, ['q/d'],
         ['=B2/C2', '=B2*(2-C2)'], None, 'propose', '=B2/C2',
         'candidate_satisfies_explicit_relation'),
        ('counter_affine', micro, 'Model', 'D2', {'B2': 'q', 'C2': 'd'}, symbols,
         ['q*(2-d)'], ['=B2/C2', '=B2*(2-C2)'], None,
         'propose', '=B2*(2-C2)', 'candidate_satisfies_explicit_relation'),
        ('counter_conflict', micro, 'Model', 'D2', {'B2': 'q', 'C2': 'd'}, symbols,
         ['q/d', 'q*(2-d)'], ['=B2/C2', '=B2*(2-C2)'], None,
         'abstain', None, 'conflicting_business_relations'),
        ('counter_missing', micro, 'Model', 'D2', {'B2': 'q', 'C2': 'd'}, symbols,
         ['q/d'], ['=B2*(2-C2)'], None,
         'abstain', None, 'no_candidate_satisfies_relation'),
    ]
    expected = {}
    manifest = []
    input_paths = [EPA, micro]
    for (case_id, workbook, sheet, cell, refs, symbol_defs, constraints,
         candidates, branch, status, suggestion, reason) in cases:
        request_path = request_dir / f'{case_id}.json'
        request = {
            'workbook': str(workbook.resolve()), 'workbook_sha256': sha(workbook),
            'sheet': sheet, 'cell': cell, 'references': refs, 'symbols': symbol_defs,
            'constraints': constraints, 'candidates': candidates,
        }
        if branch is not None:
            request['branch'] = branch
        dump(request_path, request)
        manifest.append({
            'id': case_id,
            'argv': [sys.executable, str((ROOT / 'tools/e3_contract.py').resolve()),
                     str(request_path.resolve()), '--output', 'prediction.json'],
            'outputs': ['prediction.json'],
        })
        expected[case_id] = {'status': status, 'suggested_formula': suggestion, 'reason': reason}
        input_paths.append(request_path)

    manifest_path = OUT / 'manifest.jsonl'
    manifest_path.write_text(''.join(json.dumps(row, sort_keys=True) + '\n' for row in manifest))
    expected_path = expected_dir / 'expected.json'
    dump(expected_path, expected)
    lock = {
        'schema': 'hybridguard-e3-input-lock-v1',
        'created_utc': __import__('datetime').datetime.now(__import__('datetime').timezone.utc).isoformat(),
        'source_workbook_sha256': sha(EPA), 'counterfactual_workbook_sha256': sha(micro),
        'workbooks_identical_across_counterfactual_conditions': True,
        'files': {('external:epa_oxidizers_before.xlsx' if path == EPA else
                   str(path.relative_to(OUT))): sha(path) for path in sorted(input_paths)},
        'manifest_sha256': sha(manifest_path), 'expected_sha256': sha(expected_path),
        'implementation_sha256': sha(ROOT / 'tools/e3_contract.py'),
        'runner_sha256': sha(ROOT / 'tools/batch_runner.py'),
        'input_builder_sha256': sha(__file__),
        'preregistration_sha256': sha(ROOT / 'E3_CONTRACT_PREREGISTRATION.md'),
        'environment': {'python': platform.python_version(), 'openpyxl': openpyxl.__version__,
                        'sympy': sympy.__version__},
        'prediction_reads_expected_file': False,
        'purpose': 'Development-only test of explicit rational-contract candidate selection; not natural accuracy.',
    }
    dump(OUT / 'input_lock.json', lock)
    print(json.dumps({'cases': len(cases), 'input_lock': str(OUT / 'input_lock.json'),
                      'manifest_sha256': lock['manifest_sha256'],
                      'counterfactual_workbook_sha256': lock['counterfactual_workbook_sha256']}, indent=2))


if __name__ == '__main__':
    main()
