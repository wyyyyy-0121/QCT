#!/usr/bin/env python3
"""Audit locked E1 predictions on original inputs; publish no workbook formulas."""
import argparse
from collections import Counter
from fractions import Fraction
import hashlib
import json
from pathlib import Path
import re

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries, get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'source_evidence/metrics_pair'
RAW = ROOT / 'source_evidence/raw/metrics_pair'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_json(path):
    return json.loads(path.read_text())


class Unsupported(ValueError):
    pass


class LinearAudit:
    """Independent limited interpreter for numeric cells, references, SUM and +.

    Leaves remain symbolic for dependency comparison. Values are evaluated only
    from the original workbook. Other Excel operations fail explicitly.
    """

    def __init__(self, workbook):
        self.workbook = workbook
        self.cache = {}

    def cell(self, sheet, address, visiting=()):
        key = (sheet, address.replace('$', '').upper())
        if key in visiting:
            raise Unsupported('cyclic dependency')
        if key in self.cache:
            return self.cache[key].copy()
        value = self.workbook[sheet][key[1]].value
        if isinstance(value, str) and value.startswith('='):
            result = self.formula(sheet, value, visiting + (key,))
        elif value is None or (isinstance(value, (int, float)) and not isinstance(value, bool)):
            result = Counter({key: 1})
        else:
            raise Unsupported('non-numeric reference')
        self.cache[key] = result.copy()
        return result

    def reference(self, sheet, value, visiting):
        if '!' in value:
            sheet, value = value.rsplit('!', 1)
            sheet = sheet.strip("'").replace("''", "'")
        if sheet not in self.workbook.sheetnames:
            raise Unsupported('external or missing sheet')
        value = value.replace('$', '')
        if not re.fullmatch(r'[A-Za-z]+[1-9][0-9]*(?::[A-Za-z]+[1-9][0-9]*)?', value):
            raise Unsupported('unsupported reference')
        lo_c, lo_r, hi_c, hi_r = range_boundaries(value)
        if (hi_c - lo_c + 1) * (hi_r - lo_r + 1) > 10000:
            raise Unsupported('reference too large')
        result = Counter()
        for r in range(lo_r, hi_r + 1):
            for c in range(lo_c, hi_c + 1):
                result.update(self.cell(sheet, f'{get_column_letter(c)}{r}', visiting))
        return result

    def formula(self, sheet, formula, visiting=()):
        tokens = [t for t in Tokenizer(formula).items if t.type != 'WHITE-SPACE']
        index = 0

        def expression():
            nonlocal index
            result = term()
            while index < len(tokens) and tokens[index].value == '+':
                index += 1
                result.update(term())
            return result

        def term():
            nonlocal index
            if index >= len(tokens):
                raise Unsupported('missing operand')
            token = tokens[index]
            index += 1
            if token.type == 'OPERAND' and token.subtype == 'RANGE':
                return self.reference(sheet, token.value, visiting)
            if token.type == 'OPERAND' and token.subtype == 'NUMBER':
                return Counter({(None, str(Fraction(token.value))): 1})
            if token.type == 'FUNC' and token.subtype == 'OPEN' and token.value.upper() == 'SUM(':
                result = expression()
                if index >= len(tokens) or tokens[index].type != 'FUNC' or tokens[index].subtype != 'CLOSE':
                    raise Unsupported('unsupported SUM arguments')
                index += 1
                return result
            raise Unsupported('unsupported token: ' + token.type + '/' + token.subtype)

        result = expression()
        if index != len(tokens):
            raise Unsupported('unsupported trailing operation')
        return result

    def value(self, coefficients):
        return sum(((Fraction(a) if s is None else Fraction(str(self.workbook[s][a].value or 0))) * count
                    for (s, a), count in coefficients.items()), Fraction(0))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=OUT / 'E1_BASELINE_AUDIT_V2.json')
    args = parser.parse_args()
    lock = read_json(OUT / 'prediction_lock.json')
    for relative, expected in lock['files'].items():
        assert sha(ROOT / relative) == expected, relative
    for stage in ('before', 'after'):
        receipt = read_json(OUT / f'{stage}_acquisition.json')
        assert sha(RAW / f'{stage}.xlsx') == receipt['sha256']
    prediction_path = RAW / 'original_run/tasks/metrics/prediction.json'
    prediction = read_json(prediction_path)
    assert prediction['labels_read'] == [] and prediction['original_unchanged']
    assert prediction['input_sha256'] == sha(RAW / 'before.xlsx')
    before, after = [openpyxl.load_workbook(RAW / f'{stage}.xlsx', data_only=False)
                     for stage in ('before', 'after')]
    assert before.sheetnames == after.sheetnames
    audit = LinearAudit(before)
    model_rows = {m: {r['cell']: r for r in prediction[m]['rows']} for m in ('paper', 'dev11')}
    records, changes = [], Counter()
    for sn in before.sheetnames:
        old_sheet, new_sheet = before[sn], after[sn]
        for r in range(1, max(old_sheet.max_row, new_sheet.max_row) + 1):
            for c in range(1, max(old_sheet.max_column, new_sheet.max_column) + 1):
                old, new = old_sheet.cell(r, c).value, new_sheet.cell(r, c).value
                if old == new:
                    continue
                both_formulas = all(isinstance(v, str) and v.startswith('=') for v in (old, new))
                changes['formula_changed' if both_formulas else 'nonformula_changed'] += 1
                if not both_formulas:
                    continue
                col = get_column_letter(c)
                assert r == 16 and old == f'=SUM({col}11:{col}14)' and new == f'=SUM({col}11:{col}15)'
                assert str(old_sheet['A16'].value).casefold() == 'team totals'
                key = f'{sn}!{col}{r}'
                present = old_sheet.cell(15, c).value not in (None, '')
                record = dict(cell=key, row15_present=present,
                              label='formula_defect_candidate' if present else 'blank_boundary_control')
                try:
                    old_terms, corrected_terms = [audit.formula(sn, f) for f in (old, new)]
                    old_value, corrected_value = [audit.value(t) for t in (old_terms, corrected_terms)]
                    record.update(execution='supported', original_value=str(old_value),
                                  corrected_value_on_original_inputs=str(corrected_value),
                                  realized_output_error=old_value != corrected_value)
                except Unsupported as error:
                    record.update(execution='unsupported', reason=str(error), realized_output_error=None)
                    old_terms = corrected_terms = None
                for model in ('paper', 'dev11'):
                    row = model_rows[model].get(key, {})
                    proposed = row.get('suggested_formula')
                    changed = bool(proposed) and proposed != prediction['observed_formulas'][key]
                    result = dict(changed_suggestion=changed, status=row.get('status', 'no_row'),
                                  exact_revision_match=bool(changed and proposed == new))
                    if changed:
                        try:
                            terms = audit.formula(sn, proposed)
                            result.update(execution='supported', value=str(audit.value(terms)),
                                          equals_original_dependency_closure=terms == old_terms,
                                          equals_corrected_dependency_closure=terms == corrected_terms,
                                          correct_on_current_original_inputs=(audit.value(terms) == corrected_value
                                                                             if corrected_terms is not None else None))
                        except Unsupported as error:
                            result.update(execution='unsupported', reason=str(error))
                    record[model] = result
                records.append(record)
    summary = {}
    for model in ('paper', 'dev11'):
        summary[model] = {}
        for label in ('formula_defect_candidate', 'blank_boundary_control'):
            group = [r for r in records if r['label'] == label]
            realized = [r for r in group if r['realized_output_error'] is True]
            summary[model][label] = dict(
                cells=len(group), realized_output_errors=len(realized),
                execution_unknown=sum(r['execution'] != 'supported' for r in group),
                changed_suggestions=sum(r[model]['changed_suggestion'] for r in group),
                exact_revision_matches=sum(r[model]['exact_revision_match'] for r in group),
                equivalent_corrected_formula_suggestions=sum(r[model].get('equals_corrected_dependency_closure') is True
                                                             for r in group),
                correct_suggestions_on_realized_errors=sum(
                    r[model].get('correct_on_current_original_inputs') is True for r in realized),
                suggestions_equal_original_closure=sum(r[model].get('equals_original_dependency_closure') is True
                                                       for r in group))
    result = dict(experiment='E1 frozen baseline; fixed original inputs and dependencies',
                  prediction_lock_verified=True, prediction_sha256=sha(prediction_path),
                  source_hashes={s: sha(RAW / f'{s}.xlsx') for s in ('before', 'after')},
                  scorer_sha256=sha(Path(__file__)), pair_change_counts=dict(changes),
                  labels_read_by_predictor=[], summary=summary, records=records,
                  precision_all_suggestions=None, final_validation_eligible=False,
                  limits='Local arithmetic/dependency audit only. Formula defects with zero current effect are separate; '
                         'all unknown/unscored model suggestions remain unscored. No new model or rule evaluated.')
    with args.output.open('x') as handle:
        json.dump(result, handle, ensure_ascii=False, sort_keys=True, indent=2)
        handle.write('\n')
    print(json.dumps(dict(pair_change_counts=dict(changes), summary=summary), indent=2))


if __name__ == '__main__':
    main()
