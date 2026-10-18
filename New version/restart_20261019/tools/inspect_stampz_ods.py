#!/usr/bin/env python3
"""Address-aware ODS formula diff and original conversion-fidelity audit."""
from pathlib import Path
import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.utils.cell import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'source_evidence/raw/stampz'
OUT = ROOT / 'source_evidence/stampz'
NS = dict(t='urn:oasis:names:tc:opendocument:xmlns:table:1.0',
          o='urn:oasis:names:tc:opendocument:xmlns:office:1.0',
          text='urn:oasis:names:tc:opendocument:xmlns:text:1.0')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inventory(path):
    with zipfile.ZipFile(path) as archive:
        tree = ET.fromstring(archive.read('content.xml'))
    result, formulas, captions = {}, {}, {}
    for table in tree.findall('.//t:table', NS):
        name = table.attrib['{' + NS['t'] + '}name']
        row_index = 1
        for row in table.findall('t:table-row', NS):
            repeated_rows = int(row.attrib.get('{' + NS['t'] + '}number-rows-repeated', '1'))
            col_index = 1
            for cell in row:
                if cell.tag not in ['{' + NS['t'] + '}table-cell', '{' + NS['t'] + '}covered-table-cell']:
                    continue
                repeated_cols = int(cell.attrib.get('{' + NS['t'] + '}number-columns-repeated', '1'))
                formula = cell.attrib.get('{' + NS['t'] + '}formula')
                text = '\n'.join(''.join(p.itertext()) for p in cell.findall('text:p', NS))
                kind = cell.attrib.get('{' + NS['o'] + '}value-type')
                value = cell.attrib.get('{' + NS['o'] + '}value')
                if formula or text or value is not None:
                    assert repeated_rows * repeated_cols <= 10000, 'bounded nonempty expansion'
                    assert not formula or repeated_cols * repeated_rows == 1, 'repeated formula requires relative translation'
                    for r in range(row_index, row_index + repeated_rows):
                        for c in range(col_index, col_index + repeated_cols):
                            key = name + '!' + get_column_letter(c) + str(r)
                            result[key] = dict(formula=formula, text=text, type=kind, value=value)
                            if formula:
                                formulas[key] = formula
                            elif text:
                                captions[key] = text
                col_index += repeated_cols
            row_index += repeated_rows
    return dict(cells=result, formulas=formulas, captions=captions)


def excel_formula(formula):
    assert formula.startswith('of:=')
    pieces = re.split(r'("(?:[^"]|"")*")', formula[3:])
    for index in range(0, len(pieces), 2):
        part = pieces[index].replace(';', ',')
        part = re.sub(r'\[\.(\$?[A-Z]+\$?\d+)(?::\.(\$?[A-Z]+\$?\d+))?\]',
                      lambda match: match[1] + (':' + match[2] if match[2] else ''), part)
        assert '[' not in part and ']' not in part
        pieces[index] = part
    return ''.join(pieces)


def main():
    target = OUT / 'ods_inventory_audit.json'
    assert not target.exists()
    lock = json.loads((OUT / 'prediction_lock.json').read_text())
    for relative, digest in lock['files'].items():
        assert sha(ROOT / relative) == digest
    before, after = [inventory(RAW / (version + '.ods')) for version in ('before', 'after')]
    workbook = openpyxl.load_workbook(RAW / 'before.xlsx', data_only=False)
    formulas = {sheet.title + '!' + cell.coordinate: cell.value for sheet in workbook
                for row in sheet for cell in row if cell.data_type == 'f'}
    converted = {key: excel_formula(value) for key, value in before['formulas'].items()}
    formula_differences = [key for key in converted if converted[key] != formulas.get(key)]
    changed = [dict(cell=key, before=before['formulas'].get(key), after=after['formulas'].get(key))
               for key in sorted(before['formulas'].keys() | after['formulas'].keys())
               if before['formulas'].get(key) != after['formulas'].get(key)]
    nonformula_changes = [key for key in before['cells'].keys() | after['cells'].keys()
                         if key not in before['formulas'] and key not in after['formulas']
                         and before['cells'].get(key) != after['cells'].get(key)]
    # Keep all third-party contents and changes local; upload only counts/qualification.
    with (RAW / 'ods_inventory.local.json').open('x') as handle:
        json.dump(dict(before=before, after=after, formula_changes=changed,
                       nonformula_changes=nonformula_changes), handle, indent=2)
    prediction = json.loads((RAW / 'prediction.json').read_text())
    summary = dict(before_formula_cells=len(before['formulas']), after_formula_cells=len(after['formulas']),
        formula_changes=len(changed), nonformula_changes=len(nonformula_changes),
        converted_original_formula_count=len(formulas),
        converted_formula_address_set_matches=set(converted) == set(formulas),
        converted_formula_text_mismatch_count=len(formula_differences),
        conversion_scalar_fidelity_verified=False,
        proposals={name:sum(bool(row.get('suggested_formula')) and row['suggested_formula'] != prediction['observed_formulas'][row['cell']]
                   for row in prediction[name]['rows']) for name in ['paper','dev11']},
        formula_change_coordinates=[row['cell'] for row in changed],
        original_prediction_lock_sha256=sha(OUT / 'prediction_lock.json'),
        original_ods_sha256=sha(RAW / 'before.ods'), after_ods_sha256=sha(RAW / 'after.ods'),
        certified_realized_natural_errors=0, business_truth_status='pending independent semantic and scalar-fidelity audit',
        final_validation_eligible=False)
    with target.open('x') as handle:
        json.dump(summary, handle, indent=2)
        handle.write('\n')
    print(json.dumps({key:value for key,value in summary.items() if key!='formula_change_coordinates'}, indent=2))


if __name__ == '__main__':
    main()
