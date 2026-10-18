#!/usr/bin/env python3
"""Prepare locked E1 Calc copies, or verify independent results without publishing formulas."""
import argparse
import json
import math
from pathlib import Path

import openpyxl

from score_e1_metrics import ROOT, RAW, OUT, sha, read_json

LOCAL = RAW / 'calc_audit'
LEGACY_WORKER = ROOT.parent / 'followup_20261014/scripts/libreoffice_reference_worker.py'


def save(path, payload):
    with path.open('x') as handle:
        json.dump(payload, handle, ensure_ascii=False, sort_keys=True, indent=2)
        handle.write('\n')


def prepare():
    LOCAL.mkdir(exist_ok=False)
    score = read_json(OUT / 'E1_BASELINE_AUDIT_V2.json')
    predictions = read_json(RAW / 'original_run/tasks/metrics/prediction.json')
    original = openpyxl.load_workbook(RAW / 'before.xlsx', data_only=False)
    revised = openpyxl.load_workbook(RAW / 'after.xlsx', data_only=False)
    proposals = {r['cell']: r.get('suggested_formula') for r in predictions['paper']['rows']}
    tasks = []
    for variant in ('original', 'corrected', 'paper_suggestions'):
        workbook = openpyxl.load_workbook(RAW / 'before.xlsx', data_only=False)
        expected = {}
        for row in score['records']:
            sheet, address = row['cell'].rsplit('!', 1)
            if variant == 'corrected':
                workbook[sheet][address] = revised[sheet][address].value
            elif variant == 'paper_suggestions' and row['paper']['changed_suggestion']:
                workbook[sheet][address] = proposals[row['cell']]
            if variant == 'original':
                expected[row['cell']] = float(row['original_value'])
            elif variant == 'corrected':
                expected[row['cell']] = float(row['corrected_value_on_original_inputs'])
            else:
                expected[row['cell']] = float(row['paper'].get('value', row['original_value']))
        modified = set()
        for sheet in original:
            for row in sheet:
                for cell in row:
                    new = workbook[sheet.title][cell.coordinate].value
                    if new != cell.value:
                        assert cell.data_type == 'f' and isinstance(new, str) and new.startswith('=')
                        modified.add(sheet.title + '!' + cell.coordinate)
        assert modified <= {r['cell'] for r in score['records']}
        workbook_path = LOCAL / f'{variant}.xlsx'
        workbook.save(workbook_path)
        # Validate persisted copies as well as in-memory construction.
        persisted = openpyxl.load_workbook(workbook_path, data_only=False)
        for sheet in original:
            for row in sheet:
                for cell in row:
                    assert persisted[sheet.title][cell.coordinate].value == workbook[sheet.title][cell.coordinate].value
        request = LOCAL / f'{variant}_request.json'
        save(request, dict(input=str(workbook_path), cells=[r['cell'] for r in score['records']]))
        save(LOCAL / f'{variant}_expected.json', expected)
        tasks.append(dict(id=variant, argv=['/usr/bin/python3', str(LEGACY_WORKER), str(request), 'result.json'],
                          outputs=['result.json']))
    manifest = LOCAL / 'tasks.jsonl'
    with manifest.open('x') as handle:
        for task in tasks:
            handle.write(json.dumps(task) + '\n')
    files = list(LOCAL.iterdir()) + [RAW / 'before.xlsx', RAW / 'after.xlsx', LEGACY_WORKER,
                                    OUT / 'E1_BASELINE_AUDIT_V2.json', ROOT / 'E1_EXECUTION_AMENDMENT.md',
                                    Path(__file__).resolve()]
    save(LOCAL / 'input_lock.json', dict(files={str(p): sha(p) for p in files}))
    print(manifest)


def verify():
    lock = read_json(LOCAL / 'input_lock.json')
    for path, expected in lock['files'].items():
        assert sha(Path(path)) == expected, path
    variants = {}
    for variant in ('original', 'corrected', 'paper_suggestions'):
        result_path = LOCAL / 'run/tasks' / variant / 'result.json'
        result = read_json(result_path)
        assert result['status'] == 'complete' and result['forced_calculate_all']
        assert result['read_only'] and not result['workbook_saved']
        expected = read_json(LOCAL / f'{variant}_expected.json')
        failures = []
        assert set(result['cells']) == set(expected)
        for cell, value in expected.items():
            observed = result['cells'][cell]
            if observed['kind'] != 'number' or not math.isclose(observed['value'], value, rel_tol=1e-9, abs_tol=1e-9):
                failures.append(dict(cell=cell, expected=value, observed_kind=observed['kind'],
                                     observed_value=observed.get('value')))
        variants[variant] = dict(compared=len(expected), agreements=len(expected) - len(failures),
                                 failures=failures, result_sha256=sha(result_path))
    runtime = read_json(LOCAL / 'run/runtime.json')
    receipt = dict(input_lock_sha256=sha(LOCAL / 'input_lock.json'), variants=variants,
                   runtime=runtime, final_validation_eligible=False,
                   all_agree=all(not v['failures'] for v in variants.values()),
                   note='Independent Calc on copies; original scalar inputs fixed; macros/links disabled; '
                        'no saved Calc workbook, no model changed, 3 tasks do not establish 8-worker speedup.')
    save(OUT / 'E1_CALC_VERIFICATION.json', receipt)
    print(json.dumps(variants, indent=2))
    if not receipt['all_agree']:
        raise SystemExit('Calc disagrees; retain failures and do not certify scores')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('stage', choices=['prepare', 'verify'])
    args = parser.parse_args()
    prepare() if args.stage == 'prepare' else verify()
