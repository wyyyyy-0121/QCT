#!/usr/bin/env python3
"""Prepare and run one locked 20-worker source-equation experiment."""
import argparse
from copy import deepcopy
import datetime as dt
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook

from resume_spreadsheetbench2_after_pilot import ROOT, EVIDENCE, RUNNER, verify_frozen
from spreadsheetbench2_scoring import sha256, verify_prediction_lock

EXPERIMENT = EVIDENCE / 'model_m7_source'
EPA = EVIDENCE.parent / 'epa_pairs/xlsx_export/oxidizers_before.xlsx'
MANUAL = EVIDENCE.parent / 'epa_manuals/oxidizers.txt'
PDF = EVIDENCE.parent / 'epa_manuals/oxidizers.pdf'


def write(path, value):
    with path.open('x') as handle:
        json.dump(value, handle, ensure_ascii=False, sort_keys=True, indent=2)
        handle.write('\n')


def verify(lock):
    for path, expected in lock['files_sha256'].items():
        if sha256(Path(path)) != expected:
            raise SystemExit('M7 locked artifact changed: ' + path)


def prepare():
    original, inputs = verify_prediction_lock(EVIDENCE)
    verify_frozen(inputs)
    if sys.executable != original['runtime_executable']:
        raise SystemExit('Use locked interpreter')
    if sha256(EPA) != '0acb06582e7f898b5bcbfe262b2f55294bd4b3a6d2eab63461f03c89ad86358c':
        raise SystemExit('EPA original changed')
    if sha256(PDF) != 'ad2d2ba6e14b57e7fe95999f6a9085fad783acc80fcee6d704a9541ceb818977':
        raise SystemExit('EPA manual changed')
    EXPERIMENT.mkdir(exist_ok=False)
    for name in ('requests', 'workbooks', 'scoring'):
        (EXPERIMENT / name).mkdir()
    common = dict(sheet='Design Parameters', cell='D33', symbols={'q': {'units': {'concentration': 1}},
                  'd': {'units': {'concentration': 1}}}, output_units={}, references={'D30': 'q', 'D31': 'd'},
                  aliases=[dict(text='Total Combustible Conc. In Mixture', symbol='q'), dict(text='LELmix', symbol='d')])
    adjusted = deepcopy(common)
    adjusted.update(cell='D35', references={'C18': 'q', 'S28': 'd'},
                    aliases=[dict(text='Total Adjusted Conc. With Dilution Air', symbol='q'), dict(text='LELmix', symbol='d')],
                    branch={'guard': 'D33<25', 'inactive': '"Not Applicable"'})
    rows, expected, register, files = [], {}, [], [EPA, MANUAL, PDF]

    def add(name, base, source_cell, status, reason, kind='real_source_fragment', workbook=EPA,
            source_sheet='Design Parameters', locator=None, **changes):
        request = deepcopy(base)
        request.update(changes)
        request.update(workbook=str(workbook), workbook_sha256=sha256(workbook),
                       source=locator or dict(kind='cell', sheet=source_sheet, cell=source_cell))
        path = EXPERIMENT / 'requests' / (name + '.json')
        write(path, request)
        files.append(path)
        rows.append(dict(id=name, argv=[sys.executable, str(ROOT / 'tools/source_relations_m7.py'), str(path),
                                       '--output', 'prediction.json'], outputs=['prediction.json'], timeout_seconds=120))
        expected[name] = dict(status=status, reason=reason, kind=kind,
                              extracted=name in {'epa_ratio', 'epa_adjusted', 'injected_ratio', 'injected_adjusted', 'missing_mapping'},
                              repair=status == 'propose')
        if request['source']['kind'] == 'cell':
            book = load_workbook(workbook, read_only=True, data_only=False)
            cell = book[request['source']['sheet']][request['source']['cell']]
            text, data_type = cell.value, cell.data_type
            book.close()
        else:
            loc = request['source']
            text = '\n'.join(Path(loc['path']).read_text().splitlines()[loc['start'] - 1:loc['end']])
            data_type = 'text_document'
        register.append(dict(id=name, kind=kind, source=request['source'], source_workbook_sha256=sha256(workbook),
                             verbatim=text, data_type=data_type, source_is_exposed=True,
                             annotation_origin='Same implementing agent; no independent human annotation',
                             mapping_origin='Manually supplied; not inferred or independently certified',
                             declaration=request))

    add('epa_ratio', common, 'B33', 'abstain', 'original_satisfies_contract')
    add('epa_adjusted', adjusted, 'B35', 'abstain', 'original_satisfies_contract')
    add('epa_flow_description', common, 'B36', 'unsupported', 'dynamic_or_executable_source_formula', cell='D36')
    for name, address, reason in [
        ('epa_dilution_sum', 'B34', 'source_notation_unsupported'),
        ('epa_lel_sum', 'B31', 'source_notation_unsupported'),
        ('epa_fan_exponent', 'B38', 'ambiguous_exponent_typography'),
        ('epa_approximate', 'B39', 'source_notation_unsupported'),
        ('epa_units_only', 'E36', 'no_explicit_rhs_equation')]:
        add(name, common, address, 'unsupported', reason)
    add('epa_readme', common, 'B28', 'unsupported', 'no_explicit_rhs_equation', source_sheet='Read Me', cell='D36')
    add('epa_manual_dilution', common, None, 'unsupported', 'no_explicit_rhs_equation', cell='D36',
        locator=dict(kind='text_lines', path=str(MANUAL), sha256=sha256(MANUAL), start=1523, end=1526,
                     pdf_sha256=sha256(PDF), printed_page='2-29'))

    # Real text is untouched for the two formula-fault controls; original is never saved.
    for name, base, text_cell, address, replacement in [
        ('injected_ratio', common, 'B33', 'D33', '=(D30/D31)*10'),
        ('injected_adjusted', adjusted, 'B35', 'D35', '=IF(D33<25, "Not Applicable", (C18/S28)*10)')]:
        book = load_workbook(EPA)
        book['Design Parameters'][address] = replacement
        path = EXPERIMENT / 'workbooks' / (name + '.xlsx')
        book.save(path)
        book.close()
        files.append(path)
        add(name, base, text_cell, 'propose', 'candidate_satisfies_explicit_relation', 'engineering_formula_fault', path)
    add('missing_mapping', common, 'B33', 'unsupported', 'missing_reference_mapping',
        'engineering_mapping_control', references={'D30': 'q'})
    add('ambiguous_alias', common, 'B33', 'unsupported', 'ambiguous_or_invalid_aliases',
        'engineering_mapping_control', aliases=common['aliases'] + [dict(text='LELmix', symbol='q')])
    for name, replacement, reason in [
        ('negated_source', '= Do not use (Total Combustible Conc. In Mixture/LELmix) × 100 =',
         'qualified_source_text'),
        ('conflicting_source', '= (Total Combustible Conc. In Mixture/LELmix) × 100 = 50 =',
         'multiple_or_empty_equations')]:
        book = load_workbook(EPA)
        cell = book['Design Parameters']['B33']
        cell.value, cell.data_type = replacement, 's'
        path = EXPERIMENT / 'workbooks' / (name + '.xlsx')
        book.save(path)
        book.close()
        files.append(path)
        add(name, common, 'B33', 'unsupported', reason, 'engineering_text_perturbation', path)
    if len(rows) != 16:
        raise SystemExit('Case count changed')
    manifest = EXPERIMENT / 'manifest.jsonl'
    with manifest.open('x') as handle:
        handle.write(''.join(json.dumps(row) + '\n' for row in rows))
    key, source_register = EXPERIMENT / 'scoring/expected.json', EXPERIMENT / 'source_register.json'
    write(key, expected)
    write(source_register, register)
    files += [key, source_register, manifest, Path(__file__), ROOT / 'MODEL_M7_SOURCE_PLAN.md', RUNNER]
    files += [ROOT / 'tools' / name for name in ['source_relations_m7.py', 'relation_candidates_m6.py',
              'e3_contract.py', 'trace_paper_candidates.py', 'candidate_gate_m1.py', 'score_source_m7.py',
              'score_relation_m6.py', 'test_source_relations_m7.py']]
    write(EXPERIMENT / 'input_lock.json', dict(created_utc=dt.datetime.now(dt.timezone.utc).isoformat(),
          files_sha256={str(path): sha256(path) for path in files}, requested_workers=20, task_count=16,
          final_validation_eligible=False, automatic_field_mapping=False,
          paper_input_lock_sha256=sha256(EVIDENCE / 'input_model_lock.json')))
    print('M7 INPUTS LOCKED', len(rows))


def run():
    original, inputs = verify_prediction_lock(EVIDENCE)
    verify_frozen(inputs)
    if sys.executable != original['runtime_executable']:
        raise SystemExit('Use locked interpreter')
    lock_path = EXPERIMENT / 'input_lock.json'
    lock = json.loads(lock_path.read_text())
    verify(lock)
    if (EXPERIMENT / 'runs').exists() or (EXPERIMENT / 'output_lock.json').exists():
        raise SystemExit('Preserve existing M7 execution')
    subprocess.run([sys.executable, str(RUNNER), '--manifest', str(EXPERIMENT / 'manifest.jsonl'),
                    '--output-dir', str(EXPERIMENT / 'runs'), '--workers', '20',
                    '--worker-memory-mib', '475', '--timeout', '120'], cwd=ROOT, check=False)
    runtime = json.loads((EXPERIMENT / 'runs/runtime.json').read_text())
    verify(lock)
    verify_frozen(inputs)
    if runtime['requested_worker_count'] != 20 or len(runtime['tasks']) != 16:
        raise SystemExit('Concurrency/task accounting changed')
    outputs = []
    for row in runtime['tasks']:
        path = Path(row.get('cwd', '')) / 'prediction.json'
        if row['status'] == 'complete' and sha256(path) != row['output_sha256']['prediction.json']:
            raise SystemExit('Prediction changed')
        outputs.append(dict(id=row['id'], status=row['status'], path=str(path) if path.exists() else None,
                            sha256=sha256(path) if path.exists() else None))
    write(EXPERIMENT / 'output_lock.json', dict(input_lock_sha256=sha256(lock_path),
          runtime_sha256=sha256(EXPERIMENT / 'runs/runtime.json'), outputs=outputs, scoring_may_begin=True))
    print('M7 TERMINAL OUTPUTS LOCKED', len(outputs))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('stage', choices=('prepare', 'run'))
    args = parser.parse_args()
    (prepare if args.stage == 'prepare' else run)()
