"""Create a Chinese reading copy from locked results, without editing model code."""

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
PROGRESS = BASE.parent
NAME = "HybridGuard_当前表现与后续研究建议_20261014"


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def percent(value):
    return f"{value * 100:.2f}%"


def set_font(font, size):
    font.name = "Noto Sans CJK SC"
    font.size = Pt(size)
    font._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")


def main():
    dev = read(PROGRESS / "20261014_v519_current_model/METRICS_SUMMARY.json")
    evidence = read(BASE / "DERIVED_RESULT_AUDIT.json")
    enron = read(PROGRESS / "20261014_v519_enron_retrospective/RESULT.json")["summaries"]
    overall = {r["backend"]: r for r in dev["strata"] if r["kind"] == "ALL"}
    candidates = evidence["development_candidates"]

    pages = [
        [
            ("title", "HybridGuard：当前表现与后续研究建议"),
            ("p", "2026-10-14｜当前模型 V5.1.9-development｜先评估，再适配新数据，最后修改"),
            ("h", "1. 当前模型跑完了什么"),
            ("p", "51 份开发工作簿 × 2 个定位后端 × 6 种模式，共 612 次诊断。源码与预测锁核对通过，282 个唯一解证书重放通过，结果与历史记录一致。旧模型代码未修改。"),
            ("table", [
                ["开发集指标", "V4 后端", "V5.1.1 后端"],
                ["正确接受覆盖率", f"540/735 = {percent(overall['v4']['repair_coverage'])}", f"540/735 = {percent(overall['v511']['repair_coverage'])}"],
                ["接受建议精确率", "540/540 = 100%", "540/540 = 100%"],
                ["正确格错误候选", "3,408", "27"],
                ["全部候选精确匹配比例", percent(candidates['v4']['candidate_precision']), percent(candidates['v511']['candidate_precision'])],
                ["宏平均 ordered AP", f"{overall['v4']['macro_ordered_ap']:.4f}", f"{overall['v511']['macro_ordered_ap']:.4f}"],
                ["宏平均 Recall@10", percent(overall['v4']['macro_recall_at_10']), percent(overall['v511']['macro_recall_at_10'])],
            ]),
            ("p", "候选和最终接受是两层输出。3,408/27 是候选误建议格数，系统没有把它们写回工作簿。100% 只适用于已知生成器开发集的接受建议，不能当作真实业务精确率。51 份输入都含错误，仍需全干净控制。"),
            ("h", "2. Enron 回顾性定位"),
            ("table", [
                ["180 秒/工作簿/后端", "V4", "V5.1.1"],
                ["成功工作簿", "21/25", "24/25"],
                ["Top10 事件命中", f"17/30 = {percent(enron['v4']['top10_rate_all_eligible'])}", f"16/30 = {percent(enron['v511']['top10_rate_all_eligible'])}"],
                ["超时", "4 份", "1 份"],
            ]),
            ("p", "50 个任务全部尝试，45 个成功、5 个超时；超时按未命中计入。Top10 指一个事件的任一标注源格进入前十。Enron 缺少正确公式和完整正确格真值，无法算修复精确率。"),
            ("p", "两个后端各保留 8 个 stable_copy_template 接受决策，正确性未知。后续需对这一路输出单独门控。当前应保留两个后端作外部对照。"),
        ],
        [
            ("title", "新数据：下载完成，适配和评分待做"),
            ("p", "本轮已经检查 ZIP 完整性、标注引用和 Git blob 哈希。包内数量不能直接当作独立模板数。"),
            ("table", [
                ["数据", "本地核查", "建议用途"],
                ["Modified EUSES", "705 配置；2,060 工作簿", "优先配对原表，建立注入修复真值"],
                ["Info1", "119 学生表；2 个模板", "真实学生错误补充，按模板隔离"],
                ["Integer", "231 配置；404 源格标签", "整数域补充，核对非公式标签"],
                ["Payroll/Gradebook", "349 配置；565 个用户判断错误标记", "不可靠用户测试和约束压力"],
                ["FoRepBench", "618 生成样本；59 论坛种子", "显式报错修复补充"],
                ["SpreadsheetBench 2", "文档和论文已取得；数据文件未取得", "核对 Debugging 子集后再决定"],
            ]),
            ("h", "3. 数据要先解决的四件事"),
            ("p", "① 官网 EUSES 记录 576 个单错误版本，但当前 ZIP 有 705 个配置。先确认包版本、重复、家族和原始表映射。"),
            ("p", "② properties 主要给错误源格和输出测试值，没有直接给正确源公式。先把原表与故障表配对，区分自然错误、注入错误和未知公式。"),
            ("p", "③ 在副本区转换 xls→xlsx，记录转换版本、原文件哈希、公式数量和文本差异，避免转换造成新错误。"),
            ("p", "④ 按模板家族或近重复簇划分开发与留出集；同时补充合法例外和全干净控制。冻结后先锁当前模型预测，再读标签评分。"),
            ("h", "任务边界"),
            ("p", "FoRepBench 的 618 个生成样本和 59 个种子都标有运行时错误。它提供错误公式与上下文，不能替代静默源错误定位主实验。SpreadsheetBench 2 的 Debugging 子集仍待核对。"),
            ("p", "EUSES：spreadsheets.sai.tugraz.at/index.php/corpora-for-benchmarking/euses/\nFoRepBench：github.com/microsoft/prose-benchmarks/tree/d17e278092c59ec6faaa2d88730bdcbe48bb95f2/FoRepBench"),
        ],
        [
            ("title", "后续策略：先建立外部基线，再改选择层"),
            ("h", "4. 建议研究主题"),
            ("p", "面向合法业务例外与相关跨表复制错误的选择性公式修复。主要检验：在相近正确修复覆盖率下，角色、例外和来源检查能否减少错误建议。"),
            ("h", "5. 具体执行顺序"),
            ("p", "第一步：先完成 EUSES 原表配对、去重、公式保真和试标，标出 error、correct、legitimate_exception、unknown。"),
            ("p", "第二步：用未修改的 V4、V5.1.1 跑外部 baseline。两个后端同时报告，保留失败和超时。"),
            ("p", "第三步：保留候选生成器，在选择层加入汇总/明细角色、合法例外和复制来源检查；先验证简单规则。"),
            ("p", "第四步：分别输出候选、待复核、拒绝和不支持状态，记录理由。批准约束下的联合修复单独评价，保持错误标签与约束输入隔离。"),
            ("p", "第五步：开发集调规则，留出集冻结后锁预测再评分。比较候选精确率、接受精确率、覆盖率、风险-覆盖曲线、未知比例和耗时。"),
            ("p", "开发目标建议：V5.1.1 正确格错误候选由 27 降至最多 5，正确接受覆盖下降不超过 5 个百分点，接受不安全组仍为零。该目标尚未达成，需实验前登记。"),
            ("h", "6. 目录与仓库"),
            ("p", "项目已迁移至 /home/ayaka/hybridguard；/home/ayaka/QCT 保留兼容链接。今后的输入副本、协议、结果和报告都放在 New version。/home/ayaka/code 未修改。"),
            ("p", "继续使用 github.com/wyyyyy-0121/QCT，建议新建 research/hybridguard-followup 分支。当前没有创建分支、提交或推送。"),
            ("p", "完整报告目录：\\\\wsl.localhost\\Ubuntu-26.04\\home\\ayaka\\hybridguard\\New version\n先看 README.md，再看 CURRENT_MODEL_RESULT.md、DATASET_REVIEW.md、FOLLOWUP_RECOMMENDATIONS.md。"),
            ("p", "原论文 0.3.0 与当前 V5.1.9 是不同版本。本轮运行的是当前 V5.1.9；原论文基线和旧负结果保留为历史材料。尚无新门控方法结果或新数据修复成绩。"),
        ],
    ]

    doc = Document()
    doc.core_properties.title = "HybridGuard 当前表现与后续研究建议"
    doc.core_properties.subject = "2026-10-14 已完成评估、新数据审计与后续研究安排"
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
    section.header_distance, section.footer_distance = Cm(0.6), Cm(0.65)
    normal = doc.styles["Normal"]
    set_font(normal.font, 10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.13
    for style_name, size in [("Title", 18), ("Heading 1", 12)]:
        style = doc.styles[style_name]
        set_font(style.font, size)
        style.font.color.rgb = RGBColor.from_string("203A4E")
        style.paragraph_format.space_before = Pt(6 if style_name == "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(6)
    header = section.header.paragraphs[0]
    header.text = "HybridGuard · 当前评估与后续研究 · 2026-10-14"
    set_font(header.runs[0].font, 8)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    expected_blocks, txt_pages = [], []
    for page_index, blocks in enumerate(pages):
        if page_index:
            doc.add_page_break()
        txt_blocks = []
        for kind, value in blocks:
            if kind == "table":
                table = doc.add_table(rows=0, cols=len(value[0]))
                table.style = "Table Grid"
                table.autofit = False
                for col, width in zip(table.columns, [5.4, 5.95, 5.95]):
                    col.width = Cm(width)
                for index, row_values in enumerate(value):
                    row = table.add_row()
                    keep = OxmlElement("w:cantSplit")
                    row._tr.get_or_add_trPr().append(keep)
                    if not index:
                        repeat = OxmlElement("w:tblHeader")
                        row._tr.get_or_add_trPr().append(repeat)
                    for cell, text in zip(row.cells, row_values):
                        cell.text = text
                        if not index:
                            shading = OxmlElement("w:shd")
                            shading.set(qn("w:fill"), "E7EFF4")
                            cell._tc.get_or_add_tcPr().append(shading)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(3)
                            paragraph.paragraph_format.space_before = Pt(3)
                            paragraph.paragraph_format.line_spacing = 1.04
                            for run in paragraph.runs:
                                set_font(run.font, 9)
                                run.bold = not index
                    expected_blocks.append(" | ".join(row_values))
                txt_blocks.append("\n".join(" | ".join(row) for row in value))
            else:
                style = {"title": "Title", "h": "Heading 1", "p": "Normal"}[kind]
                doc.add_paragraph(value, style)
                expected_blocks.append(value)
                txt_blocks.append(value)
        txt_pages.append("\n\n".join(txt_blocks))

    target = BASE / f"{NAME}.docx"
    if target.exists():
        raise FileExistsError(target)
    doc.save(target)
    text_target = BASE / f"{NAME}.txt"
    text_target.write_text("\n\n".join(txt_pages) + "\n", encoding="utf-8-sig")

    # Compare body paragraphs and table rows in document order with the source blocks.
    reopened = Document(target)
    actual_blocks = []
    for item in reopened.element.body:
        if item.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            text = Paragraph(item, reopened).text
            if text.strip():
                actual_blocks.append(text)
        elif item.tag == qn("w:tbl"):
            from docx.table import Table
            actual_blocks.extend(" | ".join(cell.text for cell in row.cells)
                                 for row in Table(item, reopened).rows)
    assert actual_blocks == expected_blocks
    assert len(pages) == 3
    receipt = {
        "model": "V5.1.9-development", "planned_pdf_pages": len(pages),
        "docx_blocks_verified": len(expected_blocks), "new_dataset_model_scores": "not run",
        "outputs": {p.name: {"sha256": hashlib.sha256(p.read_bytes()).hexdigest(), "bytes": p.stat().st_size}
                    for p in (target, text_target)},
    }
    (BASE / "READING_COPY_BUILD.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(receipt, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
